from __future__ import annotations


import json


import math


import os


import re


import shutil


import subprocess


import tempfile


from dataclasses import dataclass, replace


from pathlib import Path


from typing import Any, Callable


from openpyxl import load_workbook


from openpyxl.workbook.properties import CalcProperties


from docx import Document


from pypdf import PdfReader


from pptx import Presentation


from runtime.grading.semantic import (
    contains_concept,
    ordered_semantic_list_matches,
    semantic_equal,
    semantic_value_matches,
)


from runtime.grading.hybrid_semantic import semantic_requirement


from runtime.grading.rubric import apply_reward_policy, attach_default_policy


from runtime.accounting_mcp.paths import resolve_project_root, resolve_seed_root


GOLD_PATH = Path(__file__).resolve().parent / "gold" / "tasks_001_025.json"


SEED_WORKSPACE = resolve_seed_root(__file__) / "sources"


SAMPLE_TASK_IDS = frozenset(
    {"task_001", "task_004", "task_015", "task_035", "task_068"}
)


TASK_GRADING_REVISIONS = {
    "task_001": {
        "id": "task-001-conditional-credit-effectiveness-v43",
        "effective_date": "2026-09-08",
        "basis": (
            "outcome-only professional request, explicit July 2 evidence cutoff, and neutral four-page controller memo; "
            "case facts remain discoverable in ordinary accounting, forecast, commercial, "
            "policy, ordinary AP batch, and close-tracker records without source-side accounting conclusions or curated exception cues; "
            "separately filed, late-received conditional credit correspondence must be reconciled by "
            "operative execution terms, both signatures, ordinary package-release and executed-return records, "
            "the referenced credit-memo delivery condition, document and receipt dates, scope, and accounting effect "
            "against the earlier commercial working records; "
            "deterministic checks cover artifact integrity and exact objective prerequisites; "
            "an unchanged canonical template association takes the deterministic fast path, "
            "while every edited, moved, or equivalently relabeled numeric field receives "
            "criterion-scoped semantic association review before credit; authority, accounting "
            "treatment, cutoff support, source use, and action sufficiency are always semantic; "
            "no finite row-label alias list decides reward, misleading labels and unlabeled "
            "number dumps receive no rescue, and blank starter labels open no semantic gate; "
            "June completeness uses practical AP, payroll, billing, and total controls rather "
            "than hidden transaction registers; ordinary ERP statuses remain realistic but "
            "lightly weighted; and the proposed entry follows the existing June 1 auto-reversal "
            "by re-establishing the gross June project balance rather than misgrading the "
            "May-to-June change as the journal to post"
            "; a pre-cutoff vendor invoice on page 4 of an ordinary ten-invoice AP scan batch must be matched by "
            "project, vendor, work package, PO reference, and amount to a separate PM "
            "Commitment Detail row and reconciled to its absence from the posted June ledger, "
            "producing a separate AP accrual and an equal ETC-to-incurred transfer without "
            "double counting EAC or contaminating the WIP journal; a later-filed ordinary "
            "contract-records batch must also be reconciled to the PM commitment and AP invoice "
            "to recognize that PO-2409-117 is already carried inside the PCO-011 mechanical-rework "
            "cost support, so the shared $46,800 is included once rather than restored through a "
            "second forecast overlay; the same ordinary contract-records batch also contains a "
            "separately filed $27,000 credit memorandum against previously billed work, while "
            "the signed agreement assigns the other $45,000 of its $72,000 adjustment to the "
            "remaining unbilled subcontract balance and conditions both components on delivery "
            "of the referenced memorandum to ARM Commercial Records; the ordinary July 2 receipt "
            "therefore requires Finance to exclude both credit components from the June 30 close "
            "without treating the source-to-ledger difference as a missing June posting; this "
            "same invoice is a normal cumulative mixed-period progress bill with $46,800 gross billings, "
            "a $13,003.31 prior draw already posted under the accounting system's internal document number, "
            "and $33,796.69 currently due; its gross work lines allocate $40,200 to the June field window "
            "and $6,600 to July 2, requiring the preparer to match the prior draw, accrue only $27,196.69 of "
            "unbilled June work, and remove both the stale prior-billing duplicate and current June transfer "
            "from the PM forecast; this ordinary cumulative-billing and timing "
            "evidence is carried in the invoice and vendor release rather than an exception schedule; this "
            "deepens the existing close judgment without changing "
            "the prompt, requested deliverable, memo topology, or required workstreams; section-normalized scoring "
            "keeps that core recommendation, calculation, and posting work materially above "
            "routine packaging and source-list credit without imposing a score cap; "
            "unchanged canonical primary fields with a conflicting numeric value fail "
            "objectively instead of being rescued by an alternate or sensitivity value, "
            "rate gates reject unrelated bare counts, strict semantic entailments reconcile "
            "logically incompatible criterion verdicts, and source-use prompts require the "
            "actual controlling association rather than a headline source citation; an output-first Opus audit "
            "also makes the distributed posting-release boundary explicit so equivalent project-contribution, "
            "consolidated-schedule tie, unposted-status, and pending-Controller-review wording receives credit"
        ),
    },
    "task_004": {
        "id": "task-004-close-status-release-judgment-v15",
        "effective_date": "2026-09-08",
        "basis": (
            "an outcome-only business prompt and neutral two-sheet starter require unprompted discovery "
            "of four project judgments from ordinary PM, commercial, policy, prior-close, and accounting "
            "records; the shared PM register retains the real amounts, statuses, timing, and correspondence "
            "but no longer exposes an inclusion flag that announces forecast treatment; row, column, sheet, "
            "label, and ordinary unit variation are accepted; exact amounts and formula behavior remain "
            "deterministic while edited review status, accounting rationale, source use, ownership, and "
            "portfolio release wording receive independent scoped semantic review; a recognized primary "
            "recommendation cannot be rescued by a contradictory hypothetical support row, and a clearly "
            "scoped portfolio-wide AP action can satisfy the same cutoff action without four repetitions; "
            "an ordinary ARM-2318 invoice exceeds the amount carried for its PO in the PM commitment detail, "
            "so the close must transfer only the forecast overlap, recognize the residual EAC increase, and "
            "apply the existing Controller-review threshold without any prompt hint; portfolio-wide policy "
            "provenance is included in each scoped source-use review without requiring repeated filenames; "
            "an ordinary close-chat row records that the Controller already checked the ARM-2318 lift/retest "
            "estimate and retained the excluded PCO treatment; Finance must distinguish that completed exception "
            "review from the still-open AP, commercial, and ordinary close-sign-off work rather than placing every "
            "project on the same hold; this deepens the existing release judgment without changing the prompt, "
            "starter, requested deliverable, monetary answer, or criterion inventory; "
            "linear section normalization gives the cutoff-adjusted close and release judgment business-"
            "material priority while preserving proportional, uncapped criterion-level signal; an output-first "
            "Opus audit additionally makes the source-authority and owned-action conjuncts explicit, accepts "
            "PCO-006 and DB-27 as names for the same pricing matter, and prevents partial source or action work "
            "from receiving full semantic credit"
        ),
    },
    "task_011": {
        "id": "task-011-operating-review-presentation-v17",
        "effective_date": "2026-08-10",
        "basis": (
            "five-slide Q2 service operating review with source-controlled "
            "finance facts, work-mix-adjusted callback judgment, scarce-slot "
            "and expected-value decisions, native editable PowerPoint charts "
            "and tables, explicit May-to-June KPI, qualifying callback-focus, "
            "peer-baseline, and expected-value input disclosure, objective readability checks, nine-section normalized "
            "materiality-weighted proportional scoring, a cross-source July "
            "branch-capacity support, a source-backed scarce service-flex portfolio "
            "decision with dispatcher, timing, authority and fallback constraints, "
            "a probability-weighted lock-versus-wait execution policy with a "
            "source-derived crossover and native decision chart, "
            "a development-versus-holdout dispatcher-signal validation, a frozen "
            "row-level out-of-time production-readiness test with late and missing "
            "signal controls, conditional "
            "access probabilities, expected-value-of-information and downside-risk "
            "deployment decision, a required formula-driven audit workbook with "
            "cross-sheet controls, source-equivalent labeled chart, table, short-ID, "
            "date, and matrix-workbook parsing with ambiguity guards, and no reward caps"
        ),
    },
    "task_013": {
        "id": "task-013-lender-liquidity-source-control-v10",
        "effective_date": "2026-08-10",
        "basis": (
            "June 30 source-authority and cutoff judgment for two public-owner "
            "assignments, base-certificate versus sensitivity treatment, debtor-specific "
            "evidence gaps and next actions, downside liquidity and borrowing-base "
            "availability judgment without double counting, Daniel and Natalie source "
            "hierarchy for the current-base rebuild, current-base versus prior-import "
            "source control, a quantified collateral-erosion stop/go decision, and "
            "an explicit response-only workspace boundary aligned to the integrity gate, "
            "with proportional atomic scoring without reward caps or prescribed response schema"
        ),
    },
    "task_014": {
        "id": "task-014-covenant-package-release-v7",
        "effective_date": "2026-08-10",
        "basis": (
            "five-note debt-service support, full-population and exception WIP "
            "review, executed-agreement source authority, one-for-one covenant "
            "deterioration capacity, flagged-project FCCR sensitivity, June 30 "
            "borrowing-base and public-owner assignment-cutoff review, contractual "
            "package-release controls, an August peak-liquidity reliance decision "
            "using the current-base cash and latest collateral sources, a "
            "quantified current-base versus prior-import source-control comparison, exact "
            "boolean compliance conclusions, and "
            "materiality-weighted proportional scoring without reward caps"
        ),
    },
    "task_015": {"basis":"a concise objective-led lender request requires independent discovery and reperformance from the executed agreement, posted accounting records, ordinary June AP cutoff evidence, and the current proposed-WIP support; the July 2 close-pro-forma basis follows the source-supported posting sequence and does not import an older WIP snapshot. Exact values are unit-scoped deterministic hard gates, aligned text-box matrices retain their visible row and column associations for criterion-scoped semantic review, and a binding-covenant conclusion requires its four like-for-like capacity prerequisites. The untouched five-slide starter receives zero credit, correct alternate layouts and normal units remain gradeable, central finance work dominates presentation mechanics, and all reward is proportional without task-level or section-level score caps","effective_date":"2026-09-04","id":"task-015-current-close-covenant-v14"},
    "task_016": {
        "id": "task-016-controller-journal-package-v15",
        "effective_date": "2026-08-13",
        "basis": (
            "source-derived cash, AR, AP, debt, and WIP close controls; bank-feed "
            "to GL to book-cash roll-forward with supported manual cash activity; "
            "May 31-to-June 30 posted-GL account-group roll-forwards and proposed "
            "WIP entry release control; controller source-authority close-release "
            "gate across posted income, cash, and WIP reversal evidence; the "
            "source-directed ARM-2522 AP cutoff reclassification, percentage-of-completion "
            "effect, no-double-count, and downstream-release boundary; the staged "
            "AP-accrual-then-WIP posting sequence, combined balanced entry, and "
            "operating-income bridge; four-project controller risk decisions, "
            "aggregate WIP risk bridge, close-release hold, and staged "
            "financial-statement accounting-equation bridge; pre-cutoff flash, "
            "cutoff-adjusted income-statement, and board-plan release bridge; "
            "controller-ready late-AP and June-WIP journal headers and lines; "
            "account-group and posted-to-pro-forma trial-balance rollforwards; "
            "retained-workpaper to open-subledger controls; project-level AR, AP, "
            "and WIP schedules; debt reclassification support; materially balanced "
            "proportional atomic scoring without reward caps"
        ),
    },
    "task_017": {
        "id": "task-017-working-capital-evidence-feasibility-v5",
        "effective_date": "2026-08-11",
        "basis": (
            "source-derived gross-to-actionable AR and AP bridge, whole-invoice and "
            "whole-bill theoretical sizing, controller-note and AP-cutoff feasibility "
            "screening, linear section normalization that prevents repeated support "
            "rows from overwhelming the executable cash decision, proportional "
            "atomic scoring, and no reward caps"
        ),
    },
    "task_018": {
        "id": "task-018-project-delivery-capacity-v9",
        "effective_date": "2026-08-13",
        "basis": (
            "full project-level signed-backlog support, H2-start versus post-H2 "
            "pipeline timing reconciliation, workbook-status-controlled award-dated "
            "capacity-review action queue, "
            "monthly weighted-load decision, full-population ordered pipeline release "
            "control with professional-equivalence status matching plus capacity, "
            "schedule, timing and forecast boundaries, "
            "plus a source-locked FY27 probability-adjusted backlog burn versus "
            "field-labor capacity release decision, full project and project-month "
            "delivery-driver schedules, project-level shortfall exposure and release "
            "controls, and proportional deterministic scoring "
            "without reward caps"
        ),
    },
    "task_019": {
        "id": "task-019-capex-release-liquidity-overlay-v3",
        "effective_date": "2026-08-09",
        "basis": (
            "request-level approval and timing support, release-status funding "
            "allocation, and committed versus non-binding indicative downside "
            "liquidity overlays"
        ),
    },
    "task_020": {
        "id": "task-020-labor-plan-authority-decision-v9",
        "effective_date": "2026-08-11",
        "basis": (
            "corrected-register hourly/salaried reconciliation, board-packet versus "
            "old-upload headcount control, department hourly-rate review, posted-run "
            "scope, accrual-to-register release controls, source-derived loaded-rate "
            "classification and close disposition, review-versus-tied sign-off "
            "exposure control, payroll-component estimating-to-register bridge, "
            "stale-upload versus board-target labor-plan authority and approved-wage "
            "assumption decision, "
            "board-plan overtime-leakage "
            "premium bridge and nonbooking management decision, proportional atomic "
            "scoring, and no reward caps"
        ),
    },
    "task_021": {
        "id": "task-021-capital-cash-source-use-control-v13",
        "effective_date": "2026-08-12",
        "basis": (
            "fixed-asset register and equipment-note reconciliation, approved-capex "
            "placement by existing asset class, request-level committed and indicative "
            "funding bridge, a July-to-June existing-note and indicative-line debt-service "
            "release control with unsupported-interest and commitment gates, "
            "request-level capital authorization and committed-versus-indicative "
            "downside-liquidity release decisions, a weekly authorized-release versus "
            "all-approved cash-timing overlay, a June 30 posted-balance versus future "
            "capital commitment recognition and release control, a request-level "
            "authorization evidence matrix tying decisions, amounts, capitalization, "
            "cash timing, financing and release status to controlling sources, a "
            "request-to-approval-to-authorization amount funnel, an approved-but-blocked "
            "weekly cash and revolver overlay split between financing and operating "
            "conditions, a weekly base and blocked-capital cash source/use financing "
            "rollforward proving draws, repayments, ending cash and revolver, explicit response-only "
            "boundary, and proportional atomic scoring without reward caps"
        ),
    },
    "task_022": {
        "id": "task-022-service-execution-release-v6",
        "effective_date": "2026-08-12",
        "basis": (
            "Controller-adjusted service population, source-derived branch capacity, "
            "replacement-versus-growth source boundary, gross fleet overlay, "
            "work-type pricing and callback review, technician-level productivity and "
            "median-versus-mean conservative capacity stress, plus constrained service "
            "portfolio, timing, holdout, and out-of-time deployment decisions, and proportional atomic scoring "
            "without reward caps"
        ),
    },
    "task_023": {
        "id": "task-023-lender-workbook-wip-reporting-control-v3",
        "effective_date": "2026-08-10",
        "basis": (
            "full-population proposed-WIP source authority, formula-driven WIP "
            "reconciliation, and source-backed lender draft and release controls"
        ),
    },
    "task_024": {
        "id": "task-024-capex-funding-views-v2",
        "effective_date": "2026-08-09",
        "basis": (
            "professionally weighted request decisions, committed-versus-indicative "
            "funding conclusions, and formula-driven support without fixed reward caps"
        ),
    },
    "task_025": {
        "id": "task-025-close-note-debtor-assignment-v1",
        "effective_date": "2026-08-09",
        "basis": (
            "close decision note uses the source-disclosed government-debtor assignment "
            "schedule and project-scoped dispute classification"
        ),
    },
    "task_027": {
        "id": "task-027-stable-scenario-model-v4",
        "effective_date": "2026-08-26",
        "basis": (
            "controller-visible FY27 source conventions plus formula-driven branch, "
            "EBITDA, free-cash-flow, weekly liquidity, lender leverage and fixed-charge "
            "coverage release decisions with proportional section-normalized finance "
            "and formula-lineage scoring, with professional row-label equivalence and "
            "deterministic repeatability"
        ),
    },
    "task_035": {"basis":"a concise business objective and neutral working model require source-version reconciliation, unit-aware signed-backlog and capacity analysis, priority-based execution, contract and earnings consequences, and a management recommendation without announcing the hidden difficulty; exact magnitudes, date/project facts, and formula results retain deterministic hard gates while every editable metric, period, scenario, sign, label, and decision association receives criterion-scoped semantic review, so finite aliases cannot reject correct professional wording or rescue a right value attached to the wrong business fact; hidden tie-breaks are eliminated and all sections score proportionally without reward caps","effective_date":"2026-09-04","id":"task-035-authentic-backlog-capacity-decision-v13"},
    "task_037": {
        "id": "task-037-professional-capital-model-v2",
        "effective_date": "2026-08-26",
        "basis": (
            "source-complete capital-allocation model with formula-linked project and "
            "portfolio economics, normal finance labels, deterministic numeric checks, "
            "materiality-balanced committee-decision scoring, and no hidden all-or-"
            "nothing model cap"
        ),
    },
    "task_055": {
        "id": "task-055-professional-accretion-model-v2",
        "effective_date": "2026-08-26",
        "basis": (
            "source-grounded acquisition accretion model accepting normal sources-and-"
            "uses and income-statement labels, conventional expense signs, and "
            "case-and-year-aware sensitivity intersections, a disclosed release rule, "
            "and formula-backed finance outcomes without phrase matching"
        ),
    },
    "task_061": {
        "id": "task-061-source-complete-tax-provision-v1",
        "effective_date": "2026-08-26",
        "basis": (
            "tax-provision model with every graded opening balance and estimated payment "
            "disclosed, conventional valuation-allowance presentation, and independent "
            "roll-forward checks that preserve the correct economics"
        ),
    },
    "task_068": {
        "basis": (
            "a concise executive request requires independent cross-source reconciliation "
            "and judgment without announcing the source conflicts or guidance chain; the "
            "current controller book remains a reporting consolidation while the Vista "
            "close history, operating evidence, owner support, Finance recovery review, "
            "forecast sensitivity, guidance cases, signed risk register, release policy, "
            "and published range remain single process-owned company records shared across "
            "the affected reporting tasks; the close history contains posted originals, "
            "reversals, revisions, parked and next-period batches, plus later June-period "
            "postings after the July 3 controller-package cutoff, but no source-side revenue "
            "presentation effect, EBITDA effect, branch subtotal, consolidated answer, "
            "mitigation aggregate, stress result, or release conclusion; current operating "
            "evidence is a 48-row record population, and the action-support and Finance-review "
            "files retain dated executed, approved, superseded, withdrawn, draft, challenged, "
            "and sensitivity-only history rather than one-row action answers; a preparer must "
            "apply source status and cutoff, aggregate the accepted evidence, join the approved "
            "rate and margin bases, and distinguish the final Finance view from same-day "
            "sensitivity work; retained earlier summaries are visibly superseded and "
            "numerically different; exact finance magnitudes are deterministic hard gates "
            "while edited labels, placement, sign conventions, actual-versus-plan roles, "
            "close-basis association, and business meaning receive criterion-scoped semantic "
            "review; the posted-to-current bridge, business-unit gross impacts, aggregate "
            "mitigation, signed-register probability-weighted downside ranking, and action "
            "narrative have criterion-specific semantic routes; "
            "accounting and decision work carry eighty percent of reward; seeded content "
            "earns no deliverable credit; and all criteria score proportionally without caps"
        ),
        "effective_date": "2026-09-05",
        "id": "task-068-authentic-executive-performance-decision-v20",
    },
    "task_072": {
        "id": "task-072-management-narrative-equivalence-v1",
        "effective_date": "2026-08-26",
        "basis": (
            "quarterly finance narrative graded from local paragraph and table context, "
            "including common million-dollar presentation and ordinary plan, variance, "
            "guidance, and operating-driver labels"
        ),
    },
    "task_073": {
        "id": "task-073-disclosed-pro-forma-debt-basis-v1",
        "effective_date": "2026-08-26",
        "basis": (
            "credit-metrics model with the controlling pro-forma acquisition-facility "
            "draw and document-authority hierarchy stated in the task, so posted and "
            "pro-forma debt bases cannot be confused"
        ),
    },
    "task_100": {
        "id": "task-100-board-language-equivalence-v2",
        "effective_date": "2026-08-26",
        "basis": (
            "source-grounded CFO board deck with deterministic finance checks and "
            "independent semantic rescue for normal executive labels, decisions, and "
            "status wording while unsupported conclusions remain uncredited and adjacent "
            "board-card booleans cannot contaminate one another"
        ),
    },
    "task_028": {
        "id": "task-028-opex-monthly-function-support-v1",
        "effective_date": "2026-08-09",
        "basis": (
            "month-by-function operating-expense support, event timing, source "
            "lineage, and proportional formula-driven scoring"
        ),
    },
    "task_029": {
        "id": "task-029-service-recovery-release-v8",
        "effective_date": "2026-08-16",
        "basis": (
            "professionally equivalent product volume-and-mix allocations, product-level "
            "price and SLA-rate support, management recovery sizing, and posted Service-"
            "revenue population reconciliation plus constrained action selection and "
            "single-action outage, six-case execution-scenario reoptimization, and priority-"
            "recovery release coverage plus scenario-specific priority-product action allocation without reward caps"
        ),
    },
    "task_030": {
        "id": "task-030-margin-recovery-release-v7",
        "effective_date": "2026-08-16",
        "basis": (
            "worktype-level gross-profit bridge, source-derived unfavorable-driver "
            "identification, upper-bound margin-recovery sizing, retained-draft "
            "assumption change control, constrained recovery action selection, "
            "single-action outage and six-case execution-scenario reoptimization, "
            "budget-margin recovery release coverage, professional basis-point precision, and proportional atomic scoring "
            "without reward caps"
        ),
    },
    "task_034": {
        "id": "task-034-cohort-release-integration-v8",
        "effective_date": "2026-08-17",
        "basis": (
            "cohort renewal forecast, constrained retention portfolio, single-action "
            "outages, three-case renewal execution reoptimization, and an integrated "
            "ARR-to-release bridge and full source-cohort renewal/downsell/escalation/proration support with proportional source-backed scoring and no reward caps"
        ),
    },
    "task_084": {
        "id": "task-084-fx-close-release-integration-v8",
        "effective_date": "2026-08-17",
        "basis": (
            "layered hedge economics, natural offsets, counterparty controls, constrained "
            "execution protection, single-action outages, and a residual-cost close-release "
            "bridge plus exposure-level hedge-layer and counterparty-cure support with proportional source-backed scoring and no reward caps"
        ),
    },
    "task_091": {
        "id": "task-091-bid-ceiling-release-v7",
        "effective_date": "2026-08-16",
        "basis": (
            "valuation triangulation, bid ceiling, constrained bid-protection portfolio, "
            "single-action outages, three-case bid execution reoptimization, and bid-ceiling release integration with "
            "proportional source-backed scoring and no reward caps"
        ),
    },
    "task_096": {
        "id": "task-096-tax-basis-close-release-v9",
        "effective_date": "2026-08-17",
        "basis": (
            "stock-versus-asset tax economics, constrained transaction-tax protection, "
            "single-action outages, three-case tax-close reoptimization, legal/consideration/NPV release integration, "
            "an asset-class tax-basis and close-cash release control, and the gross-up-by-tax-close negotiation release matrix with "
            "proportional source-backed scoring and no reward caps"
        ),
    },
    "task_097": {
        "id": "task-097-guidance-release-integration-v7",
        "effective_date": "2026-08-17",
        "basis": (
            "probability-weighted guidance, combined stress, constrained protection, "
            "single-action outages, scenario reoptimization, and protected-EBITDA release "
            "integration with proportional source-backed scoring and no reward caps"
        ),
    },
    "task_098": {
        "id": "task-098-kpi-definition-release-v6",
        "effective_date": "2026-08-17",
        "basis": (
            "restated KPI workbook, pending-scope adjustments, investor-versus-covenant "
            "definition authority, and a formula-linked two-basis release bridge with "
            "proportional source-backed scoring and no reward caps"
        ),
    },
    "task_099": {
        "id": "task-099-capital-return-release-integration-v8",
        "effective_date": "2026-08-17",
        "basis": (
            "repurchase-versus-debt-paydown economics, constrained execution protection, "
            "single-action outages, three-case capital-return reoptimization, and an integrated "
            "two-year cadence release bridge and candidate-by-year cash/debt/share/EPS/leverage roll-forward with proportional source-backed scoring and no reward caps"
        ),
    },
}


@dataclass(frozen=True)
class Criterion:
    id: str
    description: str
    met: bool
    evidence: str
    category: str | None = None
    weight: int | None = None
    semantic: bool | None = None
    failure_cap: float | None = None


def load_apex_gold(task_id: str | None = None) -> dict[str, Any]:
    payload = json.loads(GOLD_PATH.read_text(encoding="utf-8"))
    return payload[task_id] if task_id else payload


def _normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if not isinstance(value, str):
        return None
    text = value.strip()
    percent = text.endswith("%")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("() ")
    suffix_match = re.search(r"([kmb])\s*$", text, flags=re.I)
    multiplier = {"k": 1_000.0, "m": 1_000_000.0, "b": 1_000_000_000.0}.get(
        suffix_match.group(1).lower() if suffix_match else "",
        1.0,
    )
    if suffix_match:
        text = text[: suffix_match.start()]
    text = text.replace("$", "").replace(",", "").replace("×", "").replace("x", "")
    text = text.strip("% ")
    try:
        result = float(text)
    except ValueError:
        return None
    if negative:
        result = -result
    result *= multiplier
    return result / 100 if percent else result


def _answer_mapping(answer: Any) -> dict[str, Any]:
    if isinstance(answer, dict):
        return answer
    text = str(answer or "").strip()
    candidates = [text]
    candidates.extend(re.findall(r"```(?:json)?\s*(\{.*?\})\s*```", text, flags=re.I | re.S))
    brace = re.search(r"\{.*\}", text, flags=re.S)
    if brace:
        candidates.append(brace.group(0))
    for candidate in reversed(candidates):
        try:
            parsed = json.loads(candidate)
            if isinstance(parsed, dict):
                return parsed
        except (json.JSONDecodeError, TypeError):
            pass

    mapping: dict[str, Any] = {}
    for line in text.splitlines():
        match = re.match(r"\s*[-*]?\s*([A-Za-z][A-Za-z0-9 _/-]{1,60})\s*[:=]\s*(.+?)\s*$", line)
        if match:
            mapping[re.sub(r"\s+", "_", match.group(1).strip().lower())] = match.group(2).strip()
    return mapping


def _get(mapping: dict[str, Any], key: str) -> Any:
    if key in mapping:
        return mapping[key]
    wanted = _normalize(key)
    for candidate, value in mapping.items():
        if _normalize(candidate) == wanted:
            return value
    return None


def _close(actual: Any, expected: float, *, abs_tol: float = .02, rel_tol: float = 1e-6) -> bool:
    value = _number(actual)
    if value is None:
        return False
    return abs(value - expected) <= max(abs_tol, abs(expected) * rel_tol)


def _result(criteria: list[Criterion]) -> dict[str, Any]:
    met_count = sum(criterion.met for criterion in criteria)
    reward = met_count / len(criteria) if criteria else 0.0
    return {
        "reward": round(reward, 6),
        "strict_pass": bool(criteria) and met_count == len(criteria),
        "criteria_met": met_count,
        "criteria_total": len(criteria),
        "criteria": [
            {
                "id": criterion.id,
                "description": criterion.description,
                "value": int(criterion.met),
                "evidence": criterion.evidence,
                **(
                    {
                        "category": criterion.category,
                        "weight": criterion.weight,
                        "semantic": criterion.semantic,
                        **(
                            {"failure_cap": criterion.failure_cap}
                            if criterion.failure_cap is not None else {}
                        ),
                    }
                    if criterion.category is not None else {}
                ),
            }
            for criterion in criteria
        ],
    }


def _task_011_presentation_specs(gold: dict[str, Any]) -> list[tuple[str, str]]:
    """Return the stable, artifact-native Task 011 rubric denominator."""

    specs: list[tuple[str, str]] = [
        ("artifact__slide_count", "The operating review contains exactly five slides"),
        ("artifact__slide_1_title", "Slide 1 has the Q2 service scorecard and June bridge title"),
        ("artifact__slide_2_title", "Slide 2 has the callback and work-mix review title"),
        ("artifact__slide_3_title", "Slide 3 has the management decision title"),
        ("artifact__slide_4_title", "Slide 4 has the dispatcher timing decision title"),
        ("artifact__slide_5_title", "Slide 5 has the dispatcher signal decision title"),
        ("artifact__slide_1_table", "Slide 1 contains a real PowerPoint scorecard table"),
        ("artifact__slide_2_table", "Slide 2 contains a real PowerPoint technician table"),
        ("artifact__slide_3_action_table", "Slide 3 contains a real PowerPoint action-window table"),
        ("artifact__slide_3_alternative_table", "Slide 3 contains a separate real PowerPoint alternatives table"),
        ("artifact__slide_3_capacity_table", "Slide 3 contains a real PowerPoint branch-capacity table"),
        ("artifact__slide_3_portfolio_table", "Slide 3 contains a real PowerPoint service-flex portfolio table"),
        ("artifact__slide_4_timing_table", "Slide 4 contains a real PowerPoint timing-outcome table"),
        ("artifact__slide_4_timing_chart_native", "Slide 4 contains a native editable policy expected-value chart"),
        ("artifact__slide_4_timing_chart_categories", "The timing chart identifies both lock-now and wait policies"),
        ("artifact__slide_5_signal_table", "Slide 5 contains a real PowerPoint signal-validation table"),
        ("artifact__slide_5_signal_chart_native", "Slide 5 contains a native editable policy-value chart"),
        ("artifact__slide_5_signal_chart_categories", "The signal chart identifies lock, unconditional wait, and contingent policies"),
        ("artifact__monthly_chart_native", "Slide 1 contains a native editable monthly KPI chart"),
        ("artifact__monthly_chart_categories", "The monthly chart uses April, May, and June in order"),
        ("artifact__callback_chart_native", "Slide 2 contains a native editable actual-versus-expected callback chart"),
        ("artifact__callback_chart_categories", "The callback chart uses the three watched technicians in follow-up order"),
    ]
    specs.extend(
        (
            f"finance__company__{field}",
            f"The Q2 company scorecard reports the correct {field.replace('_', ' ')}",
        )
        for field in (
            "completed_work_orders", "callbacks", "first_time_fix_rate",
            "billable_hours", "revenue_proxy",
        )
    )
    for month in gold["monthly"]:
        month_id = str(month["month"]).replace("-", "_")
        for field in ("completed_work_orders", "callbacks", "revenue_proxy"):
            specs.append((
                f"finance__monthly__{month_id}__{field}",
                f"The monthly scorecard reports {month['month']} {field.replace('_', ' ')}",
            ))
        for field in ("completed_work_orders", "callbacks"):
            specs.append((
                f"finance__monthly_chart__{month_id}__{field}",
                f"The native monthly chart plots {month['month']} {field.replace('_', ' ')}",
            ))
    for branch in gold["branches"]:
        branch_id = _normalize(branch["branch"]).replace(" ", "_")
        for field in ("completed_work_orders", "callbacks", "first_time_fix_rate", "revenue_proxy"):
            specs.append((
                f"finance__branch__{branch_id}__{field}",
                f"The branch scorecard reports {branch['branch']} {field.replace('_', ' ')}",
            ))
    for field in (
        "completed_work_orders", "billable_hours", "revenue_proxy",
    ):
        specs.append((
            f"finance__source_control__{field}",
            f"The retained workpaper reports the correct {field.replace('_', ' ')} and ties it to accounting",
        ))
    for field in (
        "review_completed_work_orders", "review_callbacks",
        "review_billable_hours", "review_revenue_proxy",
    ):
        specs.append((
            f"finance__adjusted_control__{field}",
            f"The retained-to-adjusted control reports the correct {field.replace('_', ' ')}",
        ))
    for field in (
        "completed_work_order_change", "callback_change", "first_time_fix_rate_change",
        "revenue_proxy_change", "average_revenue_per_work_order_change",
    ):
        specs.append((
            f"finance__june_vs_may__{field}",
            f"The June-versus-May view reports the correct {field.replace('_', ' ')}",
        ))
    for row in gold["june_value_bridge"][:3]:
        work_type = _normalize(row["work_type"]).replace(" ", "_")
        specs.append((
            f"finance__june_bridge__{work_type}",
            f"The June bridge reports the correct total change for {row['work_type']}",
        ))
    specs.append(("finance__june_bridge__company_tie", "The work-type bridge ties to the company June-versus-May operating-value change"))
    for field in ("work_type", "completed_work_orders", "callbacks", "callback_rate"):
        specs.append((
            f"finance__callback_focus__{field}",
            f"The callback-focus view reports the correct {field.replace('_', ' ')}",
        ))
    for row in gold["mix_adjusted_callback_summary"]:
        technician = str(row["technician_id"]).casefold()
        for field in ("completed_work_orders", "actual_callbacks", "expected_callbacks", "callback_excess"):
            specs.append((
                f"finance__technician__{technician}__{field}",
                f"The technician table reports {row['technician_id']} {field.replace('_', ' ')}",
            ))
        for field in ("actual_callbacks", "expected_callbacks"):
            specs.append((
                f"finance__callback_chart__{technician}__{field}",
                f"The native callback chart plots {row['technician_id']} {field.replace('_', ' ')}",
            ))
    for row in gold["execution_priority_decision"]["alternative_economics"]:
        action = _normalize(row["action"])
        if "sensor" in action:
            action_id = "sensor_recovery"
        elif "control board" in action:
            action_id = "control_board_recovery"
        else:
            action_id = "retention_package"
        for field in (
            "source_amount", "timely_probability", "queued_probability",
            "action_cost", "expected_net_now", "expected_net_if_queued",
            "incremental_value",
        ):
            specs.append((
                f"finance__alternative__{action_id}__{field}",
                f"The alternatives table reports {action_id.replace('_', ' ')} {field.replace('_', ' ')}",
            ))
    specs.append(("finance__fallback__sensor_threshold", "The sensor fallback probability threshold is correct"))

    capacity = gold["branch_capacity_decision"]
    for row in capacity["branches"]:
        branch = _normalize(row["branch"]).replace(" ", "_")
        for field in (
            "q2_operating_value_per_billable_hour",
            "planning_conversion",
            "realization_probability",
            "expected_contribution_per_hour",
            "service_floor_hours",
        ):
            specs.append((
                f"capacity__{branch}__{field}",
                f"The capacity table reports {row['branch']} {field.replace('_', ' ')}",
            ))
    specs.extend((
        ("capacity__total_flex_hours", "The full 120-hour flex block is allocated"),
        ("capacity__discretionary_flex_hours", "The 16 discretionary hours are identified"),
        ("capacity__planning_boundary", "The capacity scenario stays within planning and scheduling authority"),
    ))

    portfolio = gold["service_flex_portfolio_decision"]
    for row in portfolio["actions"]:
        action_id = _normalize(row["action"]).replace(" ", "_")
        specs.extend((
            (f"portfolio__action__{action_id}__branch_and_hours", f"{row['action']} uses the correct branch and eight-hour block"),
            (f"portfolio__action__{action_id}__service_contribution", f"{row['action']} reports the expected service contribution"),
            (f"portfolio__action__{action_id}__protected_contribution", f"{row['action']} reports the expected protected contribution"),
            (f"portfolio__action__{action_id}__outside_cost", f"{row['action']} reports the outside cost"),
            (f"portfolio__action__{action_id}__expected_net", f"{row['action']} reports the expected net consequence"),
        ))
    specs.extend((
        ("portfolio__dispatcher_constraint", "The one dispatcher-handoff constraint is applied"),
        ("portfolio__selected_action__bend_commissioning", "Bend controls commissioning is selected"),
        ("portfolio__selected_action__portland_retention", "Portland PM retention closeout is selected"),
        ("portfolio__selected_expected_net", "The selected portfolio expected net consequence is correct"),
        ("portfolio__next_best_actions", "The next-best feasible portfolio is identified"),
        ("portfolio__next_best_expected_net", "The next-best feasible portfolio expected net consequence is correct"),
        ("portfolio__selected_advantage", "The selected portfolio advantage is quantified"),
        ("portfolio__final_portland_hours", "The final Portland flex hours are correct"),
        ("portfolio__final_bend_hours", "The final Bend flex hours are correct"),
        ("portfolio__fallback_trigger", "The Bend site-access fallback cutoff is correct"),
        ("portfolio__fallback_actions", "The fallback portfolio is correct"),
        ("portfolio__planning_boundary", "The portfolio remains an internal planning recommendation within authority"),
    ))

    specs.extend((
        ("timing__decision_deadline", "The dispatcher policy is decided at the July 9 9:00 a.m. deadline"),
        ("timing__bend_access_probability", "The Bend access probability is correct"),
        ("timing__late_hospital_probability", "The late hospital placement probability is correct"),
        ("timing__lock_now_expected_net", "The lock-now policy expected net is correct"),
        ("timing__wait_access_branch_expected_net", "The wait-and-access branch expected net is correct"),
        ("timing__wait_no_access_branch_expected_net", "The wait-and-no-access branch expected net is correct"),
        ("timing__wait_policy_expected_net", "The probability-weighted wait policy expected net is correct"),
        ("timing__lock_now_advantage", "The lock-now expected-value advantage is correct"),
        ("timing__access_probability_breakeven", "The Bend access probability crossover is correct"),
        ("timing__chart_lock_now_expected_net", "The native timing chart plots the lock-now expected net"),
        ("timing__chart_wait_policy_expected_net", "The native timing chart plots the wait-policy expected net"),
        ("timing__recommend_lock_hospital_pm", "The executable policy locks the hospital SLA and keeps PM closeout"),
        ("timing__static_ranking_flip_reason", "The recommendation explains why probability-weighted execution reverses the static ranking"),
        ("timing__no_substitute_or_double_count", "The failed late-placement branch uses PM only without substitution or double counting"),
        ("timing__planning_boundary", "The timing policy remains within Service Operations planning authority"),
    ))

    specs.extend((
        ("signal__count__green_access", "The signal table reports green with access confirmed"),
        ("signal__count__green_no_access", "The signal table reports green without access confirmation"),
        ("signal__count__red_access", "The signal table reports red with access confirmed"),
        ("signal__count__red_no_access", "The signal table reports red without access confirmation"),
        ("signal__validation_total", "The signal validation population totals 400 decisions"),
        ("signal__green_probability", "The probability of a green signal is correct"),
        ("signal__access_given_green", "The posterior access probability after green is correct"),
        ("signal__access_given_red", "The posterior access probability after red is correct"),
        ("signal__wait_value_given_green", "The conditional value of waiting after green is correct"),
        ("signal__wait_value_given_red", "The conditional value of waiting after red is correct"),
        ("signal__policy_value__lock", "The unconditional lock policy value is correct"),
        ("signal__policy_value__wait", "The unconditional wait policy value is correct"),
        ("signal__policy_value__contingent", "The signal-contingent policy value is correct"),
        ("signal__chart_value__lock", "The native signal chart plots unconditional lock correctly"),
        ("signal__chart_value__wait", "The native signal chart plots unconditional wait correctly"),
        ("signal__chart_value__contingent", "The native signal chart plots the contingent policy correctly"),
        ("signal__advantage_vs_lock", "The contingent policy advantage over lock is correct"),
        ("signal__advantage_vs_wait", "The contingent policy advantage over unconditional wait is correct"),
        ("signal__risk_tolerance", "The PM-only risk tolerance is correct"),
        ("signal__risk__unconditional_wait", "The unconditional-wait PM-only risk is correct"),
        ("signal__risk__contingent", "The contingent-policy PM-only risk is correct"),
        ("signal__decision__green", "The recommendation waits only after a green signal"),
        ("signal__decision__red", "The recommendation locks after a red signal"),
        ("signal__decision__missing", "The recommendation locks when the signal is missing"),
        ("signal__risk_conclusion", "The recommendation distinguishes the policy that meets the PM-only tolerance"),
        ("signal__maximum_paid_cost", "The maximum justified paid signal cost is correct"),
        ("signal__evidence_boundary", "Validation evidence is kept distinct from customer confirmation"),
        ("signal__planning_boundary", "The signal policy remains within Service Operations planning authority"),
    ))

    specs.extend((
        ("holdout__count__green_access", "The holdout table reports green with access confirmed"),
        ("holdout__count__green_no_access", "The holdout table reports green without access confirmation"),
        ("holdout__count__red_access", "The holdout table reports red with access confirmed"),
        ("holdout__count__red_no_access", "The holdout table reports red without access confirmation"),
        ("holdout__access_given_green", "The holdout access probability after green is correct"),
        ("holdout__access_given_red", "The holdout access probability after red is correct"),
        ("holdout__wait_value_given_green", "The holdout conditional value of waiting after green is correct"),
        ("holdout__wait_value_given_red", "The holdout conditional value of waiting after red is correct"),
        ("holdout__policy_value__contingent", "The holdout signal-contingent policy value is correct"),
        ("holdout__advantage_vs_lock", "The holdout signal-policy value difference versus lock is correct"),
        ("holdout__risk__contingent", "The holdout signal-contingent PM-only risk is correct"),
        ("holdout__sample_control", "The development sample is reference evidence and the holdout controls go-live"),
        ("holdout__decision__no_deploy", "The holdout-supported recommendation does not deploy the signal"),
        ("holdout__decision__green", "The go-live policy locks after a green signal"),
        ("holdout__shadow_mode", "The signal is limited to shadow-mode review or retraining"),
        ("holdout__maximum_paid_cost", "The supported paid signal cost is zero"),
    ))

    specs.extend((
        ("oot__population_count", "The out-of-time review reconciles all 240 observations"),
        ("oot__unique_decision_ids", "The out-of-time decision IDs are unique"),
        ("oot__access_outcomes_complete", "The out-of-time access outcomes are complete"),
        ("oot__timely_signal_count", "The out-of-time review identifies the timely signals"),
        ("oot__late_signal_count", "The out-of-time review identifies the late signals"),
        ("oot__missing_signal_count", "The out-of-time review identifies the missing signals"),
        ("oot__timely_green_count", "The out-of-time review identifies timely green signals"),
        ("oot__timely_green_access_count", "The out-of-time review identifies timely green signals with access"),
        ("oot__timely_green_late_hospital_count", "The out-of-time review identifies timely green failures recovered by the late hospital handoff"),
        ("oot__timely_green_pm_only_count", "The out-of-time review identifies timely green PM-only outcomes"),
        ("oot__access_given_timely_green", "The out-of-time timely-green access probability is correct"),
        ("oot__policy_average_value", "The out-of-time contingent-policy realized average value is correct"),
        ("oot__policy_difference_vs_lock", "The out-of-time contingent-policy difference versus locking is correct"),
        ("oot__pm_only_risk", "The out-of-time contingent-policy PM-only risk is correct"),
        ("oot__risk_tolerance", "The out-of-time decision applies the 1.0% PM-only tolerance"),
        ("oot__late_or_missing_lock", "Late and missing signals use the ordinary lock policy"),
        ("oot__no_backfill", "Late and missing signals are not backfilled from eventual outcomes"),
        ("oot__decision_no_deploy", "The out-of-time evidence does not support live deployment"),
        ("oot__shadow_mode", "The signal remains in shadow mode for review or retraining"),
        ("oot__maximum_paid_cost", "The supported paid signal cost remains zero"),
        ("oot__sample_role", "The out-of-time sample is identified as the final production-readiness check"),
        ("oot__planning_boundary", "The out-of-time evidence remains within internal planning authority"),
        ("oot__workpaper_population", "The workbook preserves and controls the row-level out-of-time population"),
        ("oot__workpaper_policy_value", "The workbook calculates the out-of-time policy value with formulas"),
        ("oot__workpaper_risk", "The workbook calculates the out-of-time PM-only risk with formulas"),
        ("oot__workpaper_decision", "The workbook decision agrees with the out-of-time evidence"),
    ))

    specs.extend((
        ("workpaper__artifact_readable", "The formula-driven decision-support workbook is readable"),
        ("workpaper__sheet__inputs", "The workbook has a clear Inputs sheet"),
        ("workpaper__formula_count", "The workbook contains a substantial formula-driven calculation trail"),
        ("workpaper__cross_sheet_formula_count", "The workbook uses auditable cross-sheet references"),
        ("workpaper__no_formula_errors", "The workbook has no visible formula errors"),
        ("workpaper__source_labels", "The workbook names the timing note and distinguishes development from holdout evidence"),
        ("workpaper__input__holdout_counts", "The workbook contains all four holdout joint counts"),
        ("workpaper__input__risk_tolerance", "The workbook contains the PM-only risk tolerance"),
        ("workpaper__holdout__access_given_green", "The holdout access probability after green is formula-driven and correct"),
        ("workpaper__holdout__access_given_red", "The holdout access probability after red is formula-driven and correct"),
        ("workpaper__holdout__policy_value", "The holdout contingent-policy value is formula-driven and correct"),
        ("workpaper__holdout__risk", "The holdout contingent-policy downside risk is formula-driven and correct"),
        ("workpaper__decision__no_deploy", "The workbook recommendation does not deploy the signal"),
    ))

    specs.extend((
        ("management__proxy_boundary", "The operating-value proxy is not presented as GAAP revenue or recoverable loss"),
        ("management__source_adjustment_boundary", "The retained source tie is kept separate from approved review adjustments"),
        ("management__driver__preventive_maintenance", "Preventive maintenance is identified as the largest negative June operating-value driver"),
        ("management__driver__emergency_response", "Emergency response is identified as a positive June operating-value driver"),
    ))
    for action in gold["management_use_decision"]["technician_actions"]:
        technician = str(action["technician_id"]).casefold()
        specs.extend((
            (f"management__technician__{technician}__priority_work_type", f"{action['technician_id']} has the correct priority work type"),
            (f"management__technician__{technician}__evidence_work_order_1", f"{action['technician_id']} first reviewed work order is identified"),
            (f"management__technician__{technician}__evidence_work_order_2", f"{action['technician_id']} second reviewed work order is identified"),
            (f"management__technician__{technician}__evidence_assessment", f"{action['technician_id']} evidence assessment is source-supported"),
            (f"management__technician__{technician}__kpi_boundary", f"{action['technician_id']} KPI treatment stays within authority"),
            (f"management__technician__{technician}__operational_follow_up", f"{action['technician_id']} operational follow-up is source-supported"),
            (f"management__technician__{technician}__unsupported_actions", f"{action['technician_id']} unsupported performance and pricing actions are rejected"),
        ))
    specs.extend((
        ("management__review_release_status", "The operating review is correctly limited to conditional internal use"),
        ("management__remaining_controls", "The remaining evidence, reason-code, and signoff controls are identified"),
    ))
    action_window = gold["action_window_decision"]
    for slot, label in (
        ("supervisor_root_cause_slot", "supervisor"),
        ("supplier_recovery_slot", "supplier"),
        ("evidence_retrieval_slot", "evidence"),
    ):
        specs.extend((
            (f"management__action_window__{label}__case", f"The {label} slot is allocated to the correct case"),
            (f"management__action_window__{label}__deadline", f"The {label} slot has the correct deadline"),
        ))
    for index, row in enumerate(action_window["deferred_queue"], start=1):
        specs.append((
            f"management__deferred_queue__{index:02d}",
            f"Deferred action {index} has the correct case and deadline",
        ))
    specs.extend((
        ("management__coaching_boundary", "The unsupported immediate coaching action is rejected"),
        ("management__pricing_boundary", "The unsupported small-retrofit pricing action is rejected"),
        ("management__selected_use", "The four-hour block is allocated to the correct primary use"),
        ("management__selection_basis", "The selected use is supported by incremental value and timing"),
        ("management__sequence__second", "The control-board recovery is sequenced second"),
        ("management__sequence__third", "The retention package is sequenced third"),
        ("management__fallback_use", "The control-board recovery is the correct fallback"),
        ("management__fallback_trigger", "The fallback trigger is stated on a supportable basis"),
        ("management__retention__amount", "The conditional retention concession amount is correct"),
        ("management__retention__offer_booking_boundary", "The retention package is not presented as an authorized offer or booking"),
        ("management__retention__service_gm_approval", "The Service GM approval boundary is identified"),
        ("management__retention__controller_approval", "The Controller approval boundary is identified"),
        ("management__kpi_release_boundary", "The adjusted KPI and internal release boundaries remain clear"),
    ))
    for slide in range(1, 6):
        specs.extend((
            (f"readability__slide_{slide}__bounds", f"Slide {slide} content stays within the slide bounds"),
            (f"readability__slide_{slide}__overlap", f"Slide {slide} tables and charts do not materially overlap"),
            (f"readability__slide_{slide}__font_floor", f"Slide {slide} uses no explicitly undersized text"),
        ))
        specs.append((f"source__slide_{slide}", f"Slide {slide} names its controlling source families"))
    specs = [
        (
            criterion_id,
            criterion_id.replace("__", " ").replace("_", " "),
        )
        if criterion_id.startswith(("holdout__", "oot__", "workpaper__"))
        else (criterion_id, description)
        for criterion_id, description in specs
    ]
    ids = [criterion_id for criterion_id, _ in specs]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate Task 011 presentation criterion ids")
    return specs


def _legacy_file_atomic_specs(task_id: str) -> list[tuple[str, str]]:
    """Stable atomic denominator for missing or unreadable legacy artifacts."""

    gold = load_apex_gold(task_id)
    specs: list[tuple[str, str]] = []
    if task_id == "task_004":
        input_fields = (
            "current_contract", "cost_to_date", "pm_etc",
            "finance_adjustment", "billings", "prior_margin_percent",
        )
        result_fields = (
            "close_etc", "eac", "percent_complete", "earned_revenue",
            "contract_asset", "contract_liability", "margin",
            "margin_percent", "margin_movement_bps",
        )
        total_fields = (
            "current_contract", "cost_to_date", "pm_etc",
            "finance_adjustment", "close_etc", "eac", "percent_complete",
            "earned_revenue", "billings", "contract_asset",
            "contract_liability", "margin", "margin_percent",
            "prior_margin_percent", "margin_movement_bps",
        )
        for expected in gold["projects"]:
            slug = expected["project_id"].casefold()
            specs.extend(
                (f"input__{slug}__{field}", f"{expected['project_id']} {field} input is correct")
                for field in input_fields
            )
            specs.append((
                f"formula_lineage__{slug}",
                f"{expected['project_id']} close calculations respond to source inputs",
            ))
            specs.extend(
                (f"result__{slug}__{field}", f"{expected['project_id']} {field} result is correct")
                for field in result_fields
            )
            specs.extend((
                (f"review_status__{slug}", f"{expected['project_id']} review status is supportable"),
                (f"commercial_treatment__{slug}", f"{expected['project_id']} commercial treatment is supportable"),
                (f"source_support__{slug}", f"{expected['project_id']} controlling support is identifiable"),
                (f"next_action__{slug}", f"{expected['project_id']} next action and owner are sufficient"),
            ))
        specs.append(("total_formula_lineage", "Portfolio totals respond to the four project rows"))
        specs.extend(
            (f"total_result__{field}", f"Portfolio {field} total is correct")
            for field in total_fields
        )
        specs.extend((
            ("portfolio__review_queue", "The portfolio review queue is correct"),
            ("portfolio__release_decision", "The portfolio release recommendation is supportable"),
        ))
    elif task_id == "task_008":
        specs.extend((f"preservation__{_normalize(sheet).replace(' ', '_')}", f"Preserve {sheet}") for sheet in ("Base Case Import", "Scenario Assumptions"))
        specs.extend((f"assumption__{label}", f"Scenario assumption {label}") for label in ("eastbank_collection", "cascade_collection", "redmond_collection", "accelerated_po", "insurance_deposit"))
        for expected in gold["weeks"]:
            specs.extend((f"result__{expected['week_ending']}__{key}", f"{expected['week_ending']} {key}") for key in ("cash_before_financing", "revolver_draw", "ending_cash", "ending_revolver"))
        specs.extend((f"formula_lineage__row_{row}", f"Formula lineage row {row}") for row in range(6, 19))
        specs.extend((f"decision__{label}", f"Treasury review {label}") for label in (
            "base_minimum_cash", "downside_minimum_cash_before_financing",
            "cash_deterioration_vs_base", "peak_revolver", "remaining_commitment",
            "minimum_cash_week", "first_draw_week", "full_repayment_week",
            "facility_capacity_sufficient", "treasury_action",
        ))
        specs.extend((f"decision_formula__{label}", f"Treasury review formula lineage {label}") for label in (
            "base_minimum_cash", "downside_minimum_cash_before_financing",
            "cash_deterioration_vs_base", "minimum_cash_week", "first_draw_week",
            "peak_revolver", "remaining_commitment", "full_repayment_week",
            "facility_capacity_sufficient",
        ))
        specs.append(("chart", "Required liquidity chart"))
    elif task_id == "task_011":
        specs.extend(_task_011_presentation_specs(gold))
    elif task_id == "task_012":
        specs.extend((f"preservation__{_normalize(sheet).replace(' ', '_')}", f"Preserve {sheet}") for sheet in ("read me - DC marks", "Source Pull"))
        specs.extend((f"source_pull__{label}", f"Source Pull {label}") for label in ("posted_revenue", "proposed_wip_adjustment", "pro_forma_revenue", "plan_revenue", "gross_profit", "plan_gross_profit", "operating_income", "plan_operating_income"))
        specs.extend((f"bridge_value__{label}", f"Bridge value {label}") for label in ("pro_forma_revenue", "plan_revenue", "revenue_variance", "gross_profit", "plan_gross_profit", "gross_profit_variance", "operating_income", "plan_operating_income", "operating_income_variance", "revenue_variance_percent", "gross_margin_rate_variance", "operating_income_variance_percent"))
        specs.extend((f"interpretation__{label}", f"Bridge interpretation {label}") for label in ("revenue", "gross_profit", "operating_income"))
        specs.extend((f"bridge_formula__{column.casefold()}{row}", f"Bridge formula {column}{row}") for row in range(6, 9) for column in "BCDEFG")
        specs.append(("chart", "Required comparison chart"))
    elif task_id == "task_015":
        specs.append(("preservation__slide_count", "Exactly one slide appended"))
        specs.extend((f"preservation__slide_{index:02d}", f"Preserve slide {index}") for index in range(1, 6))
        specs.extend((("structure__title", "Required slide title"), ("structure__table", "Required covenant comparison")))
        specs.extend((
            ("style__visual_system", "Match the existing deck visual system"),
            ("style__title_hierarchy", "Use a clear title hierarchy consistent with the deck"),
            ("style__table_readability", "Present the covenant comparison in a reviewable table or equivalent executive layout"),
        ))
        for metric in ("leverage", "fccr", "tangible_net_worth"):
            specs.extend(
                (f"metric__{metric}__{field}", f"{metric} {field}")
                for field in (
                    "posted", "pro_forma", "threshold", "posted_headroom",
                    "pro_forma_headroom", "status",
                )
            )
        specs.extend((
            ("status__proposed", "Proposed WIP status"),
            ("status__unposted", "Unposted WIP status"),
            ("status__compliant", "Compliance status"),
            ("status__circulation", "Working-deck circulation boundary"),
            ("headroom__binding_covenant", "Binding covenant on like-for-like deterioration capacity"),
            ("source__authority", "Source note identifies the controlling calculation and accounting authorities"),
        ))
        specs.extend(
            (
                f"calculation_support__{field}",
                f"Calculation support {description}",
            )
            for field, description in (
                ("funded_debt", "funded debt"),
                ("posted_adjusted_ebitda", "posted adjusted EBITDA"),
                ("pending_ap_cutoff_expense", "pending June AP cutoff expense"),
                ("proposed_wip_adjustment", "proposed WIP adjustment"),
                ("pro_forma_adjusted_ebitda", "pro-forma adjusted EBITDA"),
                ("cash_taxes", "cash taxes"),
                ("ltm_fixed_asset_additions", "LTM fixed-asset additions"),
                ("direct_equipment_financing", "direct LTM equipment financing"),
                ("unfunded_capex", "unfunded capex"),
                ("cash_interest", "cash interest"),
                ("scheduled_principal", "scheduled principal"),
                ("fixed_charges", "fixed charges"),
                ("posted_fccr_numerator", "posted FCCR numerator"),
                ("pro_forma_fccr_numerator", "pro-forma FCCR numerator"),
            )
        )
        specs.extend(
            (
                f"calculation_support__equation__{field}",
                f"Calculation support equation {description}",
            )
            for field, description in (
                ("ebitda_bridge", "posted EBITDA less pending AP expense plus proposed WIP equals pro-forma EBITDA"),
                ("unfunded_capex", "fixed-asset additions less direct financing equals unfunded capex"),
                ("fixed_charges", "cash interest plus scheduled principal equals fixed charges"),
                ("posted_fccr_numerator", "posted EBITDA less taxes and unfunded capex equals the posted FCCR numerator"),
                ("pro_forma_fccr_numerator", "pro-forma EBITDA less taxes and unfunded capex equals the pro-forma FCCR numerator"),
                ("tnw_bridge", "posted TNW less pending AP expense plus proposed WIP equals pro-forma TNW"),
            )
        )
        specs.append(("structure__wip_reversal_sensitivity", "Required proposed-WIP reversal sensitivity table"))
        for reversal in (25, 50, 100):
            specs.extend(
                (
                    f"sensitivity__reversal_{reversal}__{field}",
                    f"{reversal}% WIP-reversal sensitivity {field.replace('_', ' ')}",
                )
                for field in (
                    "retained_wip_adjustment", "adjusted_ebitda",
                    "tangible_net_worth", "leverage", "fccr", "status",
                )
            )
    elif task_id == "task_023":
        specs.extend((f"structure__sheet_{_normalize(name).replace(' ', '_')}", f"Required sheet {name}") for name in ("Executed Terms", "Q2 Headroom Working", "Borrowing Base"))
        specs.extend((("structure__no_extra_sheets", "No extra sheets"), ("preservation__executed_terms", "Preserve executed terms"), ("preservation__q2_headroom", "Preserve headroom inputs")))
        specs.extend((f"covenant_formula__{column.casefold()}{row}", f"Covenant formula {column}{row}") for row in range(6, 9) for column in "CDEFG")
        specs.extend((f"covenant_value__{column.casefold()}{row}", f"Covenant value {column}{row}") for row in range(6, 9) for column in "CDEFG")
        labels = ("Invoice Count", "Gross Open AR", "Pre-Concentration Eligible AR", "Customer Cap", "Post-Concentration Eligible AR", "Advance Rate", "Borrowing Base")
        for label in labels:
            slug = _normalize(label).replace(" ", "_")
            specs.extend(((f"borrowing_base_label__{slug}", f"Borrowing-base label {label}"), (f"borrowing_base_value__{slug}", f"Borrowing-base value {label}")))
        specs.extend((f"borrowing_base_formula__b{row}", f"Borrowing-base formula B{row}") for row in range(5, 9))
        specs.append(("support__heading", "Covenant-support heading"))
        specs.extend((f"support__{_normalize(label).replace(' ', '_')}", f"Covenant support {label}") for label in ("Funded Debt", "Posted EBITDA", "Pro Forma EBITDA", "Posted FCCR", "Pro Forma FCCR", "Posted Tangible Net Worth", "Pro Forma Tangible Net Worth"))
        wip_labels = (
            "Full Project Count",
            "Controller Review Project Count",
            "Other Project Count",
            "Gross Underbillings",
            "Gross Overbillings",
            "Proposed Net WIP",
            "Controller Review Net WIP",
            "Other Project Net WIP",
            "Reconciliation Difference",
        )
        specs.extend(
            (
                f"wip_control_value__{_normalize(label).replace(' ', '_')}",
                f"WIP population control {label}",
            )
            for label in wip_labels
        )
        specs.extend(
            (
                f"wip_control_formula__{slug}",
                f"WIP population control formula {slug}",
            )
            for slug in (
                "other_project_count",
                "proposed_net_wip",
                "other_project_net_wip",
                "reconciliation_difference",
            )
        )
        specs.extend(
            (
                f"lender_review__{slug}",
                f"Lender review control {description}",
            )
            for slug, description in (
                ("posted_covenant_basis", "posted covenant basis"),
                ("proposed_wip_presentation", "proposed-WIP presentation"),
                ("draft_delivery_status", "draft delivery status"),
                ("delivery_requirement", "delivery requirement"),
                ("wip_posting_requirement", "WIP posting requirement"),
            )
        )
    elif task_id == "task_024":
        specs.extend((("structure__funding_screen", "Funding Screen sheet"), ("structure__sources_controls", "Sources & Controls sheet"), ("structure__no_extra_sheets", "No extra sheets"), ("structure__row_count", "Five request rows")))
        specs.extend((("funding_views__table_structure", "FundingViews table structure"), ("funding_views__metric_order", "FundingViews metric order")))
        headers = ("Request ID", "Asset", "Documented Decision", "Request Amount", "Approved Amount", "Funding", "Annual Savings", "Simple Payback", "Equipment-Line Proceeds", "Cash Required", "Finance Classification")
        specs.extend((f"structure__header_{_normalize(header).replace(' ', '_')}", f"Header {header}") for header in headers)
        for request in gold["requests"]:
            rid = request["request_id"].casefold()
            specs.append((f"request__{rid}__present", f"Request {rid} present"))
            specs.extend((f"request__{rid}__{field}", f"Request {rid} {field}") for field in ("request_amount", "approved_amount", "annual_savings", "simple_payback", "equipment_line_proceeds", "cash_required", "asset", "decision", "funding", "classification"))
            specs.extend((f"request_formula__{rid}__{field}", f"Request {rid} formula {field}") for field in ("simple_payback", "equipment_line_proceeds", "cash_required"))
        summary_labels = ("Approved Spend", "Approved Annual Savings", "Equipment-Line Proceeds", "Cash Required", "Facility Remaining", "AOP Remaining After Approved Requests", "Downside Maximum Revolver", "Revolver Capacity After Downside", "Pro Forma Leverage")
        for label in summary_labels:
            slug = _normalize(label).replace(" ", "_")
            specs.extend(((f"summary_label__{slug}", f"Summary label {label}"), (f"summary_value__{slug}", f"Summary value {label}"), (f"summary_formula__{slug}", f"Summary formula {label}")))
        funding_view_labels = (
            "Approved Spend", "Equipment-Line Proceeds", "Cash Required",
            "Pro Forma Funded Debt", "Pro Forma Cash", "Pro Forma Net Debt",
            "Pro Forma Leverage",
        )
        for label in funding_view_labels:
            slug = _normalize(label).replace(" ", "_")
            for column in ("committed", "indicative", "difference"):
                specs.extend((
                    (f"funding_views__{slug}__{column}", f"FundingViews {label} {column}"),
                    (f"funding_views_formula__{slug}__{column}", f"FundingViews formula {label} {column}"),
                ))
        for token in ("capex asks", "fy26 op plan", "fleet", "equipment line proposal", "downside assumptions", "usbank amdt2"):
            specs.append((f"source__{_normalize(token).replace(' ', '_')}", f"Source {token}"))
        specs.extend((("source__accounting_lineage", "Accounting lineage"), ("controls__formula_count", "Formula-driven controls")))
    elif task_id == "task_025":
        specs.extend((("identity__title", "Required title"), ("identity__status", "Internal-working status"), ("table__real_word_table", "Real Word table")))
        specs.extend((f"table__row_{_normalize(row).replace(' ', '_')}", f"Decision row {row}") for row in ("proposed june wip", "downside liquidity", "borrowing base", "covenant", "midyear capex"))
        specs.extend((f"wip__{field}", f"WIP {field}") for field in ("amount", "proposed", "unposted"))
        specs.extend((f"liquidity__{field}", f"Liquidity {field}") for field in ("minimum_cash", "first_draw_week", "maximum_revolver"))
        specs.extend((("borrowing_base__amount", "Borrowing-base amount"), ("covenants__posted_leverage", "Posted leverage"), ("covenants__pro_forma_leverage", "Pro-forma leverage"), ("covenants__compliance", "Covenant compliance"), ("capex__cash_required", "Capex cash required"), ("capex__classification_labels", "Capex classification labels")))
        capex = gold["capex"]
        for category, ids in (("release_now", capex["release_now"]), ("conditional", capex["conditional"]), ("hold_or_defer", capex["hold_or_defer"])):
            specs.extend((f"capex__{request_id.casefold()}__{category}", f"{request_id} {category}") for request_id in ids)
        specs.extend((f"source__{group}", f"Source {group}") for group in ("accounting", "wip", "liquidity", "borrowing_base", "covenants", "capex"))
        specs.append(("length", "Two-page limit"))
    else:
        raise KeyError(task_id)
    ids = [criterion_id for criterion_id, _ in specs]
    if len(ids) != len(set(ids)):
        raise ValueError(f"duplicate legacy atomic criterion ids for {task_id}")
    return specs


def _file_failure(task_id: str, evidence: str) -> dict[str, Any]:
    return _result([
        Criterion(criterion_id, description, False, evidence)
        for criterion_id, description in _legacy_file_atomic_specs(task_id)
    ])


def _legacy_file_semantic_ids(task_id: str) -> set[str]:
    ids = {criterion_id for criterion_id, _ in _legacy_file_atomic_specs(task_id)}
    if task_id == "task_004":
        # Every Task004 criterion is hybrid or semantic. Objective values and
        # formula behavior retain deterministic hard gates, while their
        # association with an edited business label is never decided by a
        # finite alias list. Canonical/recognized associations take the stable
        # deterministic fast path; relabelled evidence reaches scoped review.
        return ids
    if task_id == "task_011":
        return set()
    if task_id == "task_015":
        return {
            "structure__title", "structure__table",
            "structure__wip_reversal_sensitivity",
            "status__compliant", "status__proposed", "status__unposted",
            "status__circulation",
            "headroom__binding_covenant",
            "source__authority",
            "metric__leverage__status", "metric__fccr__status",
            "metric__tangible_net_worth__status",
            "calculation_support__equation__ebitda_bridge",
            "calculation_support__equation__unfunded_capex",
            "calculation_support__equation__fixed_charges",
            "calculation_support__equation__posted_fccr_numerator",
            "calculation_support__equation__pro_forma_fccr_numerator",
            "calculation_support__equation__tnw_bridge",
            "sensitivity__reversal_25__status",
            "sensitivity__reversal_50__status",
            "sensitivity__reversal_100__status",
        } | {
            criterion_id
            for criterion_id in ids
            if criterion_id.startswith((
                "metric__",
                "calculation_support__",
                "sensitivity__",
            ))
        }
    if task_id == "task_023":
        return {criterion_id for criterion_id in ids if criterion_id.startswith("borrowing_base_label__")}
    if task_id == "task_024":
        return {
            criterion_id
            for criterion_id in ids
            if (
                re.match(r"request__.+__(asset|decision|funding|classification)$", criterion_id)
                or criterion_id.startswith("summary_label__")
                or criterion_id.startswith("source__")
            )
        }
    if task_id == "task_025":
        fixed = {
            "identity__title", "identity__status", "wip__proposed", "wip__unposted",
            "liquidity__first_draw_week", "covenants__compliance",
            "capex__classification_labels",
        }
        return fixed | {
            criterion_id
            for criterion_id in ids
            if criterion_id.startswith("source__")
            or re.match(r"capex__cx-\d{2}-\d{3}__", criterion_id)
        }
    return set()


def _canonicalize_legacy_file_policy(task_id: str, result: dict[str, Any]) -> dict[str, Any]:
    """Make missing, corrupt, and complete artifact rubrics policy-identical.

    Artifact evidence descriptions necessarily differ between a missing file
    and a parsed file, so importance and semantic routing must never be inferred
    from whichever description happens to be present on that execution path.
    """

    canonical = attach_default_policy(_file_failure(task_id, "canonical policy"))
    semantic_ids = _legacy_file_semantic_ids(task_id)
    policy_by_id: dict[str, dict[str, Any]] = {}
    for row in canonical["criteria"]:
        criterion_id = str(row["id"])
        row["semantic"] = criterion_id in semantic_ids
        # Source naming is genuine semantic provenance, not formula lineage or
        # a central decision, even when its id includes "accounting_lineage".
        if criterion_id in semantic_ids and criterion_id.startswith("source__"):
            row.update({
                "category": "provenance",
                "weight": 10 if task_id == "task_015" else 3,
            })
            row.pop("failure_cap", None)
        elif task_id == "task_004":
            row.pop("failure_cap", None)
            if criterion_id.startswith("source_support__"):
                row.update({"category": "provenance", "weight": 3})
            elif criterion_id.startswith("formula_lineage__") or criterion_id == "total_formula_lineage":
                row.update({"category": "auditability", "weight": 5, "semantic": True})
            elif criterion_id.startswith((
                "review_status__", "commercial_treatment__",
                "next_action__", "portfolio__",
            )):
                row.update({"category": "decision", "weight": 10, "semantic": True})
            else:
                row.update({"category": "core_finance", "weight": 10, "semantic": True})
        elif task_id == "task_008":
            if criterion_id.startswith("preservation__"):
                row.update({"category": "integrity", "weight": 10, "failure_cap": 0.0})
            elif (
                criterion_id.startswith("formula_lineage__")
                or criterion_id.startswith("decision_formula__")
            ):
                row.update({"category": "auditability", "weight": 5, "semantic": False})
                row.pop("failure_cap", None)
            elif criterion_id.startswith("decision__"):
                row.update({"category": "decision", "weight": 10, "semantic": False})
                row.pop("failure_cap", None)
            elif criterion_id == "chart":
                row.update({"category": "structure", "weight": 1, "semantic": False})
                row.pop("failure_cap", None)
            else:
                row.update({"category": "core_finance", "weight": 10, "semantic": False})
                row.pop("failure_cap", None)
        elif task_id == "task_011":
            row.pop("failure_cap", None)
            if criterion_id.startswith("finance__"):
                row.update({
                    "category": "core_finance",
                    "weight": _task_011_criterion_weight(criterion_id),
                    "semantic": False,
                })
            elif criterion_id.startswith("management__"):
                row.update({
                    "category": "decision",
                    "weight": _task_011_criterion_weight(criterion_id),
                    "semantic": False,
                })
            elif criterion_id.startswith(("capacity__", "portfolio__", "timing__", "signal__", "holdout__", "oot__")):
                row.update({
                    "category": "decision",
                    "weight": _task_011_criterion_weight(criterion_id),
                    "semantic": False,
                })
            elif criterion_id.startswith("workpaper__"):
                row.update({
                    "category": "auditability",
                    "weight": _task_011_criterion_weight(criterion_id),
                    "semantic": False,
                })
            elif criterion_id.startswith("artifact__"):
                row.update({"category": "structure", "weight": 1, "semantic": False})
            elif criterion_id.startswith("readability__"):
                row.update({"category": "presentation_quality", "weight": 3, "semantic": False})
            elif criterion_id.startswith("source__"):
                row.update({"category": "provenance", "weight": 3, "semantic": False})
            else:  # pragma: no cover - guarded by the stable specification
                raise ValueError(f"unclassified task_011 criterion: {criterion_id}")
        elif task_id == "task_023":
            # This lender workbook earns proportional credit for each correct
            # calculation and control. Formula lineage supports the analysis,
            # while lender-release decisions and the material finance outputs
            # carry the greatest weight. No isolated miss imposes a scalar cap.
            row.pop("failure_cap", None)
            if criterion_id.startswith("structure__"):
                row.update({"category": "structure", "weight": 1, "semantic": False})
            elif criterion_id.startswith("preservation__"):
                row.update({"category": "integrity", "weight": 10, "semantic": False})
            elif (
                criterion_id.startswith("covenant_formula__")
                or criterion_id.startswith("borrowing_base_formula__")
                or criterion_id.startswith("wip_control_formula__")
            ):
                row.update({"category": "auditability", "weight": 5, "semantic": False})
            elif criterion_id.startswith("borrowing_base_label__"):
                row.update({
                    "category": "structure",
                    "weight": 1,
                    "semantic": criterion_id in semantic_ids,
                })
            elif criterion_id.startswith("support__"):
                row.update({"category": "controls", "weight": 1, "semantic": False})
            elif criterion_id.startswith("lender_review__"):
                row.update({"category": "decision", "weight": 10, "semantic": False})
            elif criterion_id.startswith("wip_control_value__"):
                field = criterion_id.removeprefix("wip_control_value__")
                row.update({
                    "category": "controls" if field.endswith("project_count") else "core_finance",
                    "weight": 5 if field.endswith("project_count") else 10,
                    "semantic": False,
                })
            else:
                row.update({"category": "core_finance", "weight": 10, "semantic": False})
        elif task_id == "task_024":
            # This is a decision workbook, so its committed and indicative
            # funding conclusions must carry materially more importance than
            # workbook labels and layout.  Supporting facts and formula lineage
            # still earn proportional credit; no individual miss applies a
            # fixed scalar-reward cap.
            row.pop("failure_cap", None)
            if (
                criterion_id.startswith("structure__")
                or criterion_id in {
                    "funding_views__table_structure",
                    "funding_views__metric_order",
                }
                or criterion_id.startswith("summary_label__")
                or criterion_id.endswith("__present")
            ):
                row.update({"category": "structure", "weight": 1})
            elif criterion_id.startswith("request_formula__"):
                row.update({"category": "auditability", "weight": 3, "semantic": False})
            elif criterion_id.startswith("summary_formula__"):
                row.update({"category": "auditability", "weight": 3, "semantic": False})
            elif criterion_id.startswith("funding_views_formula__"):
                row.update({"category": "auditability", "weight": 5, "semantic": False})
            elif criterion_id == "controls__formula_count":
                row.update({"category": "controls", "weight": 5, "semantic": False})
            elif criterion_id.startswith("request__"):
                field = criterion_id.rsplit("__", 1)[-1]
                if field in {
                    "approved_amount",
                    "equipment_line_proceeds",
                    "cash_required",
                }:
                    row.update({"category": "core_finance", "weight": 10, "semantic": False})
                elif field in {"decision", "funding", "classification"}:
                    row.update({
                        "category": "decision",
                        "weight": 10,
                        "semantic": criterion_id in semantic_ids,
                    })
                elif field in {"request_amount", "annual_savings", "simple_payback"}:
                    row.update({"category": "core_finance", "weight": 5, "semantic": False})
                elif field == "asset":
                    row.update({
                        "category": "provenance",
                        "weight": 3,
                        "semantic": criterion_id in semantic_ids,
                    })
                else:  # pragma: no cover - guarded by the canonical id set
                    raise ValueError(f"unclassified task_024 request criterion: {criterion_id}")
            elif criterion_id.startswith("summary_value__"):
                field = criterion_id.removeprefix("summary_value__")
                if field == "approved_annual_savings":
                    row.update({"category": "core_finance", "weight": 5, "semantic": False})
                else:
                    row.update({"category": "decision", "weight": 10, "semantic": False})
            elif criterion_id.startswith("funding_views__"):
                column = criterion_id.rsplit("__", 1)[-1]
                if column in {"committed", "indicative"}:
                    row.update({"category": "decision", "weight": 10, "semantic": False})
                elif column == "difference":
                    row.update({"category": "controls", "weight": 5, "semantic": False})
                else:  # pragma: no cover - guarded by the canonical id set
                    raise ValueError(f"unclassified task_024 funding view: {criterion_id}")
            else:  # pragma: no cover - sources are handled before this branch
                raise ValueError(f"unclassified task_024 criterion: {criterion_id}")
        elif task_id == "task_015":
            # The slide is graded as a finance deliverable, not as an exact
            # template-reproduction exercise.  Extra/changed slide structure,
            # styling, or a normal status synonym loses only its own weight;
            # it must never erase correct covenant calculations elsewhere.
            row.pop("failure_cap", None)
            if criterion_id.startswith("preservation__"):
                row.update({"category": "integrity", "weight": 1, "semantic": False})
            elif criterion_id.startswith("structure__"):
                row.update({"category": "structure", "weight": 1, "semantic": criterion_id in semantic_ids})
            elif criterion_id.startswith("style__"):
                row.update({"category": "presentation_quality", "weight": 3, "semantic": False})
            elif criterion_id.startswith("metric__"):
                is_status = criterion_id.endswith("__status")
                row.update({
                    "category": "decision" if is_status else "core_finance",
                    "weight": 5 if is_status else 10,
                    "semantic": criterion_id in semantic_ids,
                })
            elif criterion_id.startswith("status__"):
                row.update({
                    "category": "decision",
                    "weight": 5,
                    "semantic": criterion_id in semantic_ids,
                })
            elif criterion_id.startswith("headroom__"):
                row.update({
                    "category": "decision",
                    "weight": 10,
                    "semantic": criterion_id in semantic_ids,
                })
            elif criterion_id.startswith("calculation_support__equation__"):
                row.update({
                    "category": "controls",
                    "weight": 10,
                    "semantic": criterion_id in semantic_ids,
                })
            elif criterion_id.startswith("calculation_support__"):
                row.update({
                    "category": "core_finance",
                    "weight": 5,
                    "semantic": criterion_id in semantic_ids,
                })
            elif criterion_id.startswith("sensitivity__"):
                is_status = criterion_id.endswith("__status")
                is_equation = "__equation_" in criterion_id
                row.update({
                    "category": "decision" if is_status else ("controls" if is_equation else "core_finance"),
                    "weight": 5 if is_status or is_equation else 10,
                    "semantic": criterion_id in semantic_ids,
                })
            else:  # pragma: no cover - guarded by the stable specification
                raise ValueError(f"unclassified task_015 criterion: {criterion_id}")
        elif task_id == "task_025" and criterion_id == "identity__status":
            row.update({"category": "controls", "weight": 3})
            row.pop("failure_cap", None)
        elif task_id == "task_025" and (
            criterion_id in {
                "wip__proposed", "wip__unposted", "covenants__compliance",
                "capex__classification_labels",
            }
            or re.match(r"capex__cx-\d{2}-\d{3}__", criterion_id)
        ):
            row.update({"category": "decision", "weight": 10, "failure_cap": 0.49})
        elif row["semantic"] and row.get("category") == "decision":
            row.update({"weight": 10, "failure_cap": 0.49})
        policy_by_id[criterion_id] = {
            key: row[key]
            for key in ("category", "weight", "semantic", "failure_cap")
            if key in row
        }

    actual_ids = {str(row["id"]) for row in result.get("criteria", [])}
    if actual_ids != set(policy_by_id):
        raise ValueError(
            f"legacy artifact rubric ids drifted for {task_id}: "
            f"missing={sorted(set(policy_by_id) - actual_ids)}, "
            f"extra={sorted(actual_ids - set(policy_by_id))}"
        )
    review = result.get("semantic_review")
    if isinstance(review, dict):
        reviewed = {
            str(row.get("criterion_id"))
            for row in review.get("criteria", [])
            if isinstance(row, dict)
        }
        if reviewed != semantic_ids:
            raise ValueError(
                f"legacy semantic route drifted for {task_id}: "
                f"missing={sorted(semantic_ids - reviewed)}, extra={sorted(reviewed - semantic_ids)}"
            )
    for row in result["criteria"]:
        row.pop("failure_cap", None)
        row.update(policy_by_id[str(row["id"])])
    return result


def _attach_semantic_review(
    result: dict[str, Any],
    *,
    task_id: str,
    evidence: str,
    artifact_type: str,
    specs: list[dict[str, Any]],
    decision_failure_cap: float | None = 0.49,
    always_judge: bool = False,
    execution_mode: str | None = None,
) -> dict[str, Any]:
    # Assign the objective rubric policy before routing any criterion. Semantic
    # verification is a grading method, not an importance category: a title
    # stays weight 1, provenance stays weight 3, and only a true central
    # decision receives the weight-10 failure cap.
    result = attach_default_policy(result)
    by_id = {str(row["id"]): row for row in result.get("criteria", [])}
    reviews = []
    for spec in specs:
        criterion_id = str(spec["criterion_id"])
        criterion = by_id[criterion_id]
        # A criterion routed to the semantic verifier is semantic by
        # construction. Mark it explicitly so the released catalog, weights,
        # failure caps, and runtime result cannot drift from the review plan.
        criterion["semantic"] = True
        if criterion.get("category") == "decision":
            criterion["weight"] = 10
            if decision_failure_cap is None:
                criterion.pop("failure_cap", None)
            else:
                criterion["failure_cap"] = decision_failure_cap
        requirement = (
            f"Decide whether the submitted {artifact_type} evidence satisfies: "
            f"{criterion['description']}"
            if execution_mode == "scoped_per_criterion"
            else semantic_requirement(
                criterion_id=criterion_id,
                description=str(criterion["description"]),
                expected_facts=spec.get("expected_facts"),
                artifact_type=artifact_type,
            )
        )
        review = {
            "criterion_id": criterion_id,
            "requirement": requirement,
            "hard_gate_met": bool(spec["hard_gate_met"]),
            "hard_gate_evidence": str(spec["hard_gate_evidence"]),
            "legacy_lexical_match": bool(criterion["value"]),
            "always_judge": bool(spec.get("always_judge", always_judge)),
        }
        if execution_mode == "scoped_per_criterion":
            review.update({
                "submitted_evidence": str(spec.get("submitted_evidence") or "")[:12_000],
                "reference_context": spec.get("reference_context") or spec.get("expected_facts") or {},
                "task_context": spec.get("task_context") or {},
                "evidence_scope": str(spec.get("evidence_scope") or "criterion_field"),
            })
        reviews.append(review)
    result["semantic_review"] = {
        "version": 2,
        "mode": "deterministic_hard_gates_plus_bounded_semantic_judge",
        "task_id": task_id,
        "artifact": None,
        "evidence": evidence[:60_000],
        "criteria": reviews,
        "execution_mode": execution_mode or "legacy_batched",
        "policy": "Semantic wording cannot rescue a failed objective numeric, formula, or structure gate.",
    }
    return result


def _legacy_artifact_evidence(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        presentation = Presentation(path)
        chunks = []
        for index, slide in enumerate(presentation.slides, start=1):
            chunks.append(f"[Slide {index}]")
            for shape in slide.shapes:
                if getattr(shape, "text", ""):
                    chunks.append(shape.text)
                if getattr(shape, "has_table", False):
                    chunks.append("TABLE:")
                    chunks.extend(
                        " | ".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
        return "\n".join(chunks)[:60_000]
    if suffix == ".docx":
        return _document_text(Document(path))[:60_000]
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=False, read_only=False)
        chunks = []
        for sheet in workbook.worksheets:
            chunks.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows():
                values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                if values:
                    chunks.append(" | ".join(values))
        return "\n".join(chunks)[:60_000]
    return ""


def _formula_text(cell: Any) -> str:
    return str(getattr(cell, "value", "") or "").lower().replace("$", "").replace("'", "").replace(" ", "")


def _formula_has_large_literal(cell: Any) -> bool:
    formula = _formula_text(cell)
    if not formula.startswith("="):
        return False
    without_references = re.sub(r"(?:[a-z0-9_ ]+!)?[a-z]{1,3}[0-9]+", "", formula)
    return any(abs(float(token)) >= 10_000 for token in re.findall(r"(?<![a-z])\d+(?:\.\d+)?", without_references))


def _recalculated_data_workbook(path: Path):
    """Return a data-only workbook after a safe LibreOffice recalc copy."""
    formula_workbook = load_workbook(path, data_only=False, read_only=False)
    original_values = load_workbook(path, data_only=True, read_only=False)

    def cached_formula_values(candidate) -> int:
        count = 0
        for sheet in formula_workbook.worksheets:
            if sheet.title not in candidate.sheetnames:
                continue
            value_sheet = candidate[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        if value_sheet[cell.coordinate].value is not None:
                            count += 1
        return count

    def preserves_submitted_cache(candidate) -> bool:
        for sheet in formula_workbook.worksheets:
            if sheet.title not in original_values.sheetnames or sheet.title not in candidate.sheetnames:
                continue
            original_sheet = original_values[sheet.title]
            candidate_sheet = candidate[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    if original_sheet[cell.coordinate].value is not None and candidate_sheet[cell.coordinate].value is None:
                        return False
        return True

    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return original_values
    with tempfile.TemporaryDirectory(prefix="arm-grade-xlsx-") as directory:
        root = Path(directory)
        source_dir = root / "source"
        output_dir = root / "output"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        try:
            subprocess.run(
                [executable, f"-env:UserInstallation={root.as_uri()}/profile", "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=30,
                check=False,
                env={**os.environ, "HOME": str(root)},
            )
        except (OSError, subprocess.TimeoutExpired):
            return original_values
        recalculated = output_dir / path.name
        if recalculated.exists():
            # load_workbook reads the entire ZIP before the temporary directory closes.
            recalculated_values = load_workbook(recalculated, data_only=True, read_only=False)
            # Some LibreOffice builds successfully emit an xlsx but discard
            # cached formula results.  Never replace a more complete submitted
            # cache with a worse conversion.
            if (
                cached_formula_values(recalculated_values) >= cached_formula_values(original_values)
                and preserves_submitted_cache(recalculated_values)
            ):
                return recalculated_values
    return original_values


def _task_004_treatment_ok(value: Any, expected: dict[str, Any]) -> bool:
    text = str(value or "")
    normalized = _normalize(text)
    if expected["netted_claim_or_backcharge"]:
        # A disputed backcharge recovery is a cost-recovery judgment. Requiring
        # the word "revenue" here misstates the accounting issue and rejects
        # ordinary controller language such as "de-netted from PM ETC."
        wrong_direction = any(
            token in normalized
            for token in (
                "do not add back", "not add back", "leave netted", "remain netted",
                "keep netted", "retain the recovery offset", "retain recovery offset",
            )
        )
        treatment_ok = not wrong_direction and (
            any(token in normalized for token in ("backcharge", "recovery"))
            and any(
                token in normalized
                for token in (
                    "do not net", "not netted", "de net", "denet", "add back",
                    "added back", "restore", "restored", "remove from pm etc",
                    "removed from pm etc", "cost added",
                )
            )
        )
        return treatment_ok and _text_contains_number(
            text, expected["cost_amount"], abs_tol=1.0
        )

    wrong_revenue_direction = any(
        token in normalized
        for token in (
            "do not exclude", "not excluded", "include revenue", "revenue included",
            "recognize revenue", "recognise revenue", "revenue recognized",
            "revenue recognised",
        )
    )
    revenue_ok = not wrong_revenue_direction and (
        any(
            token in normalized
            for token in (
                "exclude", "excluded", "omit", "omitted", "no revenue",
                "not recognize", "not recognised", "unapproved", "pending revenue",
            )
        )
        and any(
            token in normalized
            for token in (
                "revenue", "recovery", "change order", "commercial value", "pco",
            )
        )
        and _text_contains_number(text, expected["revenue_amount"], abs_tol=1.0)
    )
    cost_ok = (
        any(
            token in normalized
            for token in (
                "include", "included", "add", "added", "carry", "carried",
                "record", "recorded", "overlay", "probable cost",
            )
        )
        and any(token in normalized for token in ("cost", "exposure", "etc", "overlay"))
    )
    return revenue_ok and cost_ok and _text_contains_number(text, expected["cost_amount"], abs_tol=1.0)


def _task_004_source_ok(value: Any, expected: dict[str, Any]) -> bool:
    text = _normalize(value)
    reference = _normalize(expected["reference"])
    return reference in text and (
        "fin rev 04" in text or "signed wip policy" in text or "wip policy" in text
    )


def _task_004_has_phrase(text: str, phrase: str) -> bool:
    """Match a normalized word/phrase without substring accidents.

    In particular, ``hold`` must not match the final five letters of
    ``threshold`` in an otherwise correct release conclusion.
    """

    return bool(re.search(
        rf"(?<![a-z0-9]){re.escape(_normalize(phrase))}(?![a-z0-9])",
        _normalize(text),
    ))


def _task_004_follow_up_ok(value: Any, expected: dict[str, Any]) -> bool:
    text = _normalize(value)
    return all(
        any(_normalize(token) in text for token in token_group)
        for token_group in expected["required_follow_up_tokens"]
    )


def _task_004_review_queue_ok(value: Any, expected: list[str]) -> bool:
    text = _normalize(value)
    held = {_normalize(project_id) for project_id in expected}
    released = {
        _normalize(project_id) for project_id in _TASK_004_PROJECT_IDS
    } - held
    if not all(project_id in text for project_id in held):
        return False
    held_direction_present = any(
        token in text
        for token in (
            "held", "hold", "review required", "controller review",
            "pending review", "requires review",
        )
    )
    if not held_direction_present:
        return False
    for project_id in held:
        if any(
            phrase in text
            for phrase in (
                f"{project_id} released", f"release {project_id}",
                f"{project_id} cleared", f"clear {project_id}",
            )
        ):
            return False
    if not released:
        return True
    return all(
        project_id in text
        and any(
            phrase in text
            for phrase in (
                f"{project_id} released", f"release {project_id}",
                f"{project_id} cleared", f"clear {project_id}",
                f"{project_id} no exception", f"{project_id} may release",
            )
        )
        for project_id in released
    )


def _task_004_close_release_ok(
    value: Any,
    expected_hold: list[str] | tuple[str, ...] | None = None,
) -> bool:
    text = _normalize(value)
    held = {_normalize(project_id) for project_id in (expected_hold or ())}
    released = {
        _normalize(project_id) for project_id in _TASK_004_PROJECT_IDS
    } - held
    for project_id in held:
        if any(
            _task_004_has_phrase(text, token)
            for token in (
                f"{project_id} released", f"release {project_id}",
                f"{project_id} cleared", f"clear {project_id}",
            )
        ):
            return False
    if any(
        _task_004_has_phrase(text, token)
        for token in (
            "do not hold", "not held", "release now", "release immediately",
            "cleared for release", "clear for release",
        )
    ):
        return False
    hold_supported = _task_004_has_phrase(text, "hold") and any(
        _task_004_has_phrase(text, token)
        for token in ("review", "controller", "signoff", "sign off", "approval", "pending")
    )
    if not hold_supported or not released:
        return hold_supported
    return all(
        project_id in text
        and any(
            _task_004_has_phrase(text, token)
            for token in (
                f"{project_id} released", f"release {project_id}",
                f"{project_id} cleared", f"clear {project_id}",
                f"{project_id} no exception", f"{project_id} may release",
            )
        )
        for project_id in released
    )


_TASK_004_HEADER_ALIASES: dict[str, frozenset[str]] = {
    "project_id": frozenset({
        "project id", "project", "job id", "job", "job number",
        "project number", "project code", "job code",
    }),
    "project_name": frozenset({
        "project name", "job name", "project description", "job description",
        "description",
    }),
    "current_contract": frozenset({
        "current contract", "contract value", "revised contract",
        "approved contract", "contract amount", "current contract value",
    }),
    "cost_to_date": frozenset({
        "posted cost to date", "cost to date", "job cost", "job cost to date",
        "actual cost", "incurred cost", "cost incurred",
    }),
    "pm_etc": frozenset({
        "pm etc", "submitted etc", "forecast etc", "project manager etc",
        "current etc", "pm forecast etc",
    }),
    "finance_adjustment": frozenset({
        "finance adjustment", "finance overlay", "etc adjustment",
        "documented etc overlay", "close adjustment", "risk adjustment",
        "finance etc adjustment", "forecast adjustment",
    }),
    "close_etc": frozenset({
        "close etc", "adjusted etc", "final etc", "recommended etc",
        "finance etc", "june close etc",
    }),
    "eac": frozenset({
        "eac", "close eac", "estimated cost at completion", "final cost",
        "forecast cost", "cost at completion", "june close eac",
    }),
    "percent_complete": frozenset({
        "complete", "percent complete", "completion", "pct complete",
        "percentage complete",
    }),
    "earned_revenue": frozenset({
        "earned revenue", "revenue earned", "earned", "revenue to date",
    }),
    "billings": frozenset({
        "billings", "billed to date", "cumulative billings", "total billed",
        "billing to date", "billings to date",
    }),
    "contract_asset": frozenset({
        "contract asset", "underbilling", "underbilled", "wip asset",
    }),
    "contract_liability": frozenset({
        "contract liability", "overbilling", "overbilled", "wip liability",
    }),
    "margin": frozenset({
        "margin", "total margin", "estimated margin", "gross margin",
        "close margin", "total est margin",
    }),
    "margin_percent": frozenset({
        "margin percent", "margin rate", "gm", "gross margin percent",
        "estimated margin percent", "close margin percent",
    }),
    "prior_margin_percent": frozenset({
        "prior margin percent", "may margin percent", "prior close margin percent",
        "previous margin percent", "prior gm", "may gm",
    }),
    "margin_movement_bps": frozenset({
        "margin movement bps", "bps change", "margin change bps",
        "movement bps", "change bps", "margin movement", "june movement bps", "bps",
        "margin movement pp", "margin change pp", "movement pp", "change pp",
        "june movement pp", "percentage point change", "change percentage points",
    }),
    "review_status": frozenset({
        "review status", "controller review", "controller review status",
        "close status", "status", "review conclusion", "release status",
        "release posture", "review decision", "controller status",
    }),
    "treatment": frozenset({
        "finance rationale", "accounting treatment", "commercial treatment",
        "rationale", "finance basis", "finance basis rationale", "accounting basis", "close rationale",
        "recommendation", "accounting conclusion", "finance conclusion",
        "close basis", "close treatment", "forecast treatment", "june treatment",
    }),
    "source_support": frozenset({
        "sources used", "source support", "evidence", "support", "sources",
        "source", "source s used", "controlling sources", "support used",
        "source tie", "supporting evidence", "basis source", "document trail",
    }),
    "next_action": frozenset({
        "next step owner", "next action", "action owner", "owner action",
        "next step", "action", "follow up", "follow up owner", "action required",
        "follow up owner", "follow-up", "follow-up owner", "owner next step",
        "owner and timing", "accountable follow-through",
    }),
}


def _task_004_header_field(value: Any) -> str | None:
    raw = str(value or "").casefold()
    normalized = _normalize(value)
    if "bps" in raw or ("margin" in normalized and "movement" in normalized):
        return "margin_movement_bps"
    if "%" in raw or "percent" in normalized or "pct" in normalized or "rate" in normalized:
        if any(token in normalized for token in ("prior", "may", "previous")):
            return "prior_margin_percent"
        if "complete" in normalized or "completion" in normalized:
            return "percent_complete"
        if "margin" in normalized or re.search(r"\bgm\b", normalized):
            return "margin_percent"
    unitless = re.sub(
        r"\b(?:usd|dollars?|000s?|thousands?|millions?|mm)\b", " ", normalized
    )
    unitless = re.sub(r"\s+", " ", unitless).strip()
    for field, aliases in _TASK_004_HEADER_ALIASES.items():
        if normalized in aliases or unitless in aliases:
            return field
    return None


def _task_004_project_identity(value: Any) -> str | None:
    match = re.search(r"\bARM[\s_-]?(\d{4})\b", str(value or ""), re.IGNORECASE)
    if not match:
        return None
    project_id = f"ARM-{match.group(1)}"
    return project_id if project_id in {"ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506"} else None


def _task_004_locate_table(workbook: Any) -> dict[str, Any] | None:
    best: tuple[int, int, Any, int, dict[str, int]] | None = None
    core_fields = {
        "current_contract", "cost_to_date", "pm_etc", "finance_adjustment",
        "close_etc", "eac", "earned_revenue", "billings", "margin",
    }
    for sheet_index, sheet in enumerate(workbook.worksheets):
        for row_number in range(1, min(sheet.max_row, 60) + 1):
            columns: dict[str, int] = {}
            for cell in sheet[row_number][: min(sheet.max_column, 80)]:
                field = _task_004_header_field(cell.value)
                if field is not None and field not in columns:
                    columns[field] = cell.column
            if "project_id" not in columns or len(core_fields.intersection(columns)) < 5:
                continue
            score = len(columns) * 10 + len(core_fields.intersection(columns))
            candidate = (score, -sheet_index, sheet, row_number, columns)
            if best is None or candidate[:2] > best[:2]:
                best = candidate
    if best is None:
        return None
    _, _, sheet, header_row, columns = best
    project_rows: dict[str, int] = {}
    for row_number in range(header_row + 1, min(sheet.max_row, header_row + 250) + 1):
        values = [cell.value for cell in sheet[row_number][: min(sheet.max_column, 80)]]
        identities = {identity for value in values if (identity := _task_004_project_identity(value))}
        if len(identities) == 1:
            project_rows.setdefault(identities.pop(), row_number)
    total_row = None
    total_labels = {"total", "portfolio total", "grand total", "four project total", "combined total"}
    identity_columns = [columns.get("project_id"), columns.get("project_name")]
    for row_number in range(header_row + 1, min(sheet.max_row, header_row + 250) + 1):
        labels = {
            _normalize(sheet.cell(row_number, column).value)
            for column in identity_columns
            if column is not None
        }
        if labels.intersection(total_labels):
            total_row = row_number
            break
    return {
        "sheet": sheet,
        "sheet_name": sheet.title,
        "header_row": header_row,
        "columns": columns,
        "project_rows": project_rows,
        "total_row": total_row,
    }


def _task_004_cell(sheet: Any, location: dict[str, Any] | None, row: int | None, field: str) -> Any | None:
    if location is None or row is None:
        return None
    column = location["columns"].get(field)
    return sheet.cell(row, column) if column is not None else None


def _task_004_header_text(location: dict[str, Any] | None, field: str) -> str:
    if location is None:
        return ""
    column = location["columns"].get(field)
    if column is None:
        return ""
    return str(location["sheet"].cell(location["header_row"], column).value or "")


def _task_004_currency_scale(header: str) -> float:
    raw = str(header or "").casefold()
    normalized = _normalize(header)
    if re.search(r"\bmm\b", raw) or "million" in normalized:
        return 1_000_000.0
    if "000" in raw or re.search(r"\b000s\b", raw) or "thousand" in normalized:
        return 1_000.0
    return 1.0


def _task_004_number_matches(field: str, actual: Any, expected: float, header: str = "") -> bool:
    number = _number(actual)
    if number is None:
        return False
    number = float(number)
    if field in {"percent_complete", "margin_percent", "prior_margin_percent"}:
        return any(
            _close(candidate, expected, abs_tol=0.00005, rel_tol=0.0)
            for candidate in (number, number / 100.0)
        )
    if field == "margin_movement_bps":
        return any(
            _close(candidate, expected, abs_tol=0.5, rel_tol=0.0)
            for candidate in (number, number * 100.0, number * 10_000.0)
        )
    scaled = number * _task_004_currency_scale(header)
    return _close(scaled, expected, abs_tol=0.05, rel_tol=0.0)


def _task_004_formula_cells_ok(
    sheet: Any,
    location: dict[str, Any] | None,
    row: int | None,
    fields: tuple[str, ...],
) -> tuple[bool, str]:
    evidence = []
    checks = []
    for field in fields:
        cell = _task_004_cell(sheet, location, row, field)
        formula = getattr(cell, "value", None)
        is_formula = isinstance(formula, str) and formula.startswith("=")
        no_hardcoded_output = field == "margin_movement_bps" or not _formula_has_large_literal(cell)
        checks.append(is_formula and no_hardcoded_output)
        evidence.append(f"{getattr(cell, 'coordinate', field)}={formula!r}")
    return bool(checks) and all(checks), "; ".join(evidence)


def _task_004_display_value(dollars: float, header: str) -> float:
    return dollars / _task_004_currency_scale(header)


def _task_004_percent_display(decimal_value: float, current_value: Any) -> float:
    current = _number(current_value)
    return decimal_value * 100.0 if current is not None and abs(float(current)) > 2.0 else decimal_value


def _task_004_force_recalculation(workbook: Any) -> None:
    if workbook.calculation is None:
        workbook.calculation = CalcProperties(calcMode="auto")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


_TASK_004_PROJECT_IDS = ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506")


_TASK_004_PROJECT_NAMES = frozenset({
    "columbia biologics clean utilities",
    "northline cold storage expansion",
    "cascade health ambulatory pavilion",
    "eastbank data hall cooling expansion",
})


_TASK_004_TOTAL_LABELS = frozenset({
    "total", "portfolio total", "grand total", "four project total", "combined total",
})


def _task_004_anchor_rows(
    workbook: Any,
    *,
    project_id: str | None = None,
    total: bool = False,
) -> list[tuple[Any, int]]:
    """Find candidate project/total rows without interpreting their headers."""

    rows: list[tuple[Any, int]] = []
    for sheet in workbook.worksheets:
        for row_number in range(1, min(sheet.max_row, 300) + 1):
            values = [
                cell.value
                for cell in sheet[row_number][: min(sheet.max_column, 100)]
            ]
            if project_id is not None and any(
                _task_004_project_identity(value) == project_id for value in values
            ):
                rows.append((sheet, row_number))
                continue
            if total and any(
                _normalize(value) in _TASK_004_TOTAL_LABELS for value in values
            ):
                rows.append((sheet, row_number))
    return rows


def _task_004_nearest_header_row(sheet: Any, row_number: int) -> int | None:
    candidates: list[tuple[int, int]] = []
    for candidate in range(max(1, row_number - 40), row_number):
        values = [
            cell.value
            for cell in sheet[candidate][: min(sheet.max_column, 100)]
        ]
        string_count = sum(
            isinstance(value, str) and bool(value.strip()) for value in values
        )
        if string_count >= 2:
            candidates.append((string_count, candidate))
    return max(candidates)[1] if candidates else None


def _task_004_row_evidence(sheet: Any, row_number: int) -> str:
    header_row = _task_004_nearest_header_row(sheet, row_number)
    parts: list[str] = []
    for cell in sheet[row_number][: min(sheet.max_column, 100)]:
        if cell.value is None or str(cell.value).strip() == "":
            continue
        header = (
            str(sheet.cell(header_row, cell.column).value or "").strip()
            if header_row is not None else ""
        )
        label = header or cell.coordinate
        parts.append(f"{label} [{cell.coordinate}]={cell.value!r}")
    return (
        f"[{sheet.title} row {row_number}; header row {header_row}]\n"
        + " | ".join(parts)
    )


def _task_004_header_context(sheet: Any, row_number: int, column: int) -> str:
    context: list[str] = []
    header_row = _task_004_nearest_header_row(sheet, row_number)
    if header_row is None:
        return ""
    # Include the actual header and nearby schedule-level unit captions, never
    # intervening prior data rows (whose values such as 10,000 would otherwise
    # be mistaken for a "$000" unit declaration).
    for candidate in range(max(1, header_row - 3), header_row + 1):
        direct = sheet.cell(candidate, column).value
        if direct is not None:
            context.append(str(direct))
        for cell in sheet[candidate][: min(sheet.max_column, 20)]:
            if cell.value is not None and cell.column != column:
                context.append(str(cell.value))
    return " | ".join(context)


def _task_004_numeric_candidate_evidence(
    workbook: Any,
    *,
    field: str,
    expected: float,
    project_id: str | None = None,
    total: bool = False,
) -> str:
    candidates: list[str] = []
    for sheet, row_number in _task_004_anchor_rows(
        workbook, project_id=project_id, total=total
    ):
        matches: list[str] = []
        for cell in sheet[row_number][: min(sheet.max_column, 100)]:
            if _task_004_project_identity(cell.value):
                continue
            header = _task_004_header_context(sheet, row_number, cell.column)
            if _task_004_number_matches(field, cell.value, expected, header):
                matches.append(f"{cell.coordinate}={cell.value!r}")
        if matches:
            candidates.append(
                f"MATCHED CELLS: {', '.join(matches)}\n"
                f"{_task_004_row_evidence(sheet, row_number)}"
            )
    return "\n\n".join(candidates)[:12_000]


def _task_004_authored_row_evidence(
    workbook: Any,
    *,
    project_id: str | None = None,
    total: bool = False,
) -> str:
    candidates: list[str] = []
    for sheet, row_number in _task_004_anchor_rows(
        workbook, project_id=project_id, total=total
    ):
        authored = []
        for cell in sheet[row_number][: min(sheet.max_column, 100)]:
            if not isinstance(cell.value, str) or not cell.value.strip():
                continue
            normalized = _normalize(cell.value)
            if _task_004_project_identity(cell.value):
                continue
            if normalized in _TASK_004_PROJECT_NAMES or normalized in _TASK_004_TOTAL_LABELS:
                continue
            authored.append(cell.value)
        if authored:
            candidates.append(_task_004_row_evidence(sheet, row_number))
    return "\n\n".join(candidates)[:12_000]


def _task_004_shared_cutoff_action_evidence(
    workbook: Any,
    *,
    cutoff_reference: str,
) -> str:
    """Return only shared AP/cutoff rows that can apply to one project action.

    A preparer may assign one AP owner to the four-invoice cutoff population
    instead of repeating that action in every project row.  Keep this evidence
    narrow: the row must either name the project's invoice/reference or clearly
    scope an AP/cutoff action to all four project invoices.
    """

    reference = _normalize(cutoff_reference)
    candidates: list[str] = []
    for sheet in workbook.worksheets:
        for row_number in range(1, min(sheet.max_row, 300) + 1):
            values = [
                cell.value
                for cell in sheet[row_number][: min(sheet.max_column, 100)]
                if cell.value is not None and str(cell.value).strip()
            ]
            text = _normalize(" ".join(str(value) for value in values))
            if not text:
                continue
            cutoff_context = any(
                token in text
                for token in ("ap cutoff", "accrual", "accrue", "invoice", "incurred cost")
            )
            reference_scope = bool(reference and reference in text)
            portfolio_scope = any(
                token in text
                for token in (
                    "all four", "four june dated", "four project invoices",
                    "portfolio", "each project", "every project", "all project",
                    "june project invoice", "june dated project invoice",
                    "project cutoff population", "invoice cutoff population",
                )
            )
            if cutoff_context and (reference_scope or portfolio_scope):
                candidates.append(_task_004_row_evidence(sheet, row_number))
    return "\n\n".join(candidates)[:8_000]


def _task_004_shared_policy_evidence(workbook: Any) -> str:
    """Return a narrowly scoped portfolio-wide signed-policy source row.

    A source register may identify FIN-REV-04 once and state that it applies
    throughout the workbook. Supplying that row to each source-use semantic
    review prevents placement-based false negatives without deterministically
    deciding whether the project's commercial and cutoff sources are adequate.
    """

    candidates: list[str] = []
    for sheet in workbook.worksheets:
        for row_number in range(1, min(sheet.max_row, 500) + 1):
            values = [
                cell.value
                for cell in sheet[row_number][: min(sheet.max_column, 100)]
                if cell.value is not None and str(cell.value).strip()
            ]
            text = _normalize(" ".join(str(value) for value in values))
            if not text:
                continue
            policy_named = any(
                token in text
                for token in (
                    "revenue recognition", "signed wip policy",
                    "signed revenue recognition policy", "fin rev 04",
                )
            )
            portfolio_scope = any(
                token in text
                for token in (
                    "applied throughout", "applies throughout", "used throughout",
                    "applies to all four", "governing all four", "all four",
                )
            )
            if policy_named and portfolio_scope:
                candidates.append(_task_004_row_evidence(sheet, row_number))
    return "\n\n".join(candidates)[:8_000]


def _task_004_formula_candidate_evidence(
    workbook: Any,
    *,
    project_id: str | None = None,
    total: bool = False,
    minimum_formulas: int,
) -> str:
    candidates: list[str] = []
    for sheet, row_number in _task_004_anchor_rows(
        workbook, project_id=project_id, total=total
    ):
        formulas = [
            f"{cell.coordinate}={cell.value}"
            for cell in sheet[row_number][: min(sheet.max_column, 100)]
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        if len(formulas) >= minimum_formulas:
            candidates.append(
                f"FORMULAS: {'; '.join(formulas)}\n"
                f"{_task_004_row_evidence(sheet, row_number)}"
            )
    return "\n\n".join(candidates)[:12_000]


def _task_004_perturbation_checks_v2(
    path: Path,
) -> tuple[dict[str, bool], bool, str] | None:
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        return None
    try:
        current_values = _recalculated_data_workbook(path)
        with tempfile.TemporaryDirectory(prefix="task004-lineage-") as directory:
            test_path = Path(directory) / path.name
            workbook = load_workbook(path, data_only=False, read_only=False)
            location = _task_004_locate_table(workbook)
            if location is None:
                return ({project_id: False for project_id in ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506")}, False, "main table not found")
            formula_sheet = location["sheet"]
            value_sheet = current_values[location["sheet_name"]]
            expected_by_project: dict[str, dict[str, float]] = {}
            for index, project_id in enumerate(("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506"), start=1):
                row = location["project_rows"].get(project_id)
                if row is None:
                    continue
                displayed: dict[str, float] = {}
                dollars: dict[str, float] = {}
                for field in ("current_contract", "cost_to_date", "pm_etc", "finance_adjustment", "billings"):
                    value_cell = _task_004_cell(value_sheet, location, row, field)
                    number = _number(getattr(value_cell, "value", None))
                    if number is None:
                        break
                    scale = _task_004_currency_scale(_task_004_header_text(location, field))
                    dollars[field] = float(number) * scale
                    displayed[field] = float(number)
                else:
                    prior_cell = _task_004_cell(value_sheet, location, row, "prior_margin_percent")
                    prior_number = _number(getattr(prior_cell, "value", None))
                    if prior_number is None:
                        continue
                    prior_decimal = float(prior_number) / 100.0 if abs(float(prior_number)) > 2.0 else float(prior_number)
                    dollar_deltas = {
                        "current_contract": 12_345.0 + index,
                        "cost_to_date": 3_211.0 + index,
                        "pm_etc": 1_777.0 + index,
                        "finance_adjustment": 333.0 + index,
                        "billings": 2_444.0 + index,
                    }
                    for field, delta in dollar_deltas.items():
                        new_dollars = dollars[field] + delta
                        dollars[field] = new_dollars
                        target_cell = _task_004_cell(formula_sheet, location, row, field)
                        target_cell.value = _task_004_display_value(new_dollars, _task_004_header_text(location, field))
                    prior_decimal += 0.0007 + index * 0.00001
                    prior_target = _task_004_cell(formula_sheet, location, row, "prior_margin_percent")
                    prior_target.value = _task_004_percent_display(prior_decimal, getattr(prior_cell, "value", None))
                    close_etc = dollars["pm_etc"] + dollars["finance_adjustment"]
                    eac = dollars["cost_to_date"] + close_etc
                    percent_complete = min(dollars["cost_to_date"] / eac, 1.0)
                    earned_revenue = dollars["current_contract"] * percent_complete
                    margin = dollars["current_contract"] - eac
                    margin_percent = margin / dollars["current_contract"]
                    expected_by_project[project_id] = {
                        "current_contract": dollars["current_contract"],
                        "cost_to_date": dollars["cost_to_date"],
                        "pm_etc": dollars["pm_etc"],
                        "finance_adjustment": dollars["finance_adjustment"],
                        "close_etc": close_etc,
                        "eac": eac,
                        "percent_complete": percent_complete,
                        "earned_revenue": earned_revenue,
                        "billings": dollars["billings"],
                        "contract_asset": max(earned_revenue - dollars["billings"], 0.0),
                        "contract_liability": max(dollars["billings"] - earned_revenue, 0.0),
                        "margin": margin,
                        "margin_percent": margin_percent,
                        "prior_margin_percent": prior_decimal,
                        "margin_movement_bps": (margin_percent - prior_decimal) * 10_000.0,
                    }
            _task_004_force_recalculation(workbook)
            workbook.save(test_path)
            recalculated = _recalculated_data_workbook(test_path)
            recalculated_location = _task_004_locate_table(recalculated)
            if recalculated_location is None:
                return ({project_id: False for project_id in ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506")}, False, "perturbed table not found")
            recalculated_sheet = recalculated_location["sheet"]
            project_checks: dict[str, bool] = {}
            result_fields = (
                "close_etc", "eac", "percent_complete", "earned_revenue",
                "contract_asset", "contract_liability", "margin",
                "margin_percent", "margin_movement_bps",
            )
            for project_id in ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506"):
                row = recalculated_location["project_rows"].get(project_id)
                expected = expected_by_project.get(project_id)
                project_checks[project_id] = bool(expected and row) and all(
                    _task_004_number_matches(
                        field,
                        getattr(_task_004_cell(recalculated_sheet, recalculated_location, row, field), "value", None),
                        expected[field],
                        _task_004_header_text(recalculated_location, field),
                    )
                    for field in result_fields
                )
            totals: dict[str, float] = {}
            additive_fields = (
                "current_contract", "cost_to_date", "pm_etc", "finance_adjustment",
                "close_etc", "eac", "earned_revenue", "billings", "contract_asset",
                "contract_liability", "margin",
            )
            if len(expected_by_project) == 4:
                for field in additive_fields:
                    totals[field] = sum(expected[field] for expected in expected_by_project.values())
                totals["percent_complete"] = totals["cost_to_date"] / totals["eac"]
                totals["margin_percent"] = totals["margin"] / totals["current_contract"]
                totals["prior_margin_percent"] = sum(
                    expected["current_contract"] * expected["prior_margin_percent"]
                    for expected in expected_by_project.values()
                ) / totals["current_contract"]
                totals["margin_movement_bps"] = (
                    totals["margin_percent"] - totals["prior_margin_percent"]
                ) * 10_000.0
            total_row = recalculated_location["total_row"]
            total_ok = bool(total_row and len(totals) == 15) and all(
                _task_004_number_matches(
                    field,
                    getattr(_task_004_cell(recalculated_sheet, recalculated_location, total_row, field), "value", None),
                    expected,
                    _task_004_header_text(recalculated_location, field),
                )
                for field, expected in totals.items()
            )
            return project_checks, total_ok, f"project_checks={project_checks}; total_check={total_ok}"
    except Exception as exc:
        return ({project_id: False for project_id in ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506")}, False, f"perturbation error: {exc}")


def _task_004_review_status_lexical(value: Any, required: bool) -> bool:
    text = _normalize(value)
    if not text:
        return False
    release_terms = ("release", "monitor", "within threshold", "no review", "not required", "clear", "supported")
    hold_terms = ("hold", "review", "controller", "required", "escalate", "pending", "unresolved")
    if required:
        return not any(term in text for term in ("no review", "not required", "release")) and any(term in text for term in hold_terms)
    return any(term in text for term in release_terms) and not any(term in text for term in ("hold", "escalate"))


def _grade_task_004(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Shared/Finance/Close/2026/06 June/4 WIP/WIP risk cases_7.1 847am - NB REVIEW COPY.xlsx")
    path = workspace_root / relative
    gold = load_apex_gold("task_004")
    if not path.exists():
        return _file_failure("task_004", "missing workbook")
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
        values = _recalculated_data_workbook(path)
    except Exception as exc:
        return _file_failure("task_004", str(exc))

    location = _task_004_locate_table(workbook)
    value_location = _task_004_locate_table(values)
    formula_sheet = location["sheet"] if location else None
    value_sheet = values[value_location["sheet_name"]] if value_location else None
    perturbation = _task_004_perturbation_checks_v2(path) if location else None
    project_lineage = perturbation[0] if perturbation else {}
    total_lineage = perturbation[1] if perturbation else False
    perturbation_evidence = perturbation[2] if perturbation else "LibreOffice perturbation unavailable"
    criteria: list[Criterion] = []

    source_by_id = {row["project_id"]: row for row in gold["source_inputs"]}
    project_by_id = {row["project_id"]: row for row in gold["projects"]}
    review_by_id = {row["project_id"]: row for row in gold["review_rows"]}
    decision_by_id = {row["project_id"]: row for row in gold["decision_rows"]}
    input_targets = {
        "current_contract": "current_contract",
        "cost_to_date": "cost_to_date",
        "pm_etc": "pm_etc",
        "finance_adjustment": "documented_etc_overlay",
        "billings": "billings",
        "prior_margin_percent": "prior_close_margin_percent",
    }
    result_targets = {
        "close_etc": "estimated_cost_to_complete",
        "eac": "estimated_cost_at_completion",
        "percent_complete": "percent_complete",
        "earned_revenue": "earned_revenue",
        "contract_asset": "underbilling",
        "contract_liability": "overbilling",
        "margin": "estimated_total_margin",
        "margin_percent": "estimated_margin_percent",
        "margin_movement_bps": "margin_movement_bps",
    }
    semantic_specs: list[dict[str, Any]] = []
    task_context = {
        "assignment": "Prepare the four-project June 30 WIP risk review and release recommendation.",
        "grading_boundary": (
            "Amounts and formula behavior are graded independently. Judge only the supplied field, "
            "accept normal business labels and equivalent professional wording, and do not require an authored phrase."
        ),
    }

    def add_hybrid_numeric(
        *,
        criterion_id: str,
        description: str,
        field: str,
        expected: float,
        cell: Any | None,
        project_id: str | None = None,
        total: bool = False,
    ) -> None:
        actual = getattr(cell, "value", None)
        header = _task_004_header_text(value_location, field)
        canonical_met = _task_004_number_matches(field, actual, expected, header)
        candidate_evidence = _task_004_numeric_candidate_evidence(
            values,
            field=field,
            expected=expected,
            project_id=project_id,
            total=total,
        )
        primary_has_number = cell is not None and _number(actual) is not None
        canonical_primary_conflict = primary_has_number and not canonical_met
        hard_gate_met = canonical_met or (
            bool(candidate_evidence) and not canonical_primary_conflict
        )
        criteria.append(Criterion(
            criterion_id,
            description,
            canonical_met,
            (
                f"recognized_cell={getattr(cell, 'coordinate', None)}; "
                f"actual={actual!r}; expected={expected}; "
                f"candidate_present={bool(candidate_evidence)}; "
                f"canonical_primary_conflict={canonical_primary_conflict}"
            ),
            category="core_finance",
            weight=10,
            semantic=True,
        ))
        semantic_specs.append({
            "criterion_id": criterion_id,
            "expected_facts": {
                "required_metric": description,
                "expected_value": expected,
            },
            "hard_gate_met": hard_gate_met,
            "hard_gate_evidence": (
                f"recognized field contains conflicting numeric value {actual!r}"
                if canonical_primary_conflict
                else (
                    "exact objective value is present in a project/portfolio row"
                    if hard_gate_met
                    else "exact objective value is absent from the associated row"
                )
            ),
            "submitted_evidence": candidate_evidence,
            "reference_context": {
                "required_association": description,
                "expected_value": expected,
                "accepted_presentation": (
                    "Accept whole dollars, $000s, or $mm; decimal or percentage "
                    "rates; and basis points, percentage points, or equivalent "
                    "signed rate movement. The objective value has already been "
                    "verified by the hard gate. Decide only whether it is attached "
                    "to this metric and the correct project/portfolio total."
                ),
                "rejection_boundary": (
                    "Reject a correct number associated with another metric, period, "
                    "project, scenario, or an unlabeled numerical dump."
                ),
            },
            "task_context": task_context,
            "evidence_scope": (
                f"candidate rows for {project_id or 'portfolio total'} containing "
                "the exact objective value and their adjacent headers"
            ),
            "always_judge": False,
        })

    for project_id in ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506"):
        slug = project_id.casefold()
        formula_row = location["project_rows"].get(project_id) if location else None
        value_row = value_location["project_rows"].get(project_id) if value_location else None
        source = source_by_id[project_id]
        project = project_by_id[project_id]
        review = review_by_id[project_id]
        decision = decision_by_id[project_id]
        for field, target_field in input_targets.items():
            cell = _task_004_cell(value_sheet, value_location, value_row, field) if value_sheet else None
            expected = float(source[target_field])
            add_hybrid_numeric(
                criterion_id=f"input__{slug}__{field}",
                description=f"{project_id} {field.replace('_', ' ')} input is correct",
                field=field,
                expected=expected,
                cell=cell,
                project_id=project_id,
            )
        fallback_lineage, fallback_evidence = _task_004_formula_cells_ok(
            formula_sheet, location, formula_row, tuple(result_targets)
        ) if formula_sheet else (False, "main table not found")
        lineage_met = project_lineage.get(project_id, fallback_lineage)
        required_formula_fields = set(input_targets) | set(result_targets)
        formula_mapping_complete = bool(location and formula_row) and required_formula_fields.issubset(
            set(location["columns"])
        )
        formula_candidate = _task_004_formula_candidate_evidence(
            workbook, project_id=project_id, minimum_formulas=9
        )
        lineage_hard_gate = (
            lineage_met if formula_mapping_complete else bool(formula_candidate)
        )
        lineage_id = f"formula_lineage__{slug}"
        lineage_description = f"{project_id} close calculations respond to source inputs"
        criteria.append(Criterion(
            lineage_id,
            lineage_description,
            lineage_met,
            f"{perturbation_evidence}; fallback={fallback_lineage}; {fallback_evidence}",
            category="auditability", weight=5, semantic=True,
        ))
        semantic_specs.append({
            "criterion_id": lineage_id,
            "expected_facts": {"required_formula_behavior": lineage_description},
            "hard_gate_met": lineage_hard_gate,
            "hard_gate_evidence": (
                "recognized calculation fields passed the independent perturbation check"
                if lineage_met
                else (
                    "relabelled project row contains at least nine formulas for scoped semantic association"
                    if lineage_hard_gate
                    else "recognized calculation fields failed perturbation or insufficient formulas were present"
                )
            ),
            "submitted_evidence": formula_candidate,
            "reference_context": {
                "required_formula_behavior": (
                    "The close ETC, EAC, percent complete, earned revenue, gross "
                    "contract asset/liability, margin, margin rate, and movement must "
                    "be formulas that respond to the row's source inputs."
                ),
                "association_rule": (
                    "When headers were relabelled, determine whether the visible "
                    "formulas are attached to those business metrics. Do not award "
                    "credit merely because unrelated formulas occur in the row."
                ),
            },
            "task_context": task_context,
            "evidence_scope": f"formula-bearing candidate row for {project_id}",
            "always_judge": False,
        })
        for field, target_field in result_targets.items():
            cell = _task_004_cell(value_sheet, value_location, value_row, field) if value_sheet else None
            expected = float(project[target_field])
            add_hybrid_numeric(
                criterion_id=f"result__{slug}__{field}",
                description=f"{project_id} {field.replace('_', ' ')} is correct",
                field=field,
                expected=expected,
                cell=cell,
                project_id=project_id,
            )

        text_cells = {
            field: _task_004_cell(value_sheet, value_location, value_row, field)
            if value_sheet else None
            for field in ("review_status", "treatment", "source_support", "next_action")
        }
        text_values = {
            field: getattr(cell, "value", None)
            for field, cell in text_cells.items()
        }
        lexical = {
            "review_status": _task_004_review_status_lexical(text_values["review_status"], review["controller_review_required"]),
            "commercial_treatment": _task_004_treatment_ok(text_values["treatment"], review),
            "source_support": _task_004_source_ok(text_values["source_support"], review),
            "next_action": _task_004_follow_up_ok(text_values["next_action"], decision),
        }
        criterion_details = {
            "review_status": (
                f"{project_id} review status is supportable", "decision", 10,
                {"controller_review_required": review["controller_review_required"]},
                {
                    "required_direction": (
                        "Controller review/hold"
                        if review["controller_review_required"]
                        else (
                            "No exception-driven hold. Accept release, release subject to the "
                            "ordinary Controller sign-off boundary, within-threshold, supported, "
                            "or an equivalent no-hold conclusion. The status field need not use "
                            "the word monitoring when the project action is recorded separately."
                        )
                    ),
                    "equivalence_rule": (
                        "Judge the release-versus-exception-hold direction, not an authored status "
                        "phrase. Routine Controller sign-off is compatible with release and is not "
                        "the same as a special exception hold."
                    ),
                },
            ),
            "commercial_treatment": (
                f"{project_id} commercial treatment is supportable", "decision", 10,
                {
                    "pending_revenue": review["revenue_amount"],
                    "supported_cost_or_recovery_adjustment": review["cost_amount"],
                    "disputed_recovery_de_netted": review["netted_claim_or_backcharge"],
                    "june_cutoff_cost": review["cutoff_amount"],
                    "forecast_amount_transferred_from_etc": review["cutoff_etc_transfer"],
                    "unforecast_eac_increase": review["cutoff_unforecast_amount"],
                    "cutoff_reference": review["cutoff_reference"],
                },
                {
                    "required_treatment": (
                        "Do not recognize unsigned/pending commercial revenue at June 30; "
                        + (
                            "remove the disputed forecast recovery, count the shared PO-2409-117 work package only once, "
                            "remove its already-posted prior draw from remaining cost, accrue only the $27,196.69 June "
                            "portion of invoice 88417, leave the $6,600 July 2 line in ETC, and do not recognize a recovery "
                            "from the June 30 routing email without the filed agreement terms and referenced credit memo."
                            if review["netted_claim_or_backcharge"]
                            else (
                                (
                                    f"carry the supportable cost exposure; accrue the full June {review['cutoff_reference']} "
                                    f"invoice of ${review['cutoff_amount']:,.2f}, transfer only the ${review['cutoff_etc_transfer']:,.2f} "
                                    f"amount carried for that PO in PM ETC, and recognize the ${review['cutoff_unforecast_amount']:,.2f} "
                                    "residual as an EAC increase."
                                )
                                if review["cutoff_unforecast_amount"]
                                else (
                                    f"carry the supportable cost exposure and move the June {review['cutoff_reference']} "
                                    f"amount of ${review['cutoff_amount']:,.2f} from remaining ETC to incurred cost without "
                                    "changing EAC."
                                )
                            )
                        )
                    ),
                    "equivalence_rule": (
                        "Accept concise professional equivalents and a correctly calculated row as evidence of the same "
                        "treatment; amounts already present in the associated project row need not be repeated in prose."
                    ),
                    "primary_recommendation_rule": (
                        "When the recognized Finance basis/rationale field states a treatment, judge that primary "
                        "recommendation. A hypothetical, sensitivity, or future-action row elsewhere in the workbook "
                        "cannot cure a contradictory treatment in the primary project row."
                    ),
                    "cutoff_consistency_rule": (
                        "The primary project row and its rationale must actually apply the June cutoff treatment. "
                        "A row that leaves cost to date at the posted balance, retains the invoice in remaining ETC, "
                        "or describes accrual only as a future/pending decision has not applied the treatment, even "
                        "if another sentence identifies the invoice amount or its EAC exposure."
                    ),
                },
            ),
            "source_support": (
                f"{project_id} controlling support is identifiable", "provenance", 3,
                {
                    "commercial_reference": review["reference"],
                    "cutoff_reference": review["cutoff_reference"],
                    "cutoff_sources": "July 1 AP intake workpaper and ten-page invoice batch",
                    "policy_authority": "signed WIP policy FIN-REV-04",
                    **(
                        {"commercial_routing_source": "June 30 correspondence-routing email"}
                        if project_id == "ARM-2409" else {}
                    ),
                },
                {
                    "commercial_reference": review["reference"],
                    "cutoff_reference": review["cutoff_reference"],
                    "cutoff_sources": (
                        "AP cutoff + accrual list_7.1 OL.xlsx, invoice intake - 7.1; "
                        "batch 20260701-02.pdf"
                    ),
                    **(
                        {"commercial_routing_source": "Correspondence Archive/2026/06/30/sent 1706.eml"}
                        if project_id == "ARM-2409" else {}
                    ),
                    "policy_authority": "signed WIP policy FIN-REV-04",
                    "required_conjuncts": (
                        "Full credit requires identifiable support for all of the commercial matter, the "
                        "project's June invoice cutoff/reference (or its ordinary AP/invoice-batch source), "
                        "and the signed FIN-REV-04 policy authority. Generic project summaries, job-cost, WIP, "
                        "or trial-balance citations do not substitute for a missing cutoff or policy authority."
                    ),
                    "equivalence_rule": (
                        "Accept abbreviations, invoice/PO references, filenames, and ordinary reference variants when the "
                        "commercial, cutoff, and policy authorities remain identifiable. Do not require full paths."
                    ),
                },
            ),
            "next_action": (
                f"{project_id} next action and owner are sufficient", "decision", 10,
                {
                    "required_follow_up": decision["required_follow_up"],
                    "cutoff_action": (
                        f"AP or Finance owns the June cutoff decision for {review['cutoff_reference']}"
                    ),
                },
                {
                    "required_follow_up": decision["required_follow_up"],
                    "cutoff_action": (
                        f"Assign AP or Finance to prepare, decide, or record/accrue the June "
                        f"{review['cutoff_reference']} item. When the prompt prohibits posting, an explicit owned "
                        "decision or preparation step is sufficient; the ETC-to-incurred mechanics may be documented "
                        "in the separately graded treatment or calculation rather than repeated in the action field."
                    ),
                    "owner_rule": "Name a responsible person or role and preserve any Controller approval boundary; accept a reasonable accountable owner supported by the records.",
                    "project_specific_equivalents": (
                        "For ARM-2318, assign Project Accounting or the PM to reconcile the CRS-10942 / "
                        "PO-2318-102 invoice-over-forecast variance and monitor it after the completed Controller "
                        "estimate review; ordinary final sign-off may remain open, and the action need not repeat "
                        "the separately graded treatment or release status. For ARM-2409, obtaining or filing the "
                        "executed resolution and referenced credit memo, verifying its terms/effective "
                        "amount, and updating the commercial record is the supported action; do not "
                        "require another signature or trade acceptance already evidenced by the routing "
                        "email. For ARM-2506, PCO-006 is the commercial record for the same DB-27 pricing matter: "
                        "assigning the PM, Estimating, or Commercial owner to complete, obtain, finalize, route, or "
                        "approve PCO-006 or DB-27 pricing satisfies the commercial follow-up; do not require the "
                        "authored identifier or verb. For ARM-2417, pursuing an owner/CM decision, a written "
                        "validity extension, or confirmation of current work direction is a professional equivalent "
                        "way to resolve the unexecuted directive; do not require the authored phrase."
                    ),
                    "required_conjuncts": (
                        "Full credit requires both (1) the project-specific commercial/evidence follow-up and "
                        "responsible owner and (2) an owned June invoice-cutoff decision or preparation step for "
                        "the project's cutoff item. The second component may come from one clearly scoped shared "
                        "portfolio AP action supplied in the evidence; do not pass on the commercial action alone."
                    ),
                    "shared_cutoff_action_rule": (
                        "One clearly scoped portfolio-wide AP/Finance action covering all four June invoice-cutoff "
                        "items may satisfy the cutoff-action component for each project; do not require the same owner "
                        "and transfer instruction to be repeated four times."
                    ),
                    "equivalence_rule": (
                        "Accept any concise, project-specific action that assigns and pursues the "
                        "same unresolved evidence, decision, or authorization. Do not require the "
                        "authored sentence, exact verbs, or repetition of treatment already graded "
                        "in another field."
                    ),
                },
            ),
        }
        for field, (description, category, weight, expected_facts, reference_context) in criterion_details.items():
            criterion_id = f"{field}__{slug}"
            column_field = "treatment" if field == "commercial_treatment" else field
            cell = text_cells[column_field]
            direct_value = text_values[column_field]
            candidate_row = _task_004_authored_row_evidence(
                values, project_id=project_id
            )
            primary_row = (
                _task_004_row_evidence(value_sheet, value_row)
                if value_sheet is not None and value_row is not None
                else ""
            )
            associated_evidence = primary_row if cell is not None else candidate_row
            if field == "source_support":
                shared_policy = _task_004_shared_policy_evidence(values)
                if shared_policy:
                    associated_evidence = (
                        f"{associated_evidence}\n\nWorkbook-wide policy source evidence:\n{shared_policy}"
                    )
            if field == "next_action":
                shared_cutoff = _task_004_shared_cutoff_action_evidence(
                    values,
                    cutoff_reference=review["cutoff_reference"],
                )
                if shared_cutoff:
                    associated_evidence = (
                        f"{associated_evidence}\n\nShared cutoff-action evidence:\n{shared_cutoff}"
                    )
            submitted = (
                f"Project: {project_id}\n"
                f"Submitted {field.replace('_', ' ')}: {direct_value or ''}\n"
                f"Associated evidence: {associated_evidence}"
                if cell is not None
                else candidate_row
            )
            hard_gate = (
                bool(_normalize(direct_value))
                if cell is not None
                else bool(candidate_row)
            )
            criteria.append(Criterion(
                criterion_id, description, lexical[field], f"actual={direct_value!r}",
                category=category, weight=weight, semantic=True,
            ))
            semantic_specs.append({
                "criterion_id": criterion_id,
                "expected_facts": {"project_id": project_id, **expected_facts},
                "hard_gate_met": hard_gate,
                "hard_gate_evidence": (
                    f"recognized_field={cell is not None}; authored_candidate_present={hard_gate}"
                ),
                "submitted_evidence": submitted,
                "reference_context": reference_context,
                "task_context": task_context,
                "evidence_scope": (
                    f"{value_location['sheet_name'] if value_location else 'workbook'} "
                    f"primary project row for {project_id}, with the authored {field} field identified when present"
                    + (
                        "; narrowly scoped portfolio-wide signed-policy source rows are also included"
                        if field == "source_support" else ""
                    )
                    + (
                        "; narrowly scoped shared AP/cutoff action rows are also included"
                        if field == "next_action" else ""
                    )
                ),
            })

    total_fields = (
        "current_contract", "cost_to_date", "pm_etc", "finance_adjustment",
        "close_etc", "eac", "percent_complete", "earned_revenue", "billings",
        "contract_asset", "contract_liability", "margin", "margin_percent",
        "prior_margin_percent", "margin_movement_bps",
    )
    formula_total_row = location["total_row"] if location else None
    fallback_total, fallback_total_evidence = _task_004_formula_cells_ok(
        formula_sheet, location, formula_total_row, total_fields
    ) if formula_sheet else (False, "main table not found")
    total_mapping_complete = bool(location and formula_total_row) and set(total_fields).issubset(
        set(location["columns"])
    )
    total_formula_candidate = _task_004_formula_candidate_evidence(
        workbook, total=True, minimum_formulas=15
    )
    total_lineage_met = total_lineage if perturbation else fallback_total
    total_lineage_hard_gate = (
        total_lineage_met if total_mapping_complete else bool(total_formula_candidate)
    )
    criteria.append(Criterion(
        "total_formula_lineage", "Portfolio totals respond to the four project rows",
        total_lineage_met,
        f"{perturbation_evidence}; fallback={fallback_total}; {fallback_total_evidence}",
        category="auditability", weight=5, semantic=True,
    ))
    semantic_specs.append({
        "criterion_id": "total_formula_lineage",
        "expected_facts": {
            "required_formula_behavior": "Portfolio totals respond to the four project rows"
        },
        "hard_gate_met": total_lineage_hard_gate,
        "hard_gate_evidence": (
            "recognized total fields passed the independent perturbation check"
            if total_lineage_met
            else (
                "relabelled total row contains at least fifteen formulas for scoped semantic association"
                if total_lineage_hard_gate
                else "recognized total fields failed perturbation or insufficient total formulas were present"
            )
        ),
        "submitted_evidence": total_formula_candidate,
        "reference_context": {
            "required_formula_behavior": (
                "The portfolio amounts, rates, and movement are formulas derived from "
                "the four project rows, including weighted/derived rates rather than "
                "summing percentages."
            ),
            "association_rule": (
                "For relabelled headers, confirm the visible total-row formulas are "
                "associated with the stated portfolio metrics."
            ),
        },
        "task_context": task_context,
        "evidence_scope": "formula-bearing candidate portfolio-total row",
        "always_judge": False,
    })
    total_targets = {
        "current_contract": gold["totals"]["C11"],
        "cost_to_date": gold["totals"]["D11"],
        "pm_etc": gold["totals"]["E11"],
        "finance_adjustment": gold["totals"]["F11"],
        "close_etc": gold["totals"]["G11"],
        "eac": gold["totals"]["H11"],
        "percent_complete": gold["totals"]["I11"],
        "earned_revenue": gold["totals"]["J11"],
        "billings": gold["totals"]["K11"],
        "contract_asset": gold["totals"]["L11"],
        "contract_liability": gold["totals"]["M11"],
        "margin": gold["totals"]["N11"],
        "margin_percent": gold["totals"]["O11"],
        "prior_margin_percent": gold["totals"]["P11"],
        "margin_movement_bps": gold["totals"]["Q11"],
    }
    value_total_row = value_location["total_row"] if value_location else None
    for field, expected in total_targets.items():
        cell = _task_004_cell(value_sheet, value_location, value_total_row, field) if value_sheet else None
        add_hybrid_numeric(
            criterion_id=f"total_result__{field}",
            description=f"Portfolio {field.replace('_', ' ')} total is correct",
            field=field,
            expected=float(expected),
            cell=cell,
            total=True,
        )

    row_text = []
    portfolio_fields_present = False
    portfolio_mapping_present = bool(
        value_location
        and {"review_status", "next_action"}.intersection(value_location["columns"])
    )
    if value_sheet and value_location and portfolio_mapping_present:
        for project_id in ("ARM-2318", "ARM-2409", "ARM-2417", "ARM-2506"):
            row = value_location["project_rows"].get(project_id)
            status_value = getattr(_task_004_cell(value_sheet, value_location, row, "review_status"), "value", "") or ""
            action_value = getattr(_task_004_cell(value_sheet, value_location, row, "next_action"), "value", "") or ""
            portfolio_fields_present = portfolio_fields_present or bool(_normalize(status_value)) or bool(_normalize(action_value))
            row_text.append(
                f"{project_id}: status={status_value}; action={action_value}"
            )
        if value_total_row:
            total_status = getattr(_task_004_cell(value_sheet, value_location, value_total_row, "review_status"), "value", "") or ""
            total_action = getattr(_task_004_cell(value_sheet, value_location, value_total_row, "next_action"), "value", "") or ""
            portfolio_fields_present = portfolio_fields_present or bool(_normalize(total_status)) or bool(_normalize(total_action))
            row_text.append(
                f"Portfolio total: status={total_status}; action={total_action}"
            )
    if portfolio_mapping_present:
        portfolio_evidence = "\n".join(row_text)
    else:
        portfolio_evidence = "\n\n".join(
            value
            for value in (
                *(
                    _task_004_authored_row_evidence(values, project_id=project_id)
                    for project_id in _TASK_004_PROJECT_IDS
                ),
                _task_004_authored_row_evidence(values, total=True),
            )
            if value
        )[:12_000]
        portfolio_fields_present = bool(portfolio_evidence)
    held_projects = list(gold["decision_control"]["review_queue"])
    released_projects = [
        project_id for project_id in _TASK_004_PROJECT_IDS
        if project_id not in held_projects
    ]
    portfolio_rows = (
        (
            "portfolio__review_queue", "The portfolio review queue is correct",
            _task_004_review_queue_ok(portfolio_evidence, held_projects),
            {"held_for_controller_review": held_projects, "released_without_exception_hold": released_projects},
            (
                "Identify all four projects as requiring Controller review; order and punctuation do not matter."
                if not released_projects
                else "Identify the held projects and distinguish any project released without an exception hold; order and punctuation do not matter."
            ),
        ),
        (
            "portfolio__release_decision", "The portfolio release recommendation is supportable",
            _task_004_close_release_ok(portfolio_evidence, held_projects),
            {"release_without_exception_hold": released_projects, "hold_for_controller_review": held_projects},
            (
                "Accept an unambiguous recommendation to hold all four projects for Controller review before release."
                if not released_projects
                else "Accept an unambiguous recommendation that distinguishes released projects from those held for Controller review."
            ),
        ),
    )
    for criterion_id, description, lexical_met, expected_facts, equivalence in portfolio_rows:
        criteria.append(Criterion(
            criterion_id, description, lexical_met, portfolio_evidence,
            category="decision", weight=10, semantic=True,
        ))
        semantic_specs.append({
            "criterion_id": criterion_id,
            "expected_facts": expected_facts,
            "hard_gate_met": portfolio_fields_present,
            "hard_gate_evidence": f"portfolio and project status/action evidence present={portfolio_fields_present}",
            "submitted_evidence": portfolio_evidence,
            "reference_context": {"required_conclusion": expected_facts, "equivalence_rule": equivalence},
            "task_context": task_context,
            "evidence_scope": "four project status/action fields and portfolio total status/action fields",
        })

    result = _result(criteria)
    result = _attach_semantic_review(
        result,
        task_id="task_004",
        evidence=_legacy_artifact_evidence(path),
        artifact_type="controller WIP risk-review workbook",
        specs=semantic_specs,
        decision_failure_cap=None,
        always_judge=True,
        execution_mode="scoped_per_criterion",
    )
    result["task_grading_revision"] = dict(TASK_GRADING_REVISIONS["task_004"])
    return result


def _slide_text(slide) -> str:
    return "\n".join(getattr(shape, "text", "") for shape in slide.shapes if getattr(shape, "text", ""))


def _task_011_criterion_weight(criterion_id: str) -> int:
    """Weight decision-critical review work above repetitive support atoms.

    The weights reflect professional materiality: the June movement summary,
    qualifying callback focus, peer-normalized technician conclusion, and the
    alternatives/fallback decision are the point of the operating review.
    The underlying scorecard, native-object, sourcing, and readability checks
    remain independently creditable at ordinary weight.  This is linear and
    uncapped; it does not manufacture a pass/fail threshold.
    """
    if criterion_id.startswith((
        "finance__june_vs_may__",
        "finance__callback_focus__",
        "finance__alternative__",
        "finance__fallback__",
        "management__action_window__",
        "management__deferred_queue__",
    )):
        return 10
    if criterion_id.startswith("portfolio__"):
        return 5 if criterion_id.startswith("portfolio__action__") else 10
    if criterion_id.startswith("timing__"):
        return 10
    if criterion_id.startswith("signal__"):
        return 10
    if criterion_id.startswith("oot__"):
        return 10
    if criterion_id.startswith(("holdout__", "workpaper__")):
        return 10
    if criterion_id.startswith("capacity__"):
        return 5
    if criterion_id.startswith("finance__technician__") and criterion_id.endswith(
        ("__expected_callbacks", "__callback_excess")
    ):
        return 10
    if criterion_id.startswith("finance__callback_chart__") and criterion_id.endswith(
        "__expected_callbacks"
    ):
        return 10
    if criterion_id in {
        "management__remaining_controls",
        "management__selected_use",
        "management__selection_basis",
        "management__sequence__second",
        "management__sequence__third",
        "management__fallback_use",
        "management__fallback_trigger",
    }:
        return 10
    return 5


def _pptx_color_token(color: Any) -> str | None:
    try:
        rgb = color.rgb
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is not None:
        return f"rgb:{str(rgb).upper()}"
    try:
        theme = color.theme_color
    except (AttributeError, TypeError, ValueError):
        theme = None
    if theme is None:
        return None
    try:
        brightness = round(float(color.brightness), 3)
    except (AttributeError, TypeError, ValueError):
        brightness = 0.0
    return f"theme:{theme}:{brightness}"


def _pptx_fill_signature(fill: Any) -> tuple[str | None, str | None]:
    try:
        fill_type = str(fill.type) if fill.type is not None else None
    except (AttributeError, TypeError, ValueError):
        fill_type = None
    try:
        color = _pptx_color_token(fill.fore_color) if fill_type is not None else None
    except (AttributeError, TypeError, ValueError):
        color = None
    # PowerPoint and LibreOffice normalize an unfilled shape variously as no
    # fill, background fill, or opaque white. Those representations are
    # visually equivalent in this deck and must not create a false integrity
    # failure. Meaningful deck colors remain part of the signature.
    if color in {None, "rgb:FFFFFF"}:
        return None, None
    return "color", color


def _pptx_font_signature(font: Any) -> tuple[float | None, bool | None, bool | None, str, str | None]:
    try:
        size = float(round(float(font.size.pt))) if font.size is not None else None
    except (AttributeError, TypeError, ValueError):
        size = None
    underline_value = getattr(font, "underline", None)
    underline = (
        "underlined"
        if underline_value not in {None, False} and "NONE" not in str(underline_value).upper()
        else ""
    )
    return (
        size,
        True if getattr(font, "bold", None) is True else None,
        True if getattr(font, "italic", None) is True else None,
        underline,
        _pptx_color_token(getattr(font, "color", None)),
    )


def _pptx_font_observations(text_frame: Any) -> list[tuple[float | None, bool | None, str | None]]:
    observations: list[tuple[float | None, bool | None, str | None]] = []
    for paragraph in text_frame.paragraphs:
        fonts = [run.font for run in paragraph.runs] or [paragraph.font]
        for font in fonts:
            size, bold, _, _, color = _pptx_font_signature(font)
            observations.append((size, bold, color))
    return observations


def _pptx_slide_colors(slide: Any) -> set[str]:
    colors: set[str] = set()

    def add(color: str | None) -> None:
        if color:
            colors.add(color)

    try:
        add(_pptx_color_token(slide.background.fill.fore_color))
    except (AttributeError, TypeError, ValueError):
        pass
    for shape in slide.shapes:
        try:
            add(_pptx_color_token(shape.fill.fore_color))
        except (AttributeError, TypeError, ValueError):
            pass
        try:
            add(_pptx_color_token(shape.line.color))
        except (AttributeError, TypeError, ValueError):
            pass
        if getattr(shape, "has_text_frame", False):
            for _, _, color in _pptx_font_observations(shape.text_frame):
                add(color)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    try:
                        add(_pptx_color_token(cell.fill.fore_color))
                    except (AttributeError, TypeError, ValueError):
                        pass
                    for _, _, color in _pptx_font_observations(cell.text_frame):
                        add(color)
    return colors


def _pptx_colors_match(first: str | None, second: str | None) -> bool:
    if first is None or second is None:
        return False
    if first == second:
        return True
    if not first.startswith("rgb:") or not second.startswith("rgb:"):
        return False
    try:
        first_rgb = tuple(int(first[index:index + 2], 16) for index in (4, 6, 8))
        second_rgb = tuple(int(second[index:index + 2], 16) for index in (4, 6, 8))
    except (TypeError, ValueError):
        return False
    # Small export or theme-conversion shifts remain acceptable; a different
    # presentation palette does not. This is intentionally not pixel matching.
    return math.sqrt(sum((left - right) ** 2 for left, right in zip(first_rgb, second_rgb))) <= 36


def _pptx_background_color(slide: Any) -> str | None:
    try:
        return _pptx_color_token(slide.background.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def _pptx_task_015_style_checks(
    presentation: Any,
    original: Any,
    added: Any,
    table_shape: Any | None,
) -> tuple[bool, str, bool, str, bool, str]:
    if added is None:
        return (
            False, "appended slide missing",
            False, "appended slide missing",
            False, "appended slide missing",
        )

    reference_slides = list(original.slides)[1:] or list(original.slides)
    reference_backgrounds = {
        color for color in (_pptx_background_color(slide) for slide in reference_slides) if color
    }
    reference_colors = set().union(*(_pptx_slide_colors(slide) for slide in reference_slides))
    added_colors = _pptx_slide_colors(added)
    added_background = _pptx_background_color(added)
    distinctive_reference = {
        color
        for color in reference_colors
        if color not in reference_backgrounds and color not in {"rgb:FFFFFF", "rgb:000000"}
    }
    matched_colors = {
        reference
        for reference in distinctive_reference
        if any(_pptx_colors_match(reference, candidate) for candidate in added_colors)
    }
    background_matches = any(
        _pptx_colors_match(added_background, reference)
        for reference in reference_backgrounds
    )
    # A newly appended slide may inherit its background from the deck master,
    # in which case python-pptx exposes no explicit slide-level fill.  Matching
    # the established palette is substantive evidence of visual-system use and
    # should not be erased by that harmless implementation detail.
    inherited_background = added_background is None
    visual_system_ok = (background_matches or inherited_background) and len(matched_colors) >= 2
    visual_evidence = (
        f"background={added_background}; reference_backgrounds={sorted(reference_backgrounds)}; "
        f"inherited_background={inherited_background}; "
        f"matched_reference_colors={sorted(matched_colors)}"
    )

    explicit_title = getattr(added.shapes, "title", None)
    title_candidates: list[tuple[float, float, Any]] = []
    for shape in added.shapes:
        if not (
            getattr(shape, "has_text_frame", False)
            and _normalize(getattr(shape, "text", ""))
        ):
            continue
        top_ratio = float(getattr(shape, "top", 0) or 0) / max(
            presentation.slide_height,
            1,
        )
        if top_ratio >= .28:
            continue
        sizes = [
            size
            for size, _, _ in _pptx_font_observations(shape.text_frame)
            if size is not None
        ]
        max_size = max(sizes) if sizes else 0.0
        # A decision slide can legitimately place a large KPI near its title.
        # Font size alone then mistakes the KPI for the title.  Weight the
        # candidate's horizontal title span as well as its type size so that a
        # broad headline wins over a compact metric card without requiring a
        # particular placeholder, shape name, or phrase.
        width_ratio = float(getattr(shape, "width", 0) or 0) / max(
            presentation.slide_width,
            1,
        )
        title_candidates.append((max_size * max(width_ratio, .01), max_size, shape))
    title_shape = (
        explicit_title
        if explicit_title is not None
        and _normalize(getattr(explicit_title, "text", ""))
        else max(title_candidates, key=lambda item: item[0])[2]
        if title_candidates
        else None
    )
    reference_title_sizes: list[float] = []
    reference_title_tops: list[float] = []
    for slide in reference_slides:
        candidates: list[tuple[float, Any]] = []
        for shape in slide.shapes:
            if not _normalize(getattr(shape, "text", "")):
                continue
            sizes = [
                size
                for size, _, _ in _pptx_font_observations(shape.text_frame)
                if size is not None
            ] if getattr(shape, "has_text_frame", False) else []
            if sizes and float(getattr(shape, "top", 0) or 0) / max(presentation.slide_height, 1) < .25:
                candidates.append((max(sizes), shape))
        if candidates:
            size, shape = max(candidates, key=lambda item: item[0])
            reference_title_sizes.append(size)
            reference_title_tops.append(float(shape.top) / max(presentation.slide_height, 1))
    sorted_sizes = sorted(reference_title_sizes)
    reference_title_size = (
        sorted_sizes[len(sorted_sizes) // 2]
        if sorted_sizes
        else 20.0
    )
    title_observations = (
        _pptx_font_observations(title_shape.text_frame)
        if title_shape is not None and getattr(title_shape, "has_text_frame", False)
        else []
    )
    title_sizes = [size for size, _, _ in title_observations if size is not None]
    title_bold = any(bold is True for _, bold, _ in title_observations)
    title_top = (
        float(title_shape.top) / max(presentation.slide_height, 1)
        if title_shape is not None
        else 1.0
    )
    title_top_limit = min(.25, max(reference_title_tops, default=.17) + .08)
    title_hierarchy_ok = bool(
        title_shape is not None
        and title_sizes
        and max(title_sizes) >= reference_title_size * .8
        and title_bold
        and title_top <= title_top_limit
    )
    title_evidence = (
        f"title_present={title_shape is not None}; max_size={max(title_sizes) if title_sizes else None}; "
        f"reference_size={reference_title_size}; bold={title_bold}; "
        f"top_ratio={round(title_top, 3)}; top_limit={round(title_top_limit, 3)}"
    )

    table_readability_ok = False
    table_evidence = "reviewable table or card layout missing"
    if table_shape is not None:
        table = table_shape.table
        header_cells = list(table.rows[0].cells) if len(table.rows) else []
        body_cells = [cell for row in list(table.rows)[1:] for cell in row.cells]
        header_fills = {
            color
            for cell in header_cells
            for color in [_pptx_fill_signature(cell.fill)[1]]
            if color
        }
        body_fills = {
            color
            for cell in body_cells
            for color in [_pptx_fill_signature(cell.fill)[1]]
            if color
        }
        header_bold = bool(header_cells) and all(
            any(bold is True for _, bold, _ in _pptx_font_observations(cell.text_frame))
            for cell in header_cells
        )
        header_fill_distinct = bool(header_fills) and not all(
            any(_pptx_colors_match(header, body) for body in body_fills)
            for header in header_fills
        )
        table_properties = table._tbl.tblPr
        style_nodes = table_properties.xpath("./a:tableStyleId")
        uses_header_table_style = bool(
            table_properties.get("firstRow") == "1"
            and style_nodes
            and str(style_nodes[0].text or "").strip()
        )
        within_slide = (
            table_shape.left >= 0
            and table_shape.top >= 0
            and table_shape.left + table_shape.width <= presentation.slide_width
            and table_shape.top + table_shape.height <= presentation.slide_height
        )
        below_title = bool(
            title_shape is not None
            and table_shape.top >= title_shape.top + title_shape.height
        )
        table_readability_ok = bool(
            len(table.rows) >= 4
            and len(table.columns) >= 5
            and within_slide
            and below_title
            and (header_bold or header_fill_distinct or uses_header_table_style)
        )
        table_evidence = (
            f"rows={len(table.rows)}; columns={len(table.columns)}; within_slide={within_slide}; "
            f"below_title={below_title}; header_bold={header_bold}; "
            f"header_fill_distinct={header_fill_distinct}; "
            f"uses_header_table_style={uses_header_table_style}"
        )
    else:
        title_bottom = (
            title_shape.top + title_shape.height
            if title_shape is not None
            else 0
        )
        bounded_content_blocks = [
            shape
            for shape in added.shapes
            if getattr(shape, "has_text_frame", False)
            and _normalize(getattr(shape, "text", ""))
            and shape is not title_shape
            and shape.top >= title_bottom
            and shape.left >= 0
            and shape.top >= 0
            and shape.left + shape.width <= presentation.slide_width
            and shape.top + shape.height <= presentation.slide_height
        ]
        table_readability_ok = len(bounded_content_blocks) >= 3
        table_evidence = (
            "native_table=False; bounded_content_blocks="
            f"{len(bounded_content_blocks)}; required_for_card_layout=3"
        )
    return (
        visual_system_ok, visual_evidence,
        title_hierarchy_ok, title_evidence,
        table_readability_ok, table_evidence,
    )


def _grade_task_015(workspace_root: Path) -> dict[str, Any]:
    """Grade the covenant slide without treating authored labels as passwords.

    Objective values are hard-gated deterministically. Meaning-dependent
    associations, conclusions, sources, and bridge explanations are reviewed
    semantically from criterion-scoped slide evidence.
    """

    relative = Path(
        "Shared/Finance/Treasury/Bank - covenants/2026 Q2 working/"
        "Q2 lender update - review working v3.pptx"
    )
    path = workspace_root / relative
    gold = load_apex_gold("task_015")
    if not path.exists():
        return _file_failure("task_015", "missing deck")
    try:
        presentation = Presentation(str(path))
        original = Presentation(str(SEED_WORKSPACE / relative))
    except Exception as exc:
        return _file_failure("task_015", str(exc))

    def contains_scoped_number(
        text: str,
        target: float,
        *,
        abs_tol: float,
        kind: str,
    ) -> bool:
        """Match Task015 values without letting dates or other units satisfy them.

        The general presentation matcher intentionally accepts ordinary display
        rounding.  This covenant slide mixes dates, dollars, percentages, and
        ratios in tight evidence slices, so a bare date component such as ``7``
        must not satisfy a 6.53x FCCR and a 0.37x ratio must not satisfy a zero-
        dollar retained-WIP value.
        """

        tokens = re.findall(
            r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?(?:[kmb]|x|%)?\)?",
            text,
            flags=re.I,
        )
        for token in tokens:
            normalized = token.strip().strip("()").casefold()
            has_currency = "$" in normalized
            has_money_suffix = bool(re.search(r"[kmb]\s*$", normalized))
            has_ratio_suffix = bool(re.search(r"x\s*$", normalized))
            has_percent_suffix = normalized.endswith("%")
            parsed = _number(token)
            if parsed is None:
                continue
            if kind == "money":
                if not (
                    has_currency
                    or has_money_suffix
                    or (
                        parsed == 0
                        and normalized in {"0", "-0"}
                    )
                    or (
                        not has_ratio_suffix
                        and not has_percent_suffix
                        and abs(parsed) >= 100
                    )
                ):
                    continue
            elif kind == "ratio":
                if has_currency or has_money_suffix or has_percent_suffix:
                    continue
                if (
                    not has_ratio_suffix
                    and "." not in normalized
                    and not float(target).is_integer()
                ):
                    continue
            elif kind == "percent":
                if not has_percent_suffix:
                    continue
            else:
                raise ValueError(f"unsupported Task015 numeric kind: {kind}")

            tolerance = _display_rounding_tolerance(token, floor=abs_tol)
            if kind == "ratio" and "." not in normalized:
                # Whole-number covenant displays remain acceptable for whole
                # thresholds, but not as a half-point catch-all for another
                # non-integer ratio.
                tolerance = abs_tol
            if _close(token, target, abs_tol=tolerance, rel_tol=2e-6):
                return True
        return False

    original_texts = [_normalize(_slide_text(slide)) for slide in original.slides]
    current_texts = [_normalize(_slide_text(slide)) for slide in presentation.slides]
    original_structure_signatures = [
        tuple(sorted(str(shape.shape_type) for shape in slide.shapes))
        for slide in original.slides
    ]
    current_structure_signatures = [
        tuple(sorted(str(shape.shape_type) for shape in slide.shapes))
        for slide in presentation.slides
    ]
    exactly_one_appended = len(presentation.slides) == len(original.slides) + 1
    added = presentation.slides[-1] if len(presentation.slides) > len(original.slides) else None
    added_raw_text = _slide_text(added) if added is not None else ""
    added_text = _normalize(added_raw_text)
    table_shapes = [
        shape
        for shape in added.shapes
        if getattr(shape, "has_table", False)
    ] if added is not None else []

    def table_rows(shape: Any) -> list[list[str]]:
        return [
            [str(cell.text or "").strip() for cell in row.cells]
            for row in shape.table.rows
        ]

    def row_text(values: list[str]) -> str:
        return " | ".join(value for value in values if str(value).strip())

    capacity = gold["deterioration_capacity"]
    metric_targets = {
        "leverage": {
            "posted": (gold["leverage_posted"], .0001),
            "pro_forma": (gold["leverage_pro_forma"], .0001),
            "threshold": (gold["max_leverage"], .0001),
            "posted_headroom": (capacity["leverage"]["posted"], 1.0),
            "pro_forma_headroom": (capacity["leverage"]["pro_forma"], 1.0),
        },
        "fccr": {
            "posted": (gold["fccr_posted"], .0001),
            "pro_forma": (gold["fccr_pro_forma"], .0001),
            "threshold": (gold["min_fccr"], .0001),
            "posted_headroom": (capacity["fccr"]["posted"], 1.0),
            "pro_forma_headroom": (capacity["fccr"]["pro_forma"], 1.0),
        },
        "tangible_net_worth": {
            "posted": (gold["tangible_net_worth_posted"], 1.0),
            "pro_forma": (gold["tangible_net_worth_pro_forma"], 1.0),
            "threshold": (gold["min_tangible_net_worth"], 1.0),
            "posted_headroom": (
                capacity["tangible_net_worth"]["posted"],
                1.0,
            ),
            "pro_forma_headroom": (
                capacity["tangible_net_worth"]["pro_forma"],
                1.0,
            ),
        },
    }
    main_values = [
        target
        for fields in metric_targets.values()
        for target, _ in fields.values()
    ]

    def main_table_score(shape: Any) -> int:
        rows = table_rows(shape)
        if len(rows) < 4 or len(rows[0]) < 4:
            return -1
        text = row_text([value for row in rows for value in row])
        normalized = _normalize(text)
        numeric_hits = sum(
            _text_contains_number(text, value, abs_tol=(.0001 if abs(value) < 10 else 1.0))
            for value in main_values
        )
        concept_hits = sum(
            token in normalized
            for token in (
                "leverage", "fccr", "fixed charge coverage", "tangible net worth",
                "tnw", "posted", "recorded", "pro forma", "adjusted basis",
                "threshold", "limit", "headroom", "capacity",
            )
        )
        return numeric_hits * 10 + concept_hits

    main_table_shape = None
    if table_shapes:
        candidate = max(table_shapes, key=main_table_score)
        if main_table_score(candidate) > 0:
            main_table_shape = candidate

    (
        visual_system_ok,
        visual_system_evidence,
        title_hierarchy_ok,
        title_hierarchy_evidence,
        table_readability_ok,
        table_readability_evidence,
    ) = _pptx_task_015_style_checks(
        presentation,
        original,
        added,
        main_table_shape,
    )

    non_table_shapes = [
        shape
        for shape in added.shapes
        if getattr(shape, "has_text_frame", False)
        and not getattr(shape, "has_table", False)
        and str(shape.text or "").strip()
    ] if added is not None else []
    non_table_blocks = [
        str(shape.text or "").strip()
        for shape in non_table_shapes
    ]
    table_records: list[dict[str, Any]] = []
    table_text_blocks: list[str] = []
    for shape in table_shapes:
        rows = table_rows(shape)
        if not rows:
            continue
        header = row_text(rows[0])
        table_text_blocks.append("\n".join(row_text(row) for row in rows))
        for values in rows[1:]:
            body = row_text(values)
            table_records.append({
                "body": body,
                "evidence": f"Header: {header}\nRow: {body}",
                "shape": shape,
                "layout_scope": "native_row",
            })
    # A professional PowerPoint comparison is often authored as aligned text
    # boxes rather than a native table.  Preserve the visible row/column
    # relationships for deterministic gates and criterion-scoped semantic
    # review.  Treating every text box as an isolated record strips the labels
    # from its numbers and can both deny correct work and match an unrelated
    # number elsewhere on the slide.
    aligned_grid_records: list[dict[str, Any]] = []
    if added is not None and non_table_shapes:
        slide_width = max(float(presentation.slide_width), 1.0)
        slide_height = max(float(presentation.slide_height), 1.0)
        row_tolerance = slide_height * .022
        x_tolerance = slide_width * .018

        positioned = sorted(
            non_table_shapes,
            key=lambda shape: (
                float(shape.top) + float(shape.height) / 2,
                float(shape.left),
            ),
        )
        row_groups: list[list[Any]] = []
        row_centers: list[float] = []
        for shape in positioned:
            center = float(shape.top) + float(shape.height) / 2
            if row_groups and abs(center - row_centers[-1]) <= row_tolerance:
                row_groups[-1].append(shape)
                row_centers[-1] = sum(
                    float(item.top) + float(item.height) / 2
                    for item in row_groups[-1]
                ) / len(row_groups[-1])
            else:
                row_groups.append([shape])
                row_centers.append(center)
        row_groups = [
            sorted(group, key=lambda shape: float(shape.left))
            for group in row_groups
            if len(group) >= 4
        ]

        def aligned_count(first: list[Any], second: list[Any]) -> int:
            second_lefts = [float(shape.left) for shape in second]
            return sum(
                any(abs(float(shape.left) - left) <= x_tolerance for left in second_lefts)
                for shape in first
            )

        header_group: list[Any] | None = None
        grid_rows: list[list[Any]] = []
        for index, candidate in enumerate(row_groups):
            followers = [
                group
                for group in row_groups[index + 1:]
                if aligned_count(candidate, group) >= max(3, len(candidate) // 2)
            ]
            if len(followers) >= 2:
                header_group = candidate
                grid_rows = followers
                break

        if header_group is not None:
            # Keep only genuine grid anchors.  A nearby KPI card can share the
            # header's vertical band but will not align with the repeated body
            # columns.
            header_group = [
                header
                for header in header_group
                if sum(
                    any(
                        abs(float(shape.left) - float(header.left)) <= x_tolerance
                        for shape in group
                    )
                    for group in grid_rows
                ) >= 2
            ]
            header_texts = [str(shape.text or "").strip() for shape in header_group]
            header_context = " | ".join(header_texts)
            for group in grid_rows:
                values = [str(shape.text or "").strip() for shape in group]
                body = " | ".join(values)
                aligned_grid_records.append({
                    "body": body,
                    "evidence": f"Column context: {header_context}\nRow: {body}",
                    "shape": None,
                    "layout_scope": "aligned_row",
                })

            for header in header_group:
                header_left = float(header.left)
                column_lines: list[str] = []
                for group in grid_rows:
                    label = str(group[0].text or "").strip()
                    candidates = sorted(
                        group,
                        key=lambda shape: abs(float(shape.left) - header_left),
                    )
                    if not candidates or abs(float(candidates[0].left) - header_left) > x_tolerance:
                        continue
                    value = str(candidates[0].text or "").strip()
                    column_lines.append(f"{label}: {value}")
                if len(column_lines) >= 2:
                    header_text = str(header.text or "").strip()
                    body = "\n".join([header_text, *column_lines])
                    aligned_grid_records.append({
                        "body": body,
                        "evidence": f"Column: {header_text}\n" + "\n".join(column_lines),
                        "shape": None,
                        "layout_scope": "aligned_column",
                    })

    evidence_records = [
        *table_records,
        *aligned_grid_records,
        *(
            {"body": block, "evidence": block, "shape": None}
            for block in non_table_blocks
        ),
    ]
    all_slide_evidence = "\n".join(
        [*non_table_blocks, *table_text_blocks]
    )

    metric_aliases = {
        "leverage": ("leverage", "debt / ebitda", "funded debt"),
        "fccr": ("fccr", "fixed charge coverage", "fixed-charge coverage"),
        "tangible_net_worth": ("tangible net worth", "tnw", "net worth"),
    }

    def metric_record(metric: str) -> dict[str, Any] | None:
        targets = metric_targets[metric]
        best: tuple[int, dict[str, Any]] | None = None
        for record in evidence_records:
            body = str(record["body"])
            normalized = _normalize(body)
            numeric_hits = sum(
                contains_scoped_number(
                    body,
                    target,
                    abs_tol=tolerance,
                    kind=(
                        "ratio"
                        if metric in {"leverage", "fccr"}
                        and field in {"posted", "pro_forma", "threshold"}
                        else "money"
                    ),
                )
                for field, (target, tolerance) in targets.items()
            )
            name_hint = any(alias in normalized for alias in metric_aliases[metric])
            layout_scope = str(record.get("layout_scope") or "")
            association_bonus = (
                5_000
                if layout_scope in {"native_row", "aligned_row"}
                else 0
            )
            score = association_bonus + numeric_hits * 100 + (10 if name_hint else 0)
            if score and (best is None or score > best[0]):
                best = (score, record)
        return best[1] if best else None

    appended_title = ""
    if added is not None:
        title_shape = getattr(added.shapes, "title", None)
        appended_title = str(getattr(title_shape, "text", "") or "").strip()
        if not appended_title:
            appended_title = next(
                (
                    block
                    for block in non_table_blocks
                    if "covenant" in _normalize(block)
                ),
                non_table_blocks[0] if non_table_blocks else "",
            )

    criteria: list[Criterion] = [
        Criterion(
            "preservation__slide_count",
            "Exactly one slide is appended to the original deck",
            exactly_one_appended,
            f"original={len(original.slides)}; current={len(presentation.slides)}",
        ),
        *[
            Criterion(
                f"preservation__slide_{index:02d}",
                f"Original slide {index} retains its substantive content and structure",
                exactly_one_appended
                and index <= len(current_texts)
                and current_texts[index - 1] == text
                and current_structure_signatures[index - 1]
                == original_structure_signatures[index - 1],
                (
                    f"exactly_one_appended={exactly_one_appended}; "
                    f"text_match={index <= len(current_texts) and current_texts[index - 1] == text}; "
                    "shape_types_match="
                    f"{index <= len(current_structure_signatures) and current_structure_signatures[index - 1] == original_structure_signatures[index - 1]}"
                ),
            )
            for index, text in enumerate(original_texts, start=1)
        ],
        Criterion(
            "structure__title",
            "The appended slide has a decision-ready June 30 covenant-headroom title",
            bool(appended_title),
            appended_title,
            semantic=True,
        ),
        Criterion(
            "structure__table",
            "The appended slide contains a reviewable posted-versus-pro-forma covenant comparison",
            bool(main_table_shape),
            all_slide_evidence[:4000],
            semantic=True,
        ),
        Criterion(
            "style__visual_system",
            "The appended slide uses the existing deck background and visual palette",
            visual_system_ok,
            visual_system_evidence,
        ),
        Criterion(
            "style__title_hierarchy",
            "The appended slide title follows the deck's established hierarchy",
            title_hierarchy_ok,
            title_hierarchy_evidence,
        ),
        Criterion(
            "style__table_readability",
            "The covenant comparison uses a bounded, clearly headed, readable table or equivalent executive layout",
            table_readability_ok,
            table_readability_evidence,
        ),
    ]
    semantic_specs: list[dict[str, Any]] = []
    task_context = {
        "assignment": (
            "Append one decision-ready June 30 covenant-headroom slide to the "
            "five-slide Q2 lender-update working deck."
        ),
        "grading_boundary": (
            "Exact numeric presence is hard-gated separately. Interpret free business "
            "labels, metric association, calculation meaning, source authority, and "
            "conclusions semantically; accept ordinary professional equivalents."
        ),
        "reporting_basis": (
            f"Posted means recorded June 30 results. Close pro forma deducts the "
            f"unposted June AP cutoff expense of ${gold['pending_ap_cutoff_expense']:,.2f} "
            f"and adds the proposed, unposted June WIP adjustment of "
            f"${gold['proposed_wip_adjustment']:,.2f}."
        ),
    }

    def add_semantic_spec(
        criterion_id: str,
        *,
        expected_facts: dict[str, Any],
        hard_gate: bool,
        hard_gate_evidence: str,
        submitted_evidence: str,
        reference_context: dict[str, Any],
        evidence_scope: str,
    ) -> None:
        semantic_specs.append({
            "criterion_id": criterion_id,
            "expected_facts": expected_facts,
            "hard_gate_met": hard_gate,
            "hard_gate_evidence": hard_gate_evidence,
            "submitted_evidence": submitted_evidence[:7000],
            "reference_context": reference_context,
            "task_context": task_context,
            "evidence_scope": evidence_scope,
        })

    add_semantic_spec(
        "structure__title",
        expected_facts={
            "required_meaning": "June 30 Q2 covenant headroom on posted and close-pro-forma bases",
        },
        hard_gate=bool(appended_title),
        hard_gate_evidence=f"nonblank appended-slide title={bool(appended_title)}",
        submitted_evidence=appended_title,
        reference_context={
            "equivalence_rule": (
                "Accept normal executive title variants, punctuation, and posted/recorded "
                "or pro-forma/adjusted-basis equivalents."
            ),
        },
        evidence_scope="appended slide title only",
    )
    add_semantic_spec(
        "structure__table",
        expected_facts={
            "required_comparison": (
                "leverage, fixed-charge coverage, and tangible net worth on posted and "
                "close-pro-forma bases against executed thresholds"
            ),
        },
        hard_gate=bool(added) and bool(all_slide_evidence),
        hard_gate_evidence=(
            f"appended slide and nonblank slide content present={bool(added) and bool(all_slide_evidence)}"
        ),
        submitted_evidence=all_slide_evidence,
        reference_context={
            "equivalence_rule": (
                "Accept a real table, aligned comparison matrix, or equally clear bounded "
                "executive layout; do not require authored headers or table order."
            ),
        },
        evidence_scope="appended-slide covenant comparison only",
    )

    metric_names = {
        "leverage": "Leverage",
        "fccr": "Fixed-charge coverage",
        "tangible_net_worth": "Tangible net worth",
    }
    for metric, fields in metric_targets.items():
        record = metric_record(metric)
        evidence = str(record["evidence"]) if record else "required metric evidence is missing"
        for field, (target, tolerance) in fields.items():
            met = bool(record) and contains_scoped_number(
                evidence,
                target,
                abs_tol=tolerance,
                kind=(
                    "ratio"
                    if metric in {"leverage", "fccr"}
                    and field in {"posted", "pro_forma", "threshold"}
                    else "money"
                ),
            )
            criterion_id = f"metric__{metric}__{field}"
            criteria.append(Criterion(
                criterion_id,
                f"The {metric_names[metric]} comparison contains the correct {field.replace('_', ' ')} value",
                met,
                f"submitted={evidence[:1200]!r}; expected={target}; tolerance={tolerance}",
                semantic=True,
            ))
            add_semantic_spec(
                criterion_id,
                expected_facts={
                    "metric": metric_names[metric],
                    "field": field.replace("_", " "),
                    "required_value": target,
                },
                hard_gate=met,
                hard_gate_evidence=(
                    "the exact objective value is present in the candidate metric "
                    f"evidence={met}"
                ),
                submitted_evidence=evidence,
                reference_context={
                    "association_rule": (
                        "The exact number must be visibly associated with this metric and "
                        "this posted, close-pro-forma, threshold, or capacity field. Reject "
                        "a swapped column, another metric's row, or an unlabeled number dump."
                    ),
                    "equivalence_rule": (
                        "Accept ordinary professional labels, abbreviations, reordered "
                        "columns, and amounts shown to normal displayed precision."
                    ),
                },
                evidence_scope=f"candidate {metric_names[metric]} comparison row or card only",
            )
        status_id = f"metric__{metric}__status"
        status_gate = bool(record) and bool(_normalize(evidence))
        criteria.append(Criterion(
            status_id,
            f"The {metric_names[metric]} row is correctly associated and reports a supportable status",
            status_gate,
            evidence[:1600],
            semantic=True,
        ))
        relationship = (
            "posted and close-pro-forma values are below the maximum threshold"
            if metric == "leverage"
            else "posted and close-pro-forma values are above the minimum threshold"
        )
        add_semantic_spec(
            status_id,
            expected_facts={
                "metric": metric_names[metric],
                **{field: value for field, (value, _) in fields.items()},
                "required_relationship": relationship,
                "required_status": "compliant",
            },
            hard_gate=status_gate,
            hard_gate_evidence=f"nonblank candidate metric row present={status_gate}",
            submitted_evidence=evidence,
            reference_context={
                "association_rule": (
                    "The submitted row or card must clearly associate its values and status "
                    f"with {metric_names[metric]}; reject an unlabeled number dump or a row "
                    "that labels the values as another metric."
                ),
                "equivalence_rule": (
                    "Accept compliant, pass, within threshold, no breach, or an equally "
                    "clear professional equivalent."
                ),
                "numeric_boundary": (
                    "Exact numeric accuracy is scored by separate hard-gated criteria; judge "
                    "the metric association and whether the displayed relationship supports "
                    "the submitted status."
                ),
                "consistency_rule": (
                    "The status must agree with the posted, close-pro-forma, and threshold "
                    "values displayed in this submitted row. Numeric accuracy itself is "
                    "scored separately."
                ),
            },
            evidence_scope=f"candidate {metric_names[metric]} comparison row or card only",
        )

    decision_evidence = "\n".join(non_table_blocks)
    all_comparison_evidence = "\n".join(
        [
            *(record["evidence"] for record in table_records),
            decision_evidence,
        ]
    )
    decision_specs = (
        (
            "status__proposed",
            "The WIP entry is identified as proposed",
            {"required_state": "proposed rather than approved or recorded"},
            "Accept proposed, draft, pending approval, or an unambiguous equivalent.",
            decision_evidence,
        ),
        (
            "status__unposted",
            "The WIP entry is identified as unposted",
            {"required_state": "not posted to the accounting records"},
            "Accept unposted, not yet booked, not recorded, or an unambiguous equivalent.",
            decision_evidence,
        ),
        (
            "status__compliant",
            "The slide reaches the correct overall covenant conclusion",
            {"required_conclusion": "all three covenants comply on both displayed bases"},
            "Accept pass, compliant, within threshold, no breach, or an equally clear equivalent.",
            all_comparison_evidence,
        ),
        (
            "status__circulation",
            "The slide states the correct working-deck circulation boundary",
            {"required_conclusion": "internal working review; not signed, submitted, or final"},
            "Accept draft, working, internal review, not submitted, or equivalent language that preserves the release boundary.",
            decision_evidence,
        ),
    )
    for criterion_id, description, expected_facts, rule, evidence in decision_specs:
        gate = bool(added) and bool(_normalize(evidence))
        criteria.append(Criterion(
            criterion_id,
            description,
            gate,
            evidence[:1800],
            semantic=True,
        ))
        add_semantic_spec(
            criterion_id,
            expected_facts=expected_facts,
            hard_gate=gate,
            hard_gate_evidence=f"appended slide and nonblank decision evidence present={gate}",
            submitted_evidence=evidence,
            reference_context={"equivalence_rule": rule},
            evidence_scope="appended-slide decision and status text only",
        )

    binding_conclusion_blocks = [
        block
        for block in non_table_blocks
        if any(
            token in _normalize(block)
            for token in (
                "binding", "limiting", "most restrictive", "least capacity",
                "lowest capacity", "deterioration capacity",
            )
        )
    ]
    binding_evidence = "\n".join(
        [
            *(str(record["body"]) for record in table_records),
            *binding_conclusion_blocks,
        ]
    )
    # A bare ``FCCR binds`` assertion is not a like-for-like capacity
    # comparison.  Keep the conclusion semantic, but require the four
    # objective deterioration-capacity inputs that make the conclusion
    # reviewable on both posted and close-pro-forma bases.  This also prevents
    # a broader binding verdict from passing while its numeric prerequisites
    # are absent or wrong.
    binding_capacity_values = (
        capacity["fccr"]["posted"],
        capacity["leverage"]["posted"],
        capacity["fccr"]["pro_forma"],
        capacity["leverage"]["pro_forma"],
    )
    binding_capacity_gate = all(
        contains_scoped_number(
            all_slide_evidence,
            value,
            abs_tol=1.0,
            kind="money",
        )
        for value in binding_capacity_values
    )
    binding_gate = (
        bool(added)
        and bool(_normalize(binding_evidence))
        and binding_capacity_gate
    )
    criteria.append(Criterion(
        "headroom__binding_covenant",
        "The slide identifies FCCR as binding on like-for-like dollar deterioration capacity",
        binding_gate,
        binding_evidence[:2200],
        semantic=True,
    ))
    add_semantic_spec(
        "headroom__binding_covenant",
        expected_facts={
            "binding_covenant": "FCCR",
            "comparison_basis": "lowest like-for-like adjusted-EBITDA deterioration capacity",
            "posted_fccr_capacity": capacity["fccr"]["posted"],
            "posted_leverage_capacity": capacity["leverage"]["posted"],
            "pro_forma_fccr_capacity": capacity["fccr"]["pro_forma"],
            "pro_forma_leverage_capacity": capacity["leverage"]["pro_forma"],
        },
        hard_gate=binding_gate,
        hard_gate_evidence=(
            "nonblank binding conclusion and all four posted/pro-forma FCCR/"
            f"leverage capacity values present={binding_gate}"
        ),
        submitted_evidence=binding_evidence,
        reference_context={
            "required_conclusion": (
                "FCCR is binding because its adjusted-EBITDA deterioration capacity is "
                "lower than leverage capacity on both bases. TNW headroom is a different "
                "balance-sheet cushion and is not compared as EBITDA capacity."
            ),
            "equivalence_rule": (
                "Accept FCCR, fixed-charge coverage, limiting covenant, most restrictive, "
                "least like-for-like capacity, or equivalent professional wording."
            ),
            "consistency_rule": (
                "A bare binding assertion is insufficient. The slide must show or clearly "
                "derive both FCCR and leverage deterioration capacity on the posted and "
                "close-pro-forma bases so the like-for-like conclusion is reviewable."
            ),
        },
        evidence_scope="appended-slide headroom comparison and binding conclusion only",
    )

    source_gate = bool(added) and bool(_normalize(decision_evidence))
    criteria.append(Criterion(
        "source__authority",
        "The slide identifies the controlling legal, debt, accounting, cutoff, and WIP authorities",
        source_gate,
        decision_evidence[:2200],
        semantic=True,
    ))
    add_semantic_spec(
        "source__authority",
        expected_facts={
            "legal_authority": "executed U.S. Bank amendment",
            "debt_authority": "June 30 debt schedule",
            "posted_authority": "June 30 posted accounting records",
            "cutoff_authority": "ordinary June AP cutoff and invoice support",
            "wip_authority": "current PM, commercial, policy, and proposed-WIP support",
        },
        hard_gate=source_gate,
        hard_gate_evidence=f"nonblank appended-slide note evidence present={source_gate}",
        submitted_evidence=decision_evidence,
        reference_context={
            "equivalence_rule": (
                "Accept shortened filenames, business descriptions, acronyms, or a combined "
                "source line when each controlling authority remains unambiguous. Do not "
                "require authored filenames or exact phrases."
            ),
        },
        evidence_scope="appended-slide source and calculation-basis note only",
    )

    calculation_values = {
        "funded_debt": ("funded debt", gold["funded_debt"], 1.0),
        "posted_adjusted_ebitda": (
            "posted adjusted EBITDA",
            gold["ltm_adjusted_ebitda_posted"],
            1.0,
        ),
        "pending_ap_cutoff_expense": (
            "pending June AP cutoff expense",
            gold["pending_ap_cutoff_expense"],
            1.0,
        ),
        "proposed_wip_adjustment": (
            "proposed WIP adjustment",
            gold["proposed_wip_adjustment"],
            1.0,
        ),
        "pro_forma_adjusted_ebitda": (
            "close-pro-forma adjusted EBITDA",
            gold["ltm_adjusted_ebitda_pro_forma"],
            1.0,
        ),
        "cash_taxes": ("cash taxes", gold["cash_taxes"], .0001),
        "ltm_fixed_asset_additions": (
            "LTM fixed-asset additions",
            gold["ltm_fixed_asset_additions"],
            1.0,
        ),
        "direct_equipment_financing": (
            "direct LTM equipment financing",
            gold["ltm_direct_equipment_financing"],
            1.0,
        ),
        "unfunded_capex": ("unfunded capex", gold["unfunded_capex"], 1.0),
        "cash_interest": ("cash interest", gold["cash_interest"], 1.0),
        "scheduled_principal": (
            "scheduled principal",
            gold["scheduled_principal"],
            1.0,
        ),
        "fixed_charges": ("fixed charges", gold["fixed_charges"], 1.0),
        "posted_fccr_numerator": (
            "posted FCCR numerator",
            gold["fccr_numerator_posted"],
            1.0,
        ),
        "pro_forma_fccr_numerator": (
            "close-pro-forma FCCR numerator",
            gold["fccr_numerator_pro_forma"],
            1.0,
        ),
    }

    def matching_number_evidence(value: float, tolerance: float) -> str:
        matches = [
            record["evidence"]
            for record in evidence_records
            if contains_scoped_number(
                str(record["evidence"]),
                value,
                abs_tol=tolerance,
                kind="money",
            )
        ]
        return "\n".join(dict.fromkeys(str(match) for match in matches))

    calculation_value_met: dict[str, bool] = {}
    for field, (label, expected, tolerance) in calculation_values.items():
        evidence = matching_number_evidence(expected, tolerance)
        met = bool(evidence)
        calculation_value_met[field] = met
        criterion_id = f"calculation_support__{field}"
        criteria.append(Criterion(
            criterion_id,
            f"The appended slide contains the correct {label} value in its calculation basis",
            met,
            f"submitted={evidence[:1400]!r}; expected={expected}; tolerance={tolerance}",
            semantic=True,
        ))
        add_semantic_spec(
            criterion_id,
            expected_facts={
                "calculation_component": label,
                "required_value": expected,
            },
            hard_gate=met,
            hard_gate_evidence=(
                "the exact objective value is present in candidate calculation "
                f"evidence={met}"
            ),
            submitted_evidence=evidence,
            reference_context={
                "association_rule": (
                    "The exact number must be visibly associated with this calculation "
                    "component on the appended slide. Reject a number that appears only for "
                    "another metric, period, scenario, or unlabeled dump."
                ),
                "equivalence_rule": (
                    "Accept ordinary finance abbreviations, compact equations, tables, cards, "
                    "and normal displayed precision."
                ),
            },
            evidence_scope=f"appended-slide calculation basis for {label} only",
        )

    equation_specs = {
        "ebitda_bridge": {
            "fields": (
                "posted_adjusted_ebitda", "pending_ap_cutoff_expense",
                "proposed_wip_adjustment", "pro_forma_adjusted_ebitda",
            ),
            "expected": (
                "posted adjusted EBITDA less the pending June AP cutoff expense plus "
                "the proposed June WIP adjustment equals close-pro-forma adjusted EBITDA"
            ),
        },
        "unfunded_capex": {
            "fields": (
                "ltm_fixed_asset_additions", "direct_equipment_financing",
                "unfunded_capex",
            ),
            "expected": (
                "LTM fixed-asset additions less direct equipment financing equals unfunded capex"
            ),
        },
        "fixed_charges": {
            "fields": ("cash_interest", "scheduled_principal", "fixed_charges"),
            "expected": "cash interest plus scheduled principal equals fixed charges",
        },
        "posted_fccr_numerator": {
            "fields": (
                "posted_adjusted_ebitda", "cash_taxes", "unfunded_capex",
                "posted_fccr_numerator",
            ),
            "expected": (
                "posted adjusted EBITDA less cash taxes and unfunded capex equals the posted FCCR numerator"
            ),
        },
        "pro_forma_fccr_numerator": {
            "fields": (
                "pro_forma_adjusted_ebitda", "cash_taxes", "unfunded_capex",
                "pro_forma_fccr_numerator",
            ),
            "expected": (
                "close-pro-forma adjusted EBITDA less cash taxes and unfunded capex equals the close-pro-forma FCCR numerator"
            ),
        },
        "tnw_bridge": {
            "fields": ("pending_ap_cutoff_expense", "proposed_wip_adjustment"),
            "expected": (
                f"posted tangible net worth of ${gold['tangible_net_worth_posted']:,.2f} "
                f"less the pending June AP cutoff expense plus proposed WIP equals "
                f"close-pro-forma tangible net worth of ${gold['tangible_net_worth_pro_forma']:,.2f}"
            ),
        },
    }
    for equation, spec in equation_specs.items():
        gate = all(calculation_value_met[field] for field in spec["fields"])
        if equation == "tnw_bridge":
            gate = (
                gate
                and contains_scoped_number(
                    all_slide_evidence,
                    gold["tangible_net_worth_posted"],
                    abs_tol=1.0,
                    kind="money",
                )
                and contains_scoped_number(
                    all_slide_evidence,
                    gold["tangible_net_worth_pro_forma"],
                    abs_tol=1.0,
                    kind="money",
                )
            )
        criterion_id = f"calculation_support__equation__{equation}"
        criteria.append(Criterion(
            criterion_id,
            f"The slide presents the {equation.replace('_', ' ')} calculation and association",
            gate,
            all_slide_evidence[:2400],
            semantic=True,
        ))
        add_semantic_spec(
            criterion_id,
            expected_facts={"required_equation": spec["expected"]},
            hard_gate=gate,
            hard_gate_evidence=(
                "all exact objective values required by this bridge are present="
                f"{gate}"
            ),
            submitted_evidence=all_slide_evidence,
            reference_context={
                "association_rule": (
                    "The slide must visibly associate the components with this calculation "
                    "and direction. A scattered or unlabeled number dump is insufficient."
                ),
                "equivalence_rule": (
                    "Accept a formula, compact waterfall, calculation table, or clear prose "
                    "equation using ordinary labels and units."
                ),
            },
            evidence_scope=f"appended-slide {equation.replace('_', ' ')} calculation only",
        )

    scenario_records: dict[int, dict[str, Any] | None] = {}
    for scenario in gold["wip_reversal_sensitivity"]:
        percent = int(round(float(scenario["reversal_percent"]) * 100))
        best: tuple[int, dict[str, Any]] | None = None
        for record in evidence_records:
            body = str(record["body"])
            normalized = _normalize(body)
            marker = contains_scoped_number(
                body,
                float(scenario["reversal_percent"]),
                abs_tol=.0001,
                kind="percent",
            ) or (percent == 100 and "full reversal" in normalized)
            values = (
                (scenario["retained_wip_adjustment"], 1.0, "money"),
                (scenario["adjusted_ebitda"], 1.0, "money"),
                (scenario["tangible_net_worth"], 1.0, "money"),
                (scenario["leverage"], .0001, "ratio"),
                (scenario["fccr"], .0001, "ratio"),
            )
            numeric_hits = sum(
                contains_scoped_number(
                    body,
                    value,
                    abs_tol=tolerance,
                    kind=kind,
                )
                for value, tolerance, kind in values
            )
            layout_scope = str(record.get("layout_scope") or "")
            association_bonus = (
                5_000
                if layout_scope in {"native_row", "aligned_column"}
                else 0
            )
            score = association_bonus + (1000 if marker else 0) + numeric_hits * 100
            if score and (best is None or score > best[0]):
                best = (score, record)
        scenario_records[percent] = best[1] if best else None

    sensitivity_gate = all(scenario_records.values())
    sensitivity_evidence = "\n".join(
        str(record["evidence"])
        for record in scenario_records.values()
        if record is not None
    )
    criteria.append(Criterion(
        "structure__wip_reversal_sensitivity",
        "The slide contains the requested 25%, 50%, and full proposed-WIP reversal review",
        sensitivity_gate,
        sensitivity_evidence[:2600],
        semantic=True,
    ))
    add_semantic_spec(
        "structure__wip_reversal_sensitivity",
        expected_facts={
            "required_cases": ["25% reversal", "50% reversal", "full reversal"],
        },
        hard_gate=sensitivity_gate,
        hard_gate_evidence=f"candidate evidence exists for all three cases={sensitivity_gate}",
        submitted_evidence=sensitivity_evidence,
        reference_context={
            "equivalence_rule": (
                "Accept rows, cards, or a compact matrix in any order. The cases must remain "
                "distinguishable and must not be presented as posted results."
            ),
        },
        evidence_scope="appended-slide proposed-WIP reversal review only",
    )

    sensitivity_fields = {
        "retained_wip_adjustment": ("retained WIP adjustment", 1.0, "money"),
        "adjusted_ebitda": ("adjusted EBITDA", 1.0, "money"),
        "tangible_net_worth": ("tangible net worth", 1.0, "money"),
        "leverage": ("leverage", .0001, "ratio"),
        "fccr": ("fixed-charge coverage", .0001, "ratio"),
    }
    for scenario in gold["wip_reversal_sensitivity"]:
        percent = int(round(float(scenario["reversal_percent"]) * 100))
        record = scenario_records[percent]
        evidence = str(record["evidence"]) if record else "required scenario evidence is missing"
        for field, (label, tolerance, number_kind) in sensitivity_fields.items():
            target = float(scenario[field])
            met = bool(record) and contains_scoped_number(
                evidence,
                target,
                abs_tol=tolerance,
                kind=number_kind,
            )
            criterion_id = f"sensitivity__reversal_{percent}__{field}"
            criteria.append(Criterion(
                criterion_id,
                f"The {percent}% reversal evidence contains the correct {label} value",
                met,
                f"submitted={evidence[:1400]!r}; expected={target}; tolerance={tolerance}",
                semantic=True,
            ))
            add_semantic_spec(
                criterion_id,
                expected_facts={
                    "scenario": f"{percent}% reversal of proposed WIP",
                    "metric": label,
                    "required_value": target,
                },
                hard_gate=met,
                hard_gate_evidence=(
                    "the exact objective value is present in the candidate scenario "
                    f"evidence={met}"
                ),
                submitted_evidence=evidence,
                reference_context={
                    "association_rule": (
                        "The exact number must be visibly associated with both this reversal "
                        "case and this metric. Reject a swapped column, another case's value, "
                        "or an unlabeled number dump."
                    ),
                    "equivalence_rule": (
                        "Accept full/100% wording, arbitrary row or column order, ordinary "
                        "finance abbreviations, and normal displayed precision."
                    ),
                },
                evidence_scope=f"{percent}% proposed-WIP reversal row or card only",
            )
        status_id = f"sensitivity__reversal_{percent}__status"
        status_gate = bool(record) and bool(_normalize(evidence))
        criteria.append(Criterion(
            status_id,
            f"The {percent}% reversal row is correctly associated and reports a supportable covenant status",
            status_gate,
            evidence[:1800],
            semantic=True,
        ))
        add_semantic_spec(
            status_id,
            expected_facts={
                "scenario": f"{percent}% reversal of proposed WIP",
                "retained_wip_adjustment": scenario["retained_wip_adjustment"],
                "adjusted_ebitda": scenario["adjusted_ebitda"],
                "leverage": scenario["leverage"],
                "fccr": scenario["fccr"],
                "tangible_net_worth": scenario["tangible_net_worth"],
                "required_status": "all three covenants remain compliant",
            },
            hard_gate=status_gate,
            hard_gate_evidence=f"nonblank candidate scenario evidence present={status_gate}",
            submitted_evidence=evidence,
            reference_context={
                "association_rule": (
                    "The submitted evidence must associate the values with the correct reversal "
                    "case and metric labels. Reject swapped columns, another scenario's values, "
                    "or an unlabeled number dump."
                ),
                "equivalence_rule": (
                    "Accept compliant, pass, within threshold, no breach, or an equally clear "
                    "professional conclusion."
                ),
                "calculation_rule": (
                    "Retained WIP equals the unreversed portion. Adjusted EBITDA and TNW retain "
                    "the pending June AP expense, leverage uses funded debt over adjusted EBITDA, "
                    "and FCCR uses the displayed numerator over fixed charges."
                ),
            },
            evidence_scope=f"{percent}% proposed-WIP reversal row or card only",
        )

    result = _result(criteria)
    return _attach_semantic_review(
        result,
        task_id="task_015",
        evidence=_legacy_artifact_evidence(path),
        artifact_type="lender-review presentation",
        specs=semantic_specs,
        always_judge=True,
        execution_mode="scoped_per_criterion",
    )


def _document_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


TASK_001_ARTIFACT = Path(
    "Shared/Finance/Close/2026/06 June/4 WIP/"
    "ARM-2409 June WIP controller sign-off - WORKING.docx"
)


def _docx_table_rows_by_roles(
    document: Document,
    header_aliases: dict[str, tuple[str, ...]],
    *,
    expected_row_count: int | None = None,
    allow_positional_fallback: bool = False,
) -> list[dict[str, str]]:
    """Return a table using canonical column roles rather than literal headers.

    Word workpapers are often relabeled during review (for example, ``Amount``
    instead of ``Result`` or ``Responsible party`` instead of ``Owner``).  Such
    presentation changes are harmless.  The parser therefore accepts a bounded
    set of professional header equivalents while still requiring one and only
    one physical column for every requested role.
    """

    normalized_aliases = {
        role: {_normalize(alias) for alias in aliases}
        for role, aliases in header_aliases.items()
    }
    positional_candidates: list[list[dict[str, str]]] = []
    for table in document.tables:
        if not table.rows:
            continue
        physical_headers = [_normalize(cell.text) for cell in table.rows[0].cells]
        role_indexes: dict[str, int] = {}
        ambiguous = False
        for role, aliases in normalized_aliases.items():
            matches = [
                index for index, header in enumerate(physical_headers)
                if header in aliases
            ]
            if len(matches) != 1:
                ambiguous = True
                break
            role_indexes[role] = matches[0]
        if ambiguous or len(set(role_indexes.values())) != len(role_indexes):
            if (
                allow_positional_fallback
                and len(table.rows[0].cells) == len(header_aliases)
                and (
                    expected_row_count is None
                    or len(table.rows) - 1 == expected_row_count
                )
            ):
                roles = list(header_aliases)
                positional_candidates.append([
                    {
                        role: row.cells[index].text.strip()
                        for index, role in enumerate(roles)
                    }
                    for row in table.rows[1:]
                ])
            continue
        rows = [
            {
                role: row.cells[index].text.strip()
                for role, index in role_indexes.items()
            }
            for row in table.rows[1:]
        ]
        if expected_row_count is None or len(rows) >= expected_row_count:
            return rows
    return positional_candidates[0] if len(positional_candidates) == 1 else []


def _docx_named_row_by_aliases(
    rows: list[dict[str, str]],
    label_role: str,
    aliases: tuple[str, ...],
    *,
    fallback_index: int | None = None,
) -> dict[str, str]:
    wanted = {_normalize(alias) for alias in aliases}
    matches = [row for row in rows if _normalize(row.get(label_role)) in wanted]
    if len(matches) == 1:
        return matches[0]
    if not matches and fallback_index is not None and 0 <= fallback_index < len(rows):
        return rows[fallback_index]
    return {}


def _usable_docx_value(value: Any) -> bool:
    text = str(value or "").strip()
    return bool(text) and not bool(
        re.search(
            r"\[(?:enter|select|complete|cite|record|explain|state|list)\b",
            text,
            flags=re.I,
        )
    )


def _task_001_has_cutoff_date(text: Any) -> bool:
    """Recognize ordinary business renderings of the June 30, 2026 cutoff."""

    raw = str(text or "")
    normalized = _normalize(raw)
    if any(
        value in normalized
        for value in (
            "june 30 2026", "jun 30 2026", "30 june 2026", "30 jun 2026",
        )
    ):
        return True
    return bool(
        re.search(
            r"(?<!\d)(?:0?6\s*[/.-]\s*30\s*[/.-]\s*(?:20)?26|"
            r"2026\s*[/.-]\s*0?6\s*[/.-]\s*30)(?!\d)",
            raw,
            flags=re.I,
        )
        or re.search(r"(?<!\d)0?6\s+30\s+(?:20)?26(?!\d)", normalized)
    )


def _document_has_prohibited_signature(document: Document) -> bool:
    """Detect an actual signature block without flagging cited signed sources."""

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        marker = _normalize(paragraph.text)
        if re.match(
            r"^(?:signed by|signature|electronic signature|digitally signed by)\b",
            marker,
        ):
            return True
        if re.match(r"^s [a-z][a-z ]{1,80}$", marker):
            return True
    for part in document.part.package.parts:
        part_name = str(part.partname).casefold()
        content_type = str(part.content_type).casefold()
        if "_xmlsignatures" in part_name or "digital-signature" in content_type:
            return True
    for element in document.element.body.iter():
        metadata = " ".join(
            str(element.get(key, ""))
            for key in ("name", "descr", "title")
        ).casefold()
        if "signature" in metadata:
            return True
    return False


def _display_rounding_tolerance(token: str, *, floor: float) -> float:
    """Accept a number when it rounds to the professional display precision used."""
    clean = token.strip().strip("()")
    suffix_match = re.search(r"([kmb])\s*$", clean, flags=re.I)
    multiplier = {"k": 1_000.0, "m": 1_000_000.0, "b": 1_000_000_000.0}.get(
        suffix_match.group(1).lower() if suffix_match else "",
        1.0,
    )
    if suffix_match:
        clean = clean[: suffix_match.start()]
    percent = clean.rstrip().endswith("%")
    clean = clean.rstrip("%xX× ")
    decimal_match = re.search(r"\.([0-9]+)$", clean.replace(",", ""))
    decimals = len(decimal_match.group(1)) if decimal_match else 0
    tolerance = .5 * multiplier * (10 ** -decimals)
    if percent:
        tolerance /= 100
    return max(floor, tolerance + 1e-9)


def _text_contains_number(text: str, target: float, *, abs_tol: float = .02) -> bool:
    tokens = re.findall(
        r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?(?:[kmb]|x|%)?\)?",
        text,
        flags=re.I,
    )
    return any(
        _close(
            token,
            target,
            abs_tol=_display_rounding_tolerance(token, floor=abs_tol),
            rel_tol=2e-6,
        )
        for token in tokens
    )


def _task_001_v24_row_text(row: dict[str, str]) -> str:
    return " | ".join(str(value or "").strip() for value in row.values())


def _task_001_v24_bps_matches(value: Any, target: float) -> bool:
    text = str(value or "").translate(str.maketrans({"−": "-", "–": "-", "—": "-"}))
    tokens = re.findall(
        r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?(?:[kmb]|x|%)?\)?",
        text,
        flags=re.I,
    )
    normalized = _normalize(text)
    for token in tokens:
        value_number = _number(token)
        if value_number is None:
            continue
        if "%" in token:
            candidate = value_number * 10_000
        elif re.search(
            r"\b(?:pp|percentage points?|percent points?|points?|pts?)\b",
            normalized,
        ):
            candidate = value_number * 100
        else:
            candidate = value_number
        # A controller memo commonly presents a basis-point movement as a whole
        # number (or derives it from endpoint percentages displayed to two
        # decimals).  Preserve the tighter tolerance for a precise submission,
        # while allowing the less-than-one-basis-point presentation difference
        # implied by a whole-bp value.  This accepts 19/20 bps for 19.70 and
        # (335) bps for -334.78 without accepting a materially different move.
        numeric_token = re.sub(r"[^0-9.]", "", token)
        decimal_places = (
            len(numeric_token.rsplit(".", 1)[1]) if "." in numeric_token else 0
        )
        is_percentage_point_value = bool(
            re.search(
                r"\b(?:pp|percentage points?|percent points?|points?|pts?)\b",
                normalized,
            )
        )
        displayed_resolution_bps = 1.0
        if "%" in token or is_percentage_point_value:
            displayed_resolution_bps = 100.0 / (10 ** decimal_places)
        elif decimal_places:
            displayed_resolution_bps = 1.0 / (10 ** decimal_places)
        tolerance = max(0.15, min(displayed_resolution_bps, 1.0))
        if abs(candidate - target) <= tolerance + 1e-9:
            return True
    return False


def _task_001_v24_zero(value: Any, *, allow_blank: bool = False) -> bool:
    text = str(value or "").strip()
    if not text:
        return allow_blank
    if text in {"-", "--", "–", "—", "−"}:
        return True
    normalized = _normalize(text)
    if re.search(
        r"\b(?:does\s+not|doesn\s*t|did\s+not|not)\s+"
        r"(?:tie|reconcile|balance|agree)\b|\b(?:unreconciled|out\s+of\s+balance)\b",
        normalized,
    ):
        return False
    if normalized in {
        "zero", "nil", "none", "n a", "not applicable", "tied", "ties",
        "reconciled", "reconciles", "balanced", "no variance", "zero variance",
        "pass", "passes", "ok", "clear", "no exceptions",
    }:
        return True
    # Accept an explicit zero-variance or exact-tie conclusion embedded in a
    # normal control narrative.  Do not require the entire cell to equal a
    # magic phrase, and do not let unrelated non-zero source amounts make an
    # otherwise explicit zero variance fail.
    if re.search(r"\b(?:no|zero)\s+(?:unresolved\s+)?variance\b", normalized):
        return True
    if re.match(r"^(?:tie|tied|reconciled|balanced)\b", normalized):
        return True
    if re.search(
        r"\b(?:ties?|reconciles?|balances?|agrees?)\s+exactly\b",
        normalized,
    ):
        return True
    variance_match = re.search(
        r"\bvariance\b\s*(?::|=|is|of)?\s*"
        r"(\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?%?\)?)",
        text,
        flags=re.I,
    )
    if variance_match:
        return _close(variance_match.group(1), 0.0, abs_tol=.02, rel_tol=0.0)
    leading_variance_match = re.search(
        r"(\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?%?\)?)\s+variance\b",
        text,
        flags=re.I,
    )
    if leading_variance_match:
        return _close(
            leading_variance_match.group(1), 0.0, abs_tol=.02, rel_tol=0.0
        )
    tokens = re.findall(r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?%?\)?", text)
    return bool(tokens) and all(_close(token, 0.0, abs_tol=.02, rel_tol=0.0) for token in tokens)


def _task_001_v24_contradicts_zero(value: Any) -> bool:
    """Return true only for an explicit non-zero or failed-tie statement."""

    text = str(value or "").strip()
    normalized = _normalize(text)
    if re.search(
        r"\b(?:does\s+not|doesn\s*t|did\s+not|not)\s+"
        r"(?:tie|reconcile|balance|agree)\b|\b(?:unreconciled|out\s+of\s+balance)\b",
        normalized,
    ):
        return True
    match_after = re.search(
        r"\bvariance\b\s*(?::|=|is|of)?\s*"
        r"(\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?%?\)?)",
        text,
        flags=re.I,
    )
    match_before = re.search(
        r"(\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?%?\)?)\s+variance\b",
        text,
        flags=re.I,
    )
    match = match_after or match_before
    return bool(match) and not _close(match.group(1), 0.0, abs_tol=.02, rel_tol=0.0)


def _task_001_v24_box(document: Document | None, *labels: str) -> str:
    if document is None:
        return ""
    wanted = {_normalize(label) for label in labels}
    candidates: list[tuple[str, str]] = []
    for table in document.tables:
        if len(table.rows) < 2 or len(table.rows[0].cells) != 1:
            continue
        heading = _normalize(table.rows[0].cells[0].text)
        body = "\n".join(
            row.cells[0].text.strip() for row in table.rows[1:]
        ).strip()
        candidates.append((heading, body))
    # Prefer the exact business heading.  A generic alias such as "posting
    # boundary" can also appear in the earlier close-recommendation heading;
    # returning the first substring match silently sends the wrong evidence to
    # the semantic judge.
    for heading, body in candidates:
        if heading in wanted:
            return body
    fuzzy: list[tuple[int, str]] = []
    for heading, body in candidates:
        for label in wanted:
            if not label or not heading:
                continue
            if label in heading or heading in label:
                fuzzy.append((min(len(label), len(heading)), body))
    if fuzzy:
        return max(fuzzy, key=lambda item: item[0])[1]
    return ""


def _task_001_v24_artifact_review(
    document: Document | None, path: Path, parse_error: str
) -> dict[str, Any]:
    tables: list[dict[str, Any]] = []
    paragraphs: list[str] = []
    if document is not None:
        paragraphs = [
            paragraph.text.strip() for paragraph in document.paragraphs
            if paragraph.text.strip()
        ][:80]
        for index, table in enumerate(document.tables[:30], start=1):
            rows = [
                [cell.text.strip()[:500] for cell in row.cells]
                for row in table.rows[:80]
            ]
            tables.append({"table": index, "rows": rows})
    return {
        "version": 1,
        "artifact": str(TASK_001_ARTIFACT),
        "exists": path.is_file(),
        "parse_error": parse_error,
        "paragraphs": paragraphs,
        "tables": tables,
    }


TASK_001_V27_TABLE_ALIASES: dict[str, dict[str, tuple[str, ...]]] = {
    "close": {
        "metric": ("Metric", "Measure", "Line item"),
        "source": (
            "Source / submitted", "Source or submitted", "Current / submitted",
            "Submitted", "Source value", "PM / source",
        ),
        "finance": (
            "Finance close", "Recommended close", "June close", "Adjusted close",
            "Controller-review close",
        ),
        "difference": (
            "Difference", "Change", "Variance", "Adjustment", "Change / control",
        ),
        "basis": ("Basis", "Source / basis", "Source or conclusion", "Explanation"),
    },
    "bridge": {
        "metric": ("Metric", "Measure", "Line item"),
        "may": ("May final", "Final May", "May close", "Prior close"),
        "june": ("June close", "Finance close", "Recommended close", "Current close"),
        "change": ("Change", "June less May", "Movement", "Difference", "Variance"),
        "explanation": ("Explanation", "Driver", "Comment", "Movement explanation"),
    },
    "reconciliation": {
        "control": ("Control", "Control line", "Metric", "Line item"),
        "may": ("May final", "May / opening", "Opening", "Prior close"),
        "activity": (
            "June activity", "Current-period activity", "Activity / change",
            "June activity / approved change",
        ),
        "june": ("June close", "Closing", "Ending balance", "Close"),
        "variance": ("Variance / status", "Variance", "Tie / status", "Control status"),
    },
    "population": {
        "population": ("Population", "Source / population", "Source population", "Source"),
        "cutoff": ("Cutoff", "As of", "Period / cutoff", "Through"),
        "records": ("Records", "Source records", "Record count", "Count", "Entries"),
        "source_total": ("Source total", "Source amount", "Population amount"),
        "ledger_total": ("Ledger total", "Ledger amount", "Posted amount"),
        "variance": ("Variance", "Difference", "Control variance"),
        "conclusion": (
            "Conclusion / follow-up", "Conclusion and follow-up", "Exceptions / resolution",
            "Control conclusion", "Status / resolution",
        ),
    },
    "adjustment": {
        "assumption": (
            "Assumption / component", "Assumption or component", "Component",
            "Forecast item", "Line item",
        ),
        "pm": ("PM forecast", "Submitted PM ETC", "PM ETC", "Submitted forecast"),
        "adjustment": ("Finance adjustment", "Close adjustment", "Finance correction"),
        "close": ("Finance close", "Recommended close ETC", "Adjusted ETC", "Close ETC"),
        "basis": ("Evidence / rationale", "Evidence and rationale", "Basis", "Source / treatment"),
    },
    "judgment": {
        "issue": ("Issue", "Question", "Decision", "Assessment"),
        "fact": ("Fact at cutoff", "Cutoff fact", "Current evidence", "Fact"),
        "treatment": ("Finance treatment", "Accounting treatment", "Conclusion", "Decision"),
        "evidence": ("Evidence", "Controlling evidence / policy", "Source / policy"),
        "clearance": ("Owner / clearance", "Owner and clearance", "Clearance", "Next step"),
    },
    "source": {
        "source": ("Source", "Source role", "Evidence source", "Document / system"),
        "version": ("Version / cutoff", "Version and cutoff", "As of", "Cutoff"),
        "fact": ("Fact used", "Use / fact", "Fact used and conclusion", "Purpose"),
        "status": ("Status / limitation", "Status and limitation", "Authority / status", "Limitation"),
    },
    "posting": {
        "entry": ("Entry / control", "Entry or control", "Stage", "Posting step"),
        "account": ("Account / description", "Account and description", "Account", "GL account"),
        "debit": ("Debit", "Debits"),
        "credit": ("Credit", "Credits"),
        "basis": ("Status / basis", "Status and basis", "Basis / control", "Note"),
    },
    "action": {
        "owner": ("Owner", "Responsible", "Responsible party", "Function"),
        "action": ("Required action", "Action", "Next action"),
        "evidence": ("Completion evidence", "Evidence to close", "Close evidence"),
        "due": ("Due", "Timing", "Due date"),
        "status": ("Status", "Action status"),
    },
}


def _grade_task_001_v27(workspace_root: Path, answer: Any) -> dict[str, Any]:
    """Grade a realistic ARM-2409 close memo without turning its template into an answer key."""

    del answer
    path = workspace_root / TASK_001_ARTIFACT
    document: Document | None = None
    parse_error = ""
    if path.is_file():
        try:
            document = Document(path)
        except Exception as exc:  # pragma: no cover - corrupt submissions vary
            parse_error = f"{type(exc).__name__}: {exc}"
    else:
        parse_error = "artifact missing"

    gold = load_apex_gold("task_001")
    project = gold["project"]
    prior = gold["prior_month"]
    bridge_gold = gold["bridge"]
    journal_gold = gold["proposed_journal_bridge"]
    forecast_recovery = float(gold["forecast_recovery"])
    executed_credit = float(gold["executed_credit_agreement"])
    current_cost_credit = float(gold["documented_current_cost_credit"])
    remaining_cost_credit = float(gold["documented_remaining_cost_credit"])
    credit_effective_at_cutoff = bool(gold["credit_effective_at_cutoff"])
    recognized_current_cost_credit = float(gold["executed_current_cost_credit"])
    recognized_remaining_cost_credit = float(gold["executed_remaining_cost_credit"])
    if (
        credit_effective_at_cutoff
        or recognized_current_cost_credit
        or recognized_remaining_cost_credit
    ):
        raise ValueError("Task001 conditional credit unexpectedly recognized at cutoff")
    commercial_correction = float(gold["commercial_etc_correction"])
    cutoff_invoice_gross = float(gold["cutoff_invoice_gross_amount"])
    cutoff_invoice_face = float(gold["cutoff_invoice_face_amount"])
    cutoff_invoice_prior_billing = float(gold["cutoff_invoice_prior_billing"])
    cutoff_invoice_gross_june_work = float(
        gold["cutoff_invoice_gross_june_work"]
    )
    cutoff_accrual = float(gold["cutoff_accrual"])
    post_cutoff_invoice_line = float(gold["post_cutoff_invoice_line"])
    cutoff_credit_memo = float(gold["cutoff_credit_memo"])
    net_cutoff_cost_adjustment = float(gold["net_cutoff_cost_adjustment"])
    pco_cost_overlap = float(gold["pco_cost_already_in_pm_base"])
    prior_billing_duplicated_in_pm_etc = float(
        gold["prior_billing_duplicated_in_pm_etc"]
    )
    net_etc_correction = float(gold["required_etc_adjustment"])
    posted_cost = float(project["posted_cost_to_date"])
    close_cost_basis = float(project["cost_to_date"])
    close_etc = float(project["estimated_cost_to_complete"])
    contract_asset = float(project["underbilling"])
    estimated_margin = float(project["estimated_total_margin"])
    margin_change_bps = float(bridge_gold["margin_rate_change_bps"])
    prior_contract_asset = float(prior["contract_asset"])
    proposed_debits = float(
        journal_gold["proposed_june_contribution"]["total_debits"]
    )
    proposed_credits = float(
        journal_gold["proposed_june_contribution"]["total_credits"]
    )
    period_effect = float(journal_gold["period_effect"]["contract_asset_change"])
    period_direction = "increase" if period_effect >= 0 else "decrease"
    revenue_effect_direction = "credit" if period_effect >= 0 else "debit"
    text = _document_text(document) if document is not None else ""
    normalized = _normalize(text)
    canonical_tables = {
        name: (
            _docx_table_rows_by_roles(
                document,
                {
                    role: (aliases[0],)
                    for role, aliases in role_aliases.items()
                },
                allow_positional_fallback=False,
            )
            if document is not None
            else []
        )
        for name, role_aliases in TASK_001_V27_TABLE_ALIASES.items()
    }
    raw_tables: list[dict[str, Any]] = []
    if document is not None:
        for table_index, table in enumerate(document.tables):
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
            ]
            if not rows:
                continue
            raw_tables.append({
                "index": table_index,
                "columns": len(rows[0]),
                "rows": rows,
            })

    def raw_table_block(table: dict[str, Any], *, cell_limit: int | None = None) -> str:
        rendered_rows = []
        for cells in table["rows"]:
            rendered = [
                value if cell_limit is None else value[:cell_limit]
                for value in cells
            ]
            rendered_rows.append(" | ".join(rendered))
        return f"[TABLE {table['index']}]\n" + "\n".join(rendered_rows)

    def raw_table_scope(
        *,
        columns: set[int] | None = None,
        min_rows: int | None = None,
        max_rows: int | None = None,
        cell_limit: int | None = None,
    ) -> str:
        selected = []
        for table in raw_tables:
            row_count = len(table["rows"])
            if columns is not None and table["columns"] not in columns:
                continue
            if min_rows is not None and row_count < min_rows:
                continue
            if max_rows is not None and row_count > max_rows:
                continue
            selected.append(raw_table_block(table, cell_limit=cell_limit))
        return "\n\n".join(selected)

    def compact_document_evidence() -> str:
        """Expose every row label without allowing a label parser to pre-adjudicate it."""

        paragraphs = "\n".join(
            paragraph.text.strip()
            for paragraph in (document.paragraphs if document is not None else [])
            if paragraph.text.strip()
        )
        compact_tables = raw_table_scope(cell_limit=40)
        return "\n\n".join(value for value in (paragraphs, compact_tables) if value)

    compact_evidence = compact_document_evidence()
    five_column_evidence = raw_table_scope(columns={5})
    four_column_evidence = raw_table_scope(columns={4})
    population_table_evidence = raw_table_scope(columns={7})
    narrative_evidence = raw_table_scope(columns={1})
    action_like_evidence = raw_table_scope(columns={5}, max_rows=6)

    def row_matches_numbers(
        cells: list[str],
        requirements: tuple[tuple[float, int, float], ...],
    ) -> bool:
        return all(
            sum(
                _text_contains_number(cell, target, abs_tol=abs_tol)
                for cell in cells
            ) >= required_count
            for target, required_count, abs_tol in requirements
        )

    def numeric_candidate_evidence(
        *requirements: tuple[float, int, float],
        columns: set[int] | None = None,
        bps_target: float | None = None,
    ) -> str:
        """Return exact-number candidate rows; an LLM still decides row meaning."""

        candidates: list[str] = []
        for table in raw_tables:
            if columns is not None and table["columns"] not in columns:
                continue
            header = table["rows"][0]
            for row_index, cells in enumerate(table["rows"][1:], start=1):
                if requirements and not row_matches_numbers(cells, requirements):
                    continue
                if bps_target is not None and not any(
                    _task_001_v24_bps_matches(cell, bps_target)
                    for cell in cells
                ):
                    continue
                candidates.append(
                    f"[TABLE {table['index']} ROW {row_index}]\n"
                    f"HEADER: {' | '.join(header)}\n"
                    f"SUBMITTED: {' | '.join(cells)}"
                )
        return "\n\n".join(candidates)

    def anchored_row_evidence(*anchors: str) -> str:
        wanted = tuple(_normalize(anchor) for anchor in anchors if _normalize(anchor))
        candidates: list[str] = []
        for table in raw_tables:
            header = table["rows"][0]
            for row_index, cells in enumerate(table["rows"][1:], start=1):
                normalized_row = _normalize(" | ".join(cells))
                if wanted and not any(anchor in normalized_row for anchor in wanted):
                    continue
                candidates.append(
                    f"[TABLE {table['index']} ROW {row_index}]\n"
                    f"HEADER: {' | '.join(header)}\n"
                    f"SUBMITTED: {' | '.join(cells)}"
                )
        return "\n\n".join(candidates)

    def anchored_numeric_candidate_evidence(
        anchors: tuple[str, ...],
        *requirements: tuple[float, int, float],
    ) -> str:
        """Return rows containing a stable identifier and exact numeric prerequisites."""

        wanted = tuple(_normalize(anchor) for anchor in anchors if _normalize(anchor))
        candidates: list[str] = []
        for table in raw_tables:
            header = table["rows"][0]
            for row_index, cells in enumerate(table["rows"][1:], start=1):
                normalized_row = _normalize(" | ".join(cells))
                if wanted and not all(anchor in normalized_row for anchor in wanted):
                    continue
                if requirements and not row_matches_numbers(cells, requirements):
                    continue
                candidates.append(
                    f"[TABLE {table['index']} ROW {row_index}]\n"
                    f"HEADER: {' | '.join(header)}\n"
                    f"SUBMITTED: {' | '.join(cells)}"
                )
        return "\n\n".join(candidates)

    starter_path = SEED_WORKSPACE / TASK_001_ARTIFACT
    starter_text = ""
    if starter_path.is_file():
        try:
            starter_text = _document_text(Document(starter_path))
        except Exception:
            starter_text = ""
    document_authored = bool(
        document is not None
        and _normalize(text)
        and _normalize(text) != _normalize(starter_text)
    )
    criteria: list[Criterion] = []
    semantic_specs: list[dict[str, Any]] = []

    def add(
        criterion_id: str,
        description: str,
        met: bool,
        evidence: str,
        *,
        category: str = "core_finance",
        weight: int = 10,
        failure_cap: float | None = None,
    ) -> None:
        criteria.append(Criterion(
            criterion_id,
            description,
            bool(met),
            evidence[:2_000],
            category=category,
            weight=weight,
            semantic=False,
            failure_cap=failure_cap,
        ))

    def add_semantic(
        criterion_id: str,
        description: str,
        submitted_evidence: str,
        reference_context: Any,
        *,
        category: str = "core_finance",
        weight: int = 10,
        hard_gate: bool | None = None,
        evidence_scope: str = "criterion evidence",
        failure_cap: float | None = None,
        legacy_met: bool = False,
        always_judge: bool = True,
        hard_gate_evidence: str | None = None,
    ) -> None:
        submitted = str(submitted_evidence or "").strip()
        gate = bool(document_authored) and (
            bool(submitted) if hard_gate is None else bool(hard_gate)
        )
        criteria.append(Criterion(
            criterion_id,
            description,
            bool(legacy_met),
            (
                "canonical template association matched"
                if legacy_met
                else "submitted evidence is present" if gate
                else "submitted evidence is absent"
            ),
            category=category,
            weight=weight,
            semantic=True,
            failure_cap=failure_cap,
        ))
        semantic_specs.append({
            "criterion_id": criterion_id,
            "expected_facts": reference_context,
            "reference_context": reference_context,
            "task_context": {
                "project": "ARM-2409 Northline Cold Storage Expansion",
                "cutoff": "2026-06-30",
                "work_product": "preparer controller-sign-off memorandum",
                "accounting_basis": "cost-to-cost WIP under signed FIN-REV-04",
            },
            "submitted_evidence": submitted[:12_000],
            "evidence_scope": evidence_scope,
            "hard_gate_met": gate,
            "hard_gate_evidence": (
                hard_gate_evidence
                if hard_gate_evidence is not None
                else (
                    "the exact objective prerequisite is visibly present"
                    if gate else "the exact objective prerequisite is blank or missing"
                )
            ),
            "always_judge": bool(always_judge),
        })

    def row_for(
        table_name: str,
        label_role: str,
        aliases: tuple[str, ...],
    ) -> dict[str, str]:
        # Only the exact starter header and exact starter row label qualify for
        # the deterministic fast path. Any edited language is routed to the
        # criterion-scoped semantic judge instead of being accepted or rejected
        # by a hand-maintained synonym list.
        return _docx_named_row_by_aliases(
            canonical_tables[table_name],
            label_role,
            (aliases[0],),
        )

    def row_text(row: dict[str, str]) -> str:
        return _task_001_v24_row_text(row)

    def authored_row_text(
        row: dict[str, str],
        label_role: str,
    ) -> str:
        """Return a row only when the preparer authored a non-label field."""

        return (
            row_text(row)
            if any(
                role != label_role and _usable_docx_value(value)
                for role, value in row.items()
            )
            else ""
        )

    def add_number(
        criterion_id: str,
        description: str,
        row: dict[str, str],
        roles: tuple[str, ...],
        expected: float,
        *,
        abs_tol: float = .02,
        failure_cap: float | None = None,
    ) -> None:
        primary_text = " | ".join(str(row.get(role) or "") for role in roles)
        is_rate = criterion_id in {
            "close__percent_complete",
            "close__margin_rate",
            "bridge__margin_rate__may",
        }

        def rate_matches(value: str) -> bool:
            """Match Task001 rates without treating a bare count of 1 as 100%."""

            tokens = re.findall(
                r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?(?:[kmb]|x|%)?\)?",
                value,
                flags=re.I,
            )
            for token in tokens:
                parsed = _number(token)
                if parsed is None:
                    continue
                clean = token.strip().strip("()")
                if "$" in clean or re.search(r"[kmbx×]\s*$", clean, flags=re.I):
                    continue
                if clean.endswith("%"):
                    candidates = ((parsed, _display_rounding_tolerance(
                        token, floor=abs_tol
                    )),)
                elif abs(parsed) <= 1.5:
                    # Decimal-rate notation such as .8464 or .85.
                    candidates = ((
                        parsed,
                        min(
                            _display_rounding_tolerance(token, floor=abs_tol),
                            .005000001,
                        ),
                    ),)
                elif abs(parsed) <= 100:
                    # In a rate row, 84.64 is a conventional display of
                    # 84.64%. Keep its displayed precision after conversion.
                    candidates = ((
                        parsed / 100,
                        _display_rounding_tolerance(token, floor=abs_tol) / 100,
                    ),)
                else:
                    candidates = ()
                if any(
                    abs(candidate - expected)
                    <= max(tolerance, abs(expected) * 2e-6)
                    for candidate, tolerance in candidates
                ):
                    return True
            return False

        number_matches = rate_matches if is_rate else (
            lambda value: _text_contains_number(value, expected, abs_tol=abs_tol)
        )
        canonical_met = bool(row) and number_matches(primary_text)

        if is_rate:
            candidate_rows: list[str] = []
            for table in raw_tables:
                header = table["rows"][0]
                for row_index, cells in enumerate(table["rows"][1:], start=1):
                    if not any(rate_matches(cell) for cell in cells):
                        continue
                    candidate_rows.append(
                        f"[TABLE {table['index']} ROW {row_index}]\n"
                        f"HEADER: {' | '.join(header)}\n"
                        f"SUBMITTED: {' | '.join(cells)}"
                    )
            exact_candidate_evidence = "\n\n".join(candidate_rows)
        else:
            exact_candidate_evidence = numeric_candidate_evidence(
                (expected, 1, abs_tol)
            )

        primary_has_number = bool(re.search(
            r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?(?:[kmb]|x|%)?\)?",
            primary_text,
            flags=re.I,
        ))
        canonical_primary_conflict = bool(
            row and primary_has_number and not canonical_met
        )
        submitted = exact_candidate_evidence
        reference_context: dict[str, Any] = {
            "required_association": description,
            "expected_value": expected,
            "numeric_tolerance": abs_tol,
            "grading_boundary": (
                "The code gate checks only that the exact numeric target is visibly "
                "present in at least one submitted row. Decide semantically whether "
                "that value is associated with the named metric, period, and column. "
                "Accept ordinary labels, reordered rows and tables, common finance "
                "units, and equivalent sign presentation; reject a correct number "
                "attached to a different metric or scenario."
            ),
        }
        if criterion_id in {
            "close__percent_complete",
            "close__margin_rate",
        }:
            # Percentage targets can occur accidentally inside unrelated dollar
            # magnitudes. Give the semantic judge the surrounding submitted
            # five-column schedules so its explanation can acknowledge a visible
            # wrong primary value instead of incorrectly claiming the metric is
            # absent. The exact-target evidence alone still owns the hard gate.
            submitted = "\n\n".join(
                value for value in (
                    exact_candidate_evidence,
                    five_column_evidence,
                ) if value
            )
            reference_context["explanation_fidelity"] = (
                "When a visible primary close value is wrong, name or acknowledge "
                "that submitted value and compare it with the expected value. Do not "
                "claim the metric is absent, and do not describe an alternative or "
                "sensitivity value as the only submitted value, when the supplied "
                "schedule visibly contains a primary value."
            )
        if canonical_primary_conflict:
            submitted = "\n\n".join(
                value for value in (
                    f"UNCHANGED CANONICAL PRIMARY FIELD:\n{primary_text}",
                    submitted,
                ) if value
            )
        add_semantic(
            criterion_id,
            description,
            submitted,
            reference_context,
            hard_gate=bool(exact_candidate_evidence) and not canonical_primary_conflict,
            evidence_scope=(
                "unchanged canonical primary field plus all submitted rows containing "
                "the exact numeric target"
            ),
            failure_cap=failure_cap,
            legacy_met=canonical_met,
            always_judge=False,
            hard_gate_evidence=(
                "unchanged canonical primary field contains a conflicting numeric "
                f"value: {primary_text[:500]}"
                if canonical_primary_conflict
                else (
                    "the exact objective prerequisite is visibly present"
                    if exact_candidate_evidence
                    else "the exact objective prerequisite is blank or missing"
                )
            ),
        )

    def add_zero(
        criterion_id: str,
        description: str,
        row: dict[str, str],
        role: str,
        *,
        associated_amount: float | None = None,
        failure_cap: float | None = None,
    ) -> None:
        canonical_met = bool(row) and _task_001_v24_zero(row.get(role))
        requirements: list[tuple[float, int, float]] = [(0.0, 1, .000001)]
        if associated_amount is not None:
            requirements.insert(0, (associated_amount, 1, .02))
        submitted = numeric_candidate_evidence(*requirements)
        add_semantic(
            criterion_id,
            description,
            submitted,
            {
                "required_association": description,
                "expected_value": 0.0,
                "associated_amount": associated_amount,
                "grading_boundary": (
                    "The code gate verifies only that an explicit zero appears in a "
                    "candidate row, together with the associated amount when supplied. "
                    "Decide whether the zero is actually the named ledger or variance "
                    "field. Accept an unambiguous tied/no-difference status as equivalent "
                    "when the displayed source and ledger amounts establish the same fact."
                ),
            },
            category="auditability",
            weight=5,
            hard_gate=bool(submitted),
            evidence_scope="candidate rows containing the exact amount and zero control",
            failure_cap=failure_cap,
            legacy_met=canonical_met,
            always_judge=False,
        )

    def add_bps(
        criterion_id: str,
        description: str,
        row: dict[str, str],
        role: str,
        expected: float,
    ) -> None:
        canonical_met = bool(row) and _task_001_v24_bps_matches(
            row.get(role), expected
        )
        submitted = numeric_candidate_evidence(bps_target=expected)
        add_semantic(
            criterion_id,
            description,
            submitted,
            {
                "required_association": description,
                "expected_basis_points": expected,
                "acceptable_units": (
                    "basis points, percentage points, or an equivalent signed percent "
                    "movement that converts to the same basis-point change"
                ),
                "grading_boundary": (
                    "The code gate checks only that the exact rate movement is visibly "
                    "present in some submitted row. Decide whether it is associated with "
                    "the named metric and May-to-June bridge."
                ),
            },
            hard_gate=bool(submitted),
            evidence_scope="all submitted rows containing the exact rate movement",
            legacy_met=canonical_met,
            always_judge=False,
        )

    # Editable deliverable and preparer boundary.
    add(
        "artifact__working_document_present",
        "The required editable working memorandum exists",
        path.is_file(),
        f"path={path}",
        category="structure",
        weight=1,
    )
    add(
        "artifact__document_parses",
        "The working memorandum is a readable Word document",
        document is not None,
        parse_error or "document parsed",
        category="structure",
        weight=1,
    )
    add(
        "artifact__project_id",
        "The memorandum identifies ARM-2409",
        "arm 2409" in normalized,
        "project id present" if "arm 2409" in normalized else "project id absent",
        category="structure",
        weight=1,
    )
    add(
        "artifact__project_name",
        "The memorandum identifies Northline Cold Storage Expansion",
        "northline cold storage expansion" in normalized,
        "project name present" if "northline cold storage expansion" in normalized else "project name absent",
        category="structure",
        weight=1,
    )
    add(
        "artifact__june_30_cutoff",
        "The memorandum identifies the June 30, 2026 close cutoff",
        _task_001_has_cutoff_date(text),
        "cutoff found" if _task_001_has_cutoff_date(text) else "cutoff not found",
        category="structure",
        weight=1,
    )
    metadata_evidence = "\n\n".join(
        value for value in (four_column_evidence, compact_evidence[:4_000]) if value
    )
    add_semantic(
        "artifact__preparer",
        "A preparer is identified",
        metadata_evidence,
        {
            "required_field": "an identified preparer person, role, or function",
            "grading_boundary": (
                "Accept any normal professional metadata label. A project manager, "
                "reviewer, or approver does not count unless the document identifies "
                "that party as the preparer."
            ),
        },
        category="structure",
        weight=1,
        hard_gate=document_authored and bool(metadata_evidence),
        evidence_scope="metadata-like rows and compact document field labels",
    )
    add_semantic(
        "artifact__prepared_date",
        "A preparation date is recorded",
        metadata_evidence,
        {
            "required_field": "a date explicitly associated with preparation of this memorandum",
            "grading_boundary": (
                "Accept ordinary date formats and equivalent metadata labels. A source "
                "cutoff, invoice date, posting date, or reviewer date is not a prepared date."
            ),
        },
        category="structure",
        weight=1,
        hard_gate=document_authored and bool(metadata_evidence),
        evidence_scope="metadata-like rows and compact document field labels",
    )
    add_semantic(
        "artifact__review_status",
        "The review status clearly remains pending, draft, or otherwise not Controller-approved",
        metadata_evidence,
        {
            "acceptable_meaning": (
                "pending Controller review, preparer draft, not approved, returned for "
                "review, or another unambiguous not-yet-approved status"
            ),
            "unacceptable_meaning": "approved, posted, or authorized to post",
        },
        category="structure",
        weight=1,
        hard_gate=document_authored and bool(metadata_evidence),
        evidence_scope="metadata-like rows and compact document field labels",
    )
    unresolved_prompts = re.findall(
        r"\[(?:enter|select|complete|cite|record|explain|state|list)\b",
        text,
        flags=re.I,
    )
    add(
        "artifact__no_instruction_placeholders",
        "The completed memorandum contains no unresolved bracketed instructions",
        document is not None and not unresolved_prompts,
        f"unresolved_instruction_prompts={len(unresolved_prompts)}",
        category="structure",
        weight=1,
    )
    prohibited_signature = bool(
        document and _document_has_prohibited_signature(document)
    )
    add(
        "artifact__no_simulated_controller_signature",
        "The preparer did not simulate Controller approval or a Controller signature",
        document is not None and not prohibited_signature,
        f"prohibited_signature={prohibited_signature}",
        category="integrity",
        weight=10,
        failure_cap=0.0,
    )

    # June close calculation. Only objective values are deterministic.
    close_specs = (
        ("current_contract", ("Current contract", "Executed contract", "Contract value"), ("finance",), project["current_contract"]),
        ("close_cost_basis", ("Cost to date / close basis", "Close cost basis", "Cost to date", "Posted cost to date", "Posted cost"), ("finance",), close_cost_basis),
        ("billings", ("Billings to date", "Billings", "Cumulative billings"), ("finance",), project["billings"]),
        ("pm_etc", ("Submitted PM ETC", "PM ETC", "Submitted ETC"), ("source",), gold["pm_etc"]),
        ("finance_adjustment", ("Finance ETC adjustment", "Finance adjustment", "ETC correction"), ("finance", "difference"), net_etc_correction),
        ("close_etc", ("Recommended close ETC", "Close ETC", "Adjusted ETC"), ("finance",), project["estimated_cost_to_complete"]),
        ("eac", ("Estimated cost at completion", "EAC", "Total estimated cost"), ("finance",), project["estimated_cost_at_completion"]),
        ("percent_complete", ("Percent complete", "% complete", "POC"), ("finance",), project["percent_complete"]),
        ("earned_revenue", ("Earned revenue", "Revenue earned", "Recognized revenue"), ("finance",), project["earned_revenue"]),
        ("contract_asset", ("Contract asset", "Underbilling", "Costs in excess"), ("finance",), project["underbilling"]),
        ("contract_liability", ("Contract liability", "Overbilling", "Billings in excess"), ("finance",), project["overbilling"]),
        ("estimated_margin", ("Estimated margin at completion", "Estimated margin", "Total margin"), ("finance",), project["estimated_total_margin"]),
        ("margin_rate", ("Margin rate", "Margin percent", "Gross margin rate"), ("finance",), project["estimated_margin_percent"]),
    )
    close_rows: dict[str, dict[str, str]] = {}
    for key, aliases, roles, expected in close_specs:
        row = row_for("close", "metric", aliases)
        close_rows[key] = row
        add_number(
            f"close__{key}",
            f"The June close calculation reports the correct {aliases[0].lower()}",
            row,
            roles,
            float(expected),
            # Percentage targets are stored as decimal fractions. The generic
            # two-cent currency tolerance would otherwise accept a value up to
            # two percentage points away (for example, 82.91% for 84.26%).
            abs_tol=(.00005 if key in {"percent_complete", "margin_rate"} else .02),
        )

    add_number(
        "close__posted_cost_source",
        "The close-basis row separately reports the posted June ledger cost before cutoff accrual",
        close_rows["close_cost_basis"],
        ("source",),
        posted_cost,
    )

    recommendation = _task_001_v24_box(
        document,
        "Recommendation and review boundary",
        "Close recommendation and posting boundary",
        "Recommendation",
    )
    adjustment_evidence = "\n\n".join(
        value for value in (
            anchored_row_evidence(
                "PCO-011", "88417", "PO-2409-117", "Rogue Valley"
            ),
            numeric_candidate_evidence(
                (forecast_recovery, 1, .02),
            ),
            numeric_candidate_evidence(
                (executed_credit, 1, .02),
            ),
            numeric_candidate_evidence(
                (current_cost_credit, 1, .02),
            ),
            numeric_candidate_evidence(
                (remaining_cost_credit, 1, .02),
            ),
            numeric_candidate_evidence(
                (commercial_correction, 1, .02),
            ),
            numeric_candidate_evidence(
                (cutoff_invoice_gross, 1, .02),
            ),
            numeric_candidate_evidence(
                (cutoff_invoice_face, 1, .02),
            ),
            numeric_candidate_evidence(
                (cutoff_invoice_prior_billing, 1, .02),
            ),
            numeric_candidate_evidence(
                (cutoff_invoice_gross_june_work, 1, .02),
            ),
            numeric_candidate_evidence(
                (cutoff_accrual, 1, .02),
            ),
            numeric_candidate_evidence(
                (net_cutoff_cost_adjustment, 1, .02),
            ),
            five_column_evidence,
        ) if value
    )
    judgment_conclusion = _task_001_v24_box(
        document,
        "Judgment conclusion",
        "Close judgment",
        "Accounting judgment",
    )
    central_evidence = "\n\n".join(
        value for value in (
            narrative_evidence,
            adjustment_evidence,
            recommendation,
            judgment_conclusion,
            compact_evidence,
        ) if value
    ).strip()
    add_semantic(
        "decision__recommendation",
        (
            f"The recommendation recognizes that the signed ${executed_credit:,.2f} "
            f"subcontract adjustment was not effective at June 30, resolves "
            f"the unsupported ${commercial_correction:,.2f} PM forecast credit "
            f"without duplicating the "
            f"${pco_cost_overlap:,.2f} PO already in PM base, records the invoice cutoff without "
            "prematurely recording the conditional credit, uses "
            f"${close_etc:,.2f} ETC, and holds the resulting WIP contribution for "
            "Controller disposition"
        ),
        central_evidence,
        {
            "required_finance_meaning": (
                "The PM forecast includes a full $185,000 PCO-011 credit. The June 29 "
                "correspondence shows both signatures, while a separate contemporaneous "
                "routing record documents the ARM package release and return of the "
                "signed counterpart by June 30. The agreement also makes Cascade's "
                "delivery of CCS-CM-0629-17 to ARM Commercial Records an operative "
                "condition. The ordinary receipt stamp is July 2, and the credit memo "
                "states that it has no independent effect before the agreement conditions "
                "are satisfied. The conditional document allocates $27,000 against "
                "previously billed work and $45,000 to the unbilled subcontract balance. "
                "Neither component is effective at June 30. The agreement "
                "also identifies the PCO-011 package by field reports "
                "FR-2409-118 through FR-2409-121, while the separate vendor release carries "
                "PO-2409-117 for the same field-report package at $46,800. That PO is "
                "already carried across posted cost and PM base subcontract ETC. Invoice "
                "88417 is cumulative: $46,800 gross less $13,003.31 of prior billings "
                "leaves $33,796.69 currently due. Accounting identifies the prior draw as "
                "vendor invoice 88102 posted under AP-0004287, even though the PM snapshot "
                "still forecasts the gross PO. Finance "
                "therefore reverses the unsupported forecast credit without restoring that "
                "shared cost twice, applies no conditional credit at the cutoff, "
                "allocates invoice 88417 by line and field window, accrues only $27,196.69 "
                "of unbilled June work, removes that amount and the stale $13,003.31 prior "
                "billing from remaining ETC, and retains the $6,600 July 2 line. This "
                "produces $682,000 ETC and a $3,758,696.69 close cost basis. The resulting "
                "$301,726.52 WIP contribution "
                "remains proposed and unposted."
            ),
            "acceptable_wording": (
                "Any normal business wording is acceptable. The submission need not "
                "use 'restore,' 'unsupported,' or the reference phrasing."
            ),
            "wrong_decisions": (
                "Recognizing either component of the conditional $72,000 adjustment in "
                "the June close, retaining any unsupported commercial recovery in ETC, "
                "counting PO-2409-117 in both "
                "PM base and the PCO overlay, accruing the invoice's $6,600 July line in "
                "June, double counting the $46,800 PO, "
                "increasing customer contract value for PCO-011, "
                "or representing the "
                "workpaper as approved or posted."
            ),
        },
        category="decision",
        evidence_scope="recommendation, forecast-adjustment table, and judgment conclusion",
    )
    add_semantic(
        "decision__cutoff_accrual",
        (
            f"The recommendation reconciles the ${cutoff_invoice_gross:,.2f} cumulative bill to its "
            f"${cutoff_invoice_prior_billing:,.2f} posted prior draw and ${cutoff_invoice_face:,.2f} current amount due, records only the "
            f"${cutoff_accrual:,.2f} unbilled June portion, excludes the conditional "
            f"${current_cost_credit:,.2f} billed-work credit from June, keeps the "
            f"${post_cutoff_invoice_line:,.2f} July line in ETC, removes the prior duplicate and current June transfer from ETC, and keeps the AP cutoff outside "
            "the WIP journal"
        ),
        central_evidence,
        {
            "required_finance_meaning": (
                "Invoice 88417 is a cumulative progress bill received July 1: $46,800 "
                "gross billings less $13,003.31 of prior billings leaves $33,796.69 "
                "currently due. Accounting identifies vendor invoice 88102 / AP-0004287 "
                "as that posted prior draw; current invoice 88417 "
                "has no matching June posted job-cost, AP, or journal record. The invoice "
                "line tickets and vendor release allocate $28,800 and $11,400 to the June "
                "22-28 field window and $6,600 to July 2. The PM Commitment Detail sheet "
                "shows $33,796.69 open but still carries the full $46,800 as forecast cost. "
                "CCS-CM-0629-17 is dated June 30 and conditionally applies $27,000 against "
                "previously billed work, but the agreement requires its delivery to ARM "
                "Commercial Records and the receipt stamp is July 2. It has no matching June "
                "ledger record and no June entry is required. Finance must therefore record "
                "only the +$27,196.69 invoice cutoff in the close cost basis; remove "
                "the $13,003.31 prior-billing duplicate and the $27,196.69 current June "
                "transfer from remaining ETC, retain the $6,600 July line, and do not reduce "
                "current cost or ETC for the conditional credit at June 30. The invoice "
                "transfer has zero EAC effect and the AP cutoff does not belong in the WIP journal."
            ),
            "acceptable_wording": (
                "Accept ordinary cutoff, accrual, incurred-cost, commitment-relief, and "
                "double-counting language. Do not require the reference labels."
            ),
            "wrong_decisions": (
                "Ignoring either document, prematurely recording the conditional credit, "
                "netting away the gross AP cutoff evidence, "
                "leaving the invoice both in incurred cost and ETC, reducing ETC for the "
                "conditional billed-work credit, including the July 2 line in June cost, "
                "or putting the AP cutoff into the WIP journal."
            ),
        },
        category="decision",
        evidence_scope="recommendation, forecast-adjustment table, and judgment conclusion",
    )
    add_semantic(
        "decision__forecast_overlap",
        (
            "The recommendation reconciles PO-2409-117 across PM base ETC, the $185,000 "
            "PCO cost build, and invoice 88417 so the same $46,800 cost is included once"
        ),
        central_evidence,
        {
            "source_facts": (
                "The PM Commitment Detail sheet shows PO-2409-117 at $33,796.69 open but "
                "still forecasts $46,800 inside base subcontract ETC. Invoice 88417 deducts "
                "$13,003.31 of prior billings; accounting records identify that draw as "
                "vendor invoice 88102 / AP-0004287. The July 2 contract-records batch contains a PCO-011 "
                "agreement covering field reports FR-2409-118 through FR-2409-121 and, on "
                "a separate page, a $46,800 Rogue Valley release under PO-2409-117 for "
                "that same field-report package. Invoice 88417 carries the same project, "
                "vendor package, PO, released line items, and amount."
            ),
            "required_finance_meaning": (
                "Only $138,200 of the $185,000 gross PCO cost exposure is incremental to "
                "the project base. Of the invoice's $40,200 gross June work, $13,003.31 is "
                "already posted and $27,196.69 requires a current accrual; both amounts "
                "leave stale PM ETC while the $6,600 July 2 line remains there. The full "
                "$46,800 PO must not also be restored through the PCO overlay. "
                "The signed agreement's $27,000 billed-work and $45,000 unbilled-balance "
                "components are not effective until the July 2 Commercial Records delivery, "
                "so neither reduces June current cost or ETC. The resulting net ETC "
                "correction is $98,000."
            ),
            "grading_boundary": (
                "Judge the cross-source cost reconciliation semantically. Accept any "
                "professional wording or schedule that includes the shared PO cost once. "
                "Do not require the words overlap, duplicate, incremental, or overlay."
            ),
        },
        category="decision",
        evidence_scope=(
            "recommendation, forecast-adjustment table, cutoff issue, source record, "
            "and judgment conclusion"
        ),
    )

    # Prior-close bridge. June endpoints are already scored above, so this
    # section grades May comparatives and movements without triple-counting them.
    bridge_aliases = {
        "close_cost_basis": ("Cost to date / close basis", "Close cost basis", "Cost to date", "Posted cost", "Posted cost to date"),
        "close_etc": ("Remaining cost / ETC", "Remaining cost", "Close ETC", "ETC"),
        "eac": ("Estimated cost at completion", "EAC", "Total estimated cost"),
        "earned_revenue": ("Earned revenue", "Revenue earned", "Recognized revenue"),
        "billings": ("Billings", "Cumulative billings"),
        "contract_asset": ("Contract asset", "Underbilling", "Costs in excess"),
        "estimated_margin": ("Estimated margin at completion", "Estimated margin", "Total margin"),
        "margin_rate": ("Margin rate", "Margin percent", "Gross margin rate"),
    }
    bridge_values = {
        "close_cost_basis": (prior["posted_cost"], bridge_gold["close_cost_basis_change"]),
        "close_etc": (prior["close_etc"], bridge_gold["close_etc_change"]),
        "eac": (prior["eac"], bridge_gold["eac_change"]),
        "earned_revenue": (prior["earned_revenue"], bridge_gold["earned_revenue_change"]),
        "billings": (prior["billings"], bridge_gold["billings_change"]),
        "contract_asset": (prior["contract_asset"], bridge_gold["contract_asset_change"]),
        "estimated_margin": (prior["estimated_margin"], bridge_gold["estimated_margin_change"]),
        "margin_rate": (prior["margin_percent"], bridge_gold["margin_rate_change_bps"]),
    }
    bridge_rows: dict[str, dict[str, str]] = {}
    for key, aliases in bridge_aliases.items():
        row = row_for("bridge", "metric", aliases)
        bridge_rows[key] = row
        may_value, change_value = bridge_values[key]
        add_number(
            f"bridge__{key}__may",
            f"The bridge reports the correct final-May {aliases[0].lower()}",
            row,
            ("may",),
            float(may_value),
            # Decimal-fraction rate targets need rate precision, not the
            # generic two-cent currency tolerance. Otherwise a visibly wrong
            # two-decimal percentage can receive credit merely because it is
            # within two percentage points of the target.
            abs_tol=(.00005 if key == "margin_rate" else .02),
        )
        if key == "margin_rate":
            add_bps(
                "bridge__margin_rate__change",
                "The bridge reports the correct margin-rate movement",
                row,
                "change",
                float(change_value),
            )
        else:
            add_number(
                f"bridge__{key}__change",
                f"The bridge reports the correct change in {aliases[0].lower()}",
                row,
                ("change",),
                float(change_value),
            )

    movement_box = _task_001_v24_box(
        document,
        "Material movement explanation",
        "Movement explanation",
        "Bridge explanation",
    )
    movement_numeric_rows = "\n\n".join(
        numeric_candidate_evidence(
            (float(may_value), 1, .00005 if key == "margin_rate" else .02),
        )
        for key, (may_value, _change_value) in bridge_values.items()
    )
    movement_evidence = "\n\n".join(
        value for value in (
            movement_box,
            narrative_evidence,
            movement_numeric_rows,
            five_column_evidence,
        ) if value
    ).strip()
    add_semantic(
        "bridge__material_movement",
        "The memo explains the material May-to-June cost, billing, EAC, contract-asset, and margin movements without contradicting the reported numbers",
        movement_evidence,
        {
            "expected_causal_chain": (
                "June posted cost increases $161,892.13 and already includes vendor invoice "
                "88102 / AP-0004287 for $13,003.31. Invoice 88417 reports $46,800 gross, "
                "deducts $13,003.31 of prior billings, and leaves $33,796.69 due. Gross June work is $40,200 and the "
                "$6,600 July 2 line stays in ETC, so the separate $27,196.69 current-invoice "
                "cutoff raises the close cost basis by $189,088.82 from May. The signed "
                "$72,000 adjustment is not effective at June 30 because its referenced "
                "credit memorandum reached ARM Commercial Records July 2. Removing both the prior-billing duplicate "
                "and current June transfer from remaining ETC prevents double counting. Billings increase "
                "$79,513.55. Only $138,200 of the $185,000 PCO cost support is incremental "
                "to PM base because PO-2409-117 is already carried there, and no conditional "
                "credit component reduces June cost. EAC increases and estimated margin "
                "decreases by $125,196.69 from May. "
                "The contract asset increases $26,921.36 because earned-revenue growth "
                "exceeds billing growth."
            ),
            "grading_boundary": (
                "Judge the causal finance meaning, not wording or whether every number "
                "is repeated in prose; the numeric bridge is scored separately."
            ),
        },
        evidence_scope="material-movement box and complete bridge rows",
    )

    # Compact accounting rollforward.
    reconciliation_specs = (
        (
            "job_cost",
            ("Posted job cost", "Job cost", "Total job cost", "Total project cost"),
            3_569_607.87,
            161_892.13,
            3_731_500.00,
        ),
        (
            "billings",
            ("Posted billings", "Billings", "Project billings"),
            4_296_079.58,
            79_513.55,
            4_375_593.13,
        ),
        (
            "contract",
            ("Current contract", "Executed contract", "Contract value"),
            5_526_000.00,
            0.0,
            5_526_000.00,
        ),
    )
    for key, aliases, may_value, activity, june_value in reconciliation_specs:
        row = row_for("reconciliation", "control", aliases)
        canonical_rollforward_met = bool(row) and all((
            _text_contains_number(row.get("may"), may_value),
            _text_contains_number(row.get("activity"), activity),
            _text_contains_number(row.get("june"), june_value),
        ))
        requirements: list[tuple[float, int, float]] = []
        if _close(may_value, june_value, abs_tol=.000001, rel_tol=0.0):
            requirements.append((may_value, 2, .02))
        else:
            requirements.extend(((may_value, 1, .02), (june_value, 1, .02)))
        requirements.append((activity, 1, .02))
        submitted = numeric_candidate_evidence(*requirements)
        add_semantic(
            f"accounting__{key}__rollforward",
            f"The {aliases[0].lower()} control rolls May plus June activity to the June close",
            submitted,
            {
                "required_control": aliases[0],
                "may_final": may_value,
                "june_activity": activity,
                "june_close": june_value,
                "required_equation": "May final plus June activity equals June close",
                "grading_boundary": (
                    "The code gate verifies only that the exact three values occur in "
                    "one submitted row. Decide whether the row is the named accounting "
                    "control and the values occupy the correct period roles."
                ),
            },
            hard_gate=bool(submitted),
            evidence_scope="all rows containing the exact rollforward values",
            legacy_met=canonical_rollforward_met,
            always_judge=False,
        )
        canonical_variance_met = bool(row) and (
            _task_001_v24_zero(row.get("variance"))
            or (
                canonical_rollforward_met
                and not _task_001_v24_contradicts_zero(row.get("variance"))
            )
        )
        add_semantic(
            f"accounting__{key}__variance",
            f"The {aliases[0].lower()} control has zero variance",
            submitted,
            {
                "required_control": aliases[0],
                "may_final": may_value,
                "june_activity": activity,
                "june_close": june_value,
                "expected_variance": 0.0,
                "acceptable_equivalent": (
                    "An explicit zero, tied/OK status, or an uncontradicted exact "
                    "May-plus-activity-equals-June equation in the row."
                ),
                "grading_boundary": (
                    "Decide the meaning of the variance/status field; a zero belonging "
                    "to a different row or metric is insufficient."
                ),
            },
            category="auditability",
            weight=5,
            hard_gate=bool(submitted),
            evidence_scope="all rows containing the exact rollforward values",
            legacy_met=canonical_variance_met,
            always_judge=False,
        )

    # Practical populations: amounts and ties, not hidden transaction registers.
    population_specs = (
        (
            "ap",
            (
                "AP project cost", "AP vendor bills", "Accounts payable project cost",
                "Vendor bills", "Job cost - AP", "AP",
            ),
            107_393.79,
        ),
        (
            "payroll",
            (
                "Payroll project cost", "Payroll", "Field payroll", "Labor payroll",
                "Job cost - PAY", "PAY (field labor)",
                "June job cost - PAY (field labor)",
            ),
            54_498.34,
        ),
        (
            "job_cost_total",
            (
                "Total job cost", "Job cost total", "Project cost total",
                "Total / control", "Total control", "June job cost",
                "June job cost - total", "Job cost - June",
            ),
            161_892.13,
        ),
        (
            "billing",
            ("June billing", "Billing", "Billings", "Invoice population"),
            79_513.55,
        ),
    )
    population_rows: dict[str, dict[str, str]] = {}
    for key, aliases, expected in population_specs:
        row = row_for("population", "population", aliases)
        population_rows[key] = row
        submitted_pair = numeric_candidate_evidence((expected, 2, .02))
        canonical_source_to_ledger_met = bool(row) and all((
            _text_contains_number(row.get("source_total"), expected),
            _text_contains_number(row.get("ledger_total"), expected),
        ))
        add_semantic(
            f"population__{key}__source_to_ledger",
            f"The {aliases[0].lower()} source and ledger totals both equal the supported amount",
            submitted_pair,
            {
                "required_population": aliases[0],
                "expected_source_total": expected,
                "expected_ledger_total": expected,
                "grading_boundary": (
                    "The code gate verifies only that the exact supported amount appears "
                    "at least twice in one submitted row. Decide whether that row is the "
                    "named population and the two values are respectively its source and "
                    "posted-ledger totals. Normal population labels and table layouts are "
                    "fully acceptable."
                ),
            },
            hard_gate=bool(submitted_pair),
            evidence_scope="all rows containing two instances of the supported amount",
            legacy_met=canonical_source_to_ledger_met,
            always_judge=False,
        )
        variance_evidence = numeric_candidate_evidence(
            (expected, 2, .02),
            (0.0, 1, .000001),
        )
        add_semantic(
            f"population__{key}__variance",
            f"The {aliases[0].lower()} population has zero variance",
            variance_evidence,
            {
                "required_population": aliases[0],
                "expected_source_total": expected,
                "expected_ledger_total": expected,
                "expected_variance": 0.0,
                "grading_boundary": (
                    "The code gate verifies only that a row contains the two exact totals "
                    "and an explicit zero. Decide whether the zero is the variance for "
                    "that same like-for-like population rather than another field."
                ),
            },
            category="auditability",
            weight=5,
            hard_gate=bool(variance_evidence),
            evidence_scope="all rows containing the two exact totals and a zero",
            legacy_met=(
                canonical_source_to_ledger_met
                and _task_001_v24_zero(row.get("variance"))
            ),
            always_judge=False,
        )

    population_box = _task_001_v24_box(
        document,
        "Accounting completeness and cutoff conclusion",
        "Accounting completeness conclusion",
        "Population control conclusion",
    )
    population_evidence = "\n".join([
        population_box,
        population_table_evidence,
        numeric_candidate_evidence((107_393.79, 1, .02)),
        numeric_candidate_evidence((54_498.34, 1, .02)),
        numeric_candidate_evidence((161_892.13, 1, .02)),
        numeric_candidate_evidence((79_513.55, 1, .02)),
        numeric_candidate_evidence((cutoff_invoice_gross, 1, .02)),
        numeric_candidate_evidence((cutoff_invoice_face, 1, .02)),
        numeric_candidate_evidence((cutoff_invoice_prior_billing, 1, .02)),
        numeric_candidate_evidence((cutoff_invoice_gross_june_work, 1, .02)),
        numeric_candidate_evidence((cutoff_accrual, 1, .02)),
        numeric_candidate_evidence((post_cutoff_invoice_line, 1, .02)),
        numeric_candidate_evidence((current_cost_credit, 1, .02)),
        numeric_candidate_evidence((-current_cost_credit, 1, .02)),
        narrative_evidence,
    ]).strip()
    add_semantic(
        "population__completeness",
        "The memo reaches a supported completeness conclusion for June AP, payroll, total job cost, billing, and the identified late documents rather than making an unsupported blanket assertion",
        population_evidence,
        {
            "supported_control": (
                "AP project cost $107,393.79, payroll project cost $54,498.34, "
                "total project job cost $161,892.13, and June billing $79,513.55 "
                "tie to posted records. Invoice 88417 is separately identified as a "
                "cumulative mixed-period progress bill with $46,800 gross billings, "
                "prior invoice 88102 of $13,003.31, and $33,796.69 currently due. The "
                "prior draw is posted as AP-0004287 while current invoice 88417 is absent; "
                "gross June work is $40,200, its July 2 line is $6,600, and only $27,196.69 "
                "of unbilled June work requires a current accrual. "
                "CCS-CM-0629-17 is separately identified as a conditional $27,000 billed-work "
                "credit received by Commercial Records July 2 and absent from the June ledger. A compact "
                "control is sufficient; no 23-line or 13-journal register is required."
            ),
            "grading_boundary": (
                "Judge the authored control across the population table, accounting "
                "conclusion, and related rollforward evidence. A separately quantified "
                "AP/payroll split plus a supported total-job-cost tie may establish this "
                "overall completeness conclusion even when the AP and payroll amounts "
                "are not each repeated in separate source-versus-ledger rows; the distinct "
                "AP and payroll source-to-ledger criteria score that stronger control. "
                "A conclusion that treats the entire cumulative invoice or current amount "
                "due as unposted June work "
                "does not establish a complete cutoff population. "
                "Do not require a particular table order or repeated reference wording. "
                "If the submission identifies CCS-CM-0629-17 as a conditional $27,000 "
                "billed-work credit received July 2 and states that it is absent from the June posted ledger, "
                "acknowledge that supplied control. Do not call that credit unsupported "
                "merely because another row describes its $45,000 remaining-balance "
                "component. Whether the submission applies those amounts in June is scored "
                "by the timing and commercial-treatment criteria. If the overall criterion still fails, identify the actual "
                "missing AP/payroll split, prior-posting link, timing allocation, or "
                "other visible completeness defect."
            ),
            "unsupported_conclusion": (
                "Saying only 'complete' or 'no exceptions' without identifying the "
                "controlled populations and evidence."
            ),
        },
        category="auditability",
        weight=5,
        evidence_scope="population table and accounting-completeness conclusion",
    )

    cutoff_row = row_for(
        "population",
        "population",
        (
            "Late vendor invoice", "Invoice 88417", "Vendor invoice cutoff",
            "Cutoff accrual", "Late AP invoice",
        ),
    )
    cutoff_source_reconciliation_evidence = "\n\n".join(
        value for value in (
            anchored_row_evidence("88417"),
            numeric_candidate_evidence((cutoff_invoice_gross, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_prior_billing, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_face, 1, .02)),
            population_table_evidence,
            narrative_evidence,
        ) if value
    )
    canonical_cutoff_source_met = bool(cutoff_row) and all(
        _text_contains_number(row_text(cutoff_row), value)
        for value in (
            cutoff_invoice_gross,
            cutoff_invoice_prior_billing,
            cutoff_invoice_face,
        )
    )
    cutoff_source_objective_gate = all(
        _text_contains_number(text, value)
        for value in (
            cutoff_invoice_gross,
            cutoff_invoice_prior_billing,
            cutoff_invoice_face,
        )
    ) and "88417" in normalized
    add_semantic(
        "population__cutoff_invoice__source_amount",
        "The cutoff control reconciles invoice 88417's gross billings, prior invoice, and current amount due",
        cutoff_source_reconciliation_evidence,
        {
            "required_reconciliation": (
                "$46,800 gross billings less $13,003.31 of prior billings equals "
                "the $33,796.69 amount currently due on invoice 88417"
            ),
            "grading_boundary": (
                "Accept the three values in a table, narrative, or linked rows and any "
                "normal cumulative-billing terminology. The values must be associated "
                "with invoice 88417 and its prior-billings deduction; a standalone $46,800 amount is "
                "not sufficient."
            ),
        },
        category="auditability",
        weight=10,
        hard_gate=cutoff_source_objective_gate,
        evidence_scope="cutoff population and accounting-completeness evidence",
        legacy_met=canonical_cutoff_source_met,
        always_judge=False,
    )
    cutoff_current_absence_evidence = "\n\n".join(
        value for value in (
            anchored_row_evidence("88417"),
            numeric_candidate_evidence(
                (cutoff_invoice_face, 1, .02), (0.0, 1, .000001)
            ),
            population_table_evidence,
            narrative_evidence,
        ) if value
    )
    canonical_cutoff_current_absence_met = bool(cutoff_row) and (
        _text_contains_number(row_text(cutoff_row), cutoff_invoice_face)
        and (
            _task_001_v24_zero(cutoff_row.get("ledger_total"))
            or "0 current" in _normalize(cutoff_row.get("ledger_total"))
        )
    )
    cutoff_current_absence_objective_gate = (
        _text_contains_number(text, cutoff_invoice_face)
        and "88417" in normalized
    )
    add_semantic(
        "population__cutoff_invoice__ledger_absence",
        "The cutoff control establishes that current invoice 88417 is absent from the June posted ledger",
        cutoff_current_absence_evidence,
        {
            "required_association": (
                "The $33,796.69 current amount due on invoice 88417 has no matching "
                "June job-cost, vendor-bill, or posted-journal record"
            ),
            "grading_boundary": (
                "The prior $13,003.31 draw is posted, so do not require the entire "
                "cumulative invoice to be absent. Accept an explicit zero-current amount "
                "or equivalent professional narrative tied to invoice 88417."
            ),
        },
        category="auditability",
        weight=5,
        hard_gate=cutoff_current_absence_objective_gate,
        evidence_scope="cutoff population and accounting-completeness evidence",
        legacy_met=canonical_cutoff_current_absence_met,
        always_judge=False,
    )
    prior_posting_evidence = "\n\n".join(
        value for value in (
            anchored_row_evidence("88102", "AP-0004287", "Rogue Valley"),
            numeric_candidate_evidence((cutoff_invoice_prior_billing, 1, .02)),
            population_table_evidence,
            narrative_evidence,
            five_column_evidence,
        ) if value
    )
    canonical_prior_posting_met = bool(cutoff_row) and all((
        _text_contains_number(row_text(cutoff_row), cutoff_invoice_prior_billing),
        "ap 0004287" in _normalize(row_text(cutoff_row)),
        "prior" in _normalize(row_text(cutoff_row)),
    ))
    prior_posting_objective_gate = all((
        _text_contains_number(text, cutoff_invoice_prior_billing),
        "88102" in normalized,
        "ap 0004287" in normalized,
    ))
    add_semantic(
        "population__cutoff_invoice__prior_posting",
        "The cutoff control matches prior invoice 88102 to the $13,003.31 June posting AP-0004287",
        prior_posting_evidence,
        {
            "required_match": (
                "Prior vendor invoice 88102 for $13,003.31 is already posted in June "
                "job cost, AP, and the journal under internal document AP-0004287"
            ),
            "grading_boundary": (
                "Accept ordinary wording that connects the vendor's prior invoice number, "
                "amount, and the accounting system's internal document number. Do not "
                "penalize the use of 'prior draw,' 'previous billing,' or equivalent labels."
            ),
        },
        category="auditability",
        weight=10,
        hard_gate=prior_posting_objective_gate,
        evidence_scope="cutoff, population, and posted-accounting evidence",
        legacy_met=canonical_prior_posting_met,
        always_judge=False,
    )
    eligible_invoice_evidence = "\n\n".join(
        value for value in (
            numeric_candidate_evidence((cutoff_accrual, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_prior_billing, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_gross_june_work, 1, .02)),
            numeric_candidate_evidence((post_cutoff_invoice_line, 1, .02)),
            anchored_row_evidence("88417", "RV-2409-0628", "RV-2409-0702"),
            population_table_evidence,
            narrative_evidence,
        ) if value
    )
    canonical_eligible_invoice_met = bool(cutoff_row) and all(
        _text_contains_number(row_text(cutoff_row), value)
        for value in (
            cutoff_invoice_gross_june_work,
            cutoff_invoice_prior_billing,
            cutoff_accrual,
            post_cutoff_invoice_line,
        )
    )
    eligible_invoice_objective_gate = all(
        _text_contains_number(text, value)
        for value in (
            cutoff_invoice_gross_june_work,
            cutoff_invoice_prior_billing,
            cutoff_accrual,
            post_cutoff_invoice_line,
        )
    ) and "88417" in normalized
    add_semantic(
        "population__cutoff_invoice__june_eligible_amount",
        "The cutoff control derives the $27,196.69 current June accrual from gross June work and the posted prior draw while retaining the $6,600 July 2 line",
        eligible_invoice_evidence,
        {
            "source_line_allocation": (
                "Invoice 88417 reports $46,800 gross billings, $13,003.31 previously "
                "billed, and $33,796.69 currently due. Field tickets and the matching vendor "
                "release assign $28,800 and $11,400 to the June 22-28 field window and "
                "$6,600 to the July 2 night shift. Gross June work is $40,200; after "
                "matching the posted prior draw, the current June accrual is $27,196.69, "
                "and the final $6,600 remains in ETC at June 30."
            ),
            "grading_boundary": (
                "Accept any normal business presentation that unambiguously reaches the "
                "$46,800 gross / $13,003.31 prior / $33,796.69 due reconciliation and "
                "$40,200 gross June / $27,196.69 current June / $6,600 post-cutoff allocation. The exact numbers alone "
                "are insufficient when attached to another invoice or period."
            ),
        },
        category="auditability",
        weight=10,
        hard_gate=eligible_invoice_objective_gate,
        evidence_scope="invoice-cutoff rows and authored line-level timing analysis",
        legacy_met=canonical_eligible_invoice_met,
        always_judge=False,
    )
    credit_memo_row = row_for(
        "population",
        "population",
        (
            "Late vendor invoice", "CCS-CM-0629-17",
            "Vendor credit cutoff", "Credit memorandum", "Late vendor credit",
        ),
    )
    credit_amount_evidence = "\n\n".join(
        value for value in (
            numeric_candidate_evidence((current_cost_credit, 1, .02)),
            numeric_candidate_evidence((-current_cost_credit, 1, .02)),
        ) if value
    )
    canonical_credit_source_met = bool(credit_memo_row) and (
        _text_contains_number(credit_memo_row.get("source_total"), current_cost_credit)
        or _text_contains_number(credit_memo_row.get("source_total"), -current_cost_credit)
    )
    add_semantic(
        "population__cutoff_credit_memo__source_amount",
        "The cutoff control reports the $27,000 source credit-memorandum amount",
        credit_amount_evidence,
        {
            "required_association": (
                "CCS-CM-0629-17 source amount is a $27,000 vendor credit against "
                "previously billed ARM-2409 work"
            ),
            "expected_magnitude": current_cost_credit,
            "expected_direction": "credit or reduction of cost",
            "grading_boundary": (
                "Accept parentheses, a minus sign, a positive magnitude paired with an "
                "unambiguous credit label, or another normal finance sign presentation. "
                "The exact number alone is insufficient unless associated with this document."
            ),
        },
        category="auditability",
        weight=10,
        hard_gate=bool(credit_amount_evidence),
        evidence_scope="all submitted rows containing the $27,000 magnitude",
        legacy_met=canonical_credit_source_met,
        always_judge=False,
    )
    credit_absence_evidence = "\n\n".join(
        value for value in (
            numeric_candidate_evidence(
                (current_cost_credit, 1, .02), (0.0, 1, .000001)
            ),
            numeric_candidate_evidence(
                (-current_cost_credit, 1, .02), (0.0, 1, .000001)
            ),
            anchored_row_evidence(
                "CCS-CM-0629-17", "credit memo", "credit memorandum"
            ),
            population_table_evidence,
            narrative_evidence,
        ) if value
    )
    canonical_credit_absence_met = bool(credit_memo_row) and all((
        canonical_credit_source_met,
        (
            _task_001_v24_zero(credit_memo_row.get("ledger_total"))
            or "0 credit" in _normalize(credit_memo_row.get("ledger_total"))
        ),
    ))
    add_semantic(
        "population__cutoff_credit_memo__ledger_absence",
        "The cutoff control reports no matching credit memorandum in the June posted ledger",
        credit_absence_evidence,
        {
            "required_association": (
                "CCS-CM-0629-17 has a $27,000 source credit and zero matching June "
                "posted-ledger amount"
            ),
            "grading_boundary": (
                "Accept either a dedicated zero in the June-ledger field or distributed "
                "professional narrative that unambiguously says this $27,000 credit is "
                "not yet in AP, job cost, or the June posted ledger. Do not require the "
                "absence conclusion to share a row with the source amount, and do not "
                "mistake an ETC zero for a ledger zero. Fail when the submission only "
                "identifies the document, schedules future processing, or shows a zero "
                "for another field without establishing June-ledger absence."
            ),
        },
        category="auditability",
        weight=5,
        hard_gate=bool(credit_absence_evidence),
        evidence_scope=(
            "credit-memorandum rows plus the complete population and accounting-"
            "completeness narrative"
        ),
        legacy_met=canonical_credit_absence_met,
        always_judge=False,
    )
    canonical_credit_variance_met = bool(credit_memo_row) and (
        _text_contains_number(credit_memo_row.get("variance"), current_cost_credit)
        or _text_contains_number(credit_memo_row.get("variance"), -current_cost_credit)
    )
    add_semantic(
        "population__cutoff_credit_memo__variance",
        "The cutoff control reports the $27,000 source-to-ledger credit difference",
        credit_amount_evidence,
        {
            "required_association": (
                "The source-to-ledger difference for CCS-CM-0629-17 is a $27,000 credit"
            ),
            "expected_magnitude": current_cost_credit,
            "expected_direction": "credit or reduction of cost",
            "grading_boundary": (
                "Accept ordinary debit/credit, variance, exception, parentheses, and sign "
                "presentation. Decide whether the amount is the credit memo's difference, "
                "not merely its source amount repeated elsewhere."
            ),
        },
        category="auditability",
        weight=10,
        hard_gate=bool(credit_amount_evidence),
        evidence_scope="all submitted rows containing the $27,000 magnitude",
        legacy_met=canonical_credit_variance_met,
        always_judge=False,
    )
    cutoff_issue_row = row_for(
        "judgment",
        "issue",
        (
            "Invoice 88417 cutoff", "Invoice 88417", "Late vendor invoice",
            "Vendor invoice cutoff", "Cutoff accrual",
        ),
    )
    cutoff_source_row = row_for(
        "source",
        "source",
        (
            "AP intake queue and invoice scan batch p. 4", "Invoice 88417",
            "batch 20260701-02 page 4", "Rogue Valley invoice", "Late vendor invoice",
        ),
    )
    cutoff_evidence = "\n\n".join(
        value for value in (
            anchored_row_evidence("88417", "PO-2409-117", "Rogue Valley"),
            anchored_row_evidence("88102", "AP-0004287"),
            anchored_row_evidence("CCS-CM-0629-17", "credit memo"),
            numeric_candidate_evidence((cutoff_invoice_gross, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_face, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_prior_billing, 1, .02)),
            numeric_candidate_evidence((cutoff_invoice_gross_june_work, 1, .02)),
            numeric_candidate_evidence((cutoff_accrual, 1, .02)),
            numeric_candidate_evidence((post_cutoff_invoice_line, 1, .02)),
            credit_amount_evidence,
            population_table_evidence,
            narrative_evidence,
            adjustment_evidence,
            compact_evidence,
        ) if value
    ).strip()
    add_semantic(
        "population__cutoff_accrual_treatment",
        "The cutoff conclusion records invoice 88417 as a separate gross June cutoff, treats CCS-CM-0629-17 according to its post-cutoff effectiveness, and classifies the ETC effects without double counting",
        cutoff_evidence,
        {
            "source_facts": (
                "Invoice 88417 is a cumulative ARM-2409 progress bill dated June 29 and "
                "received in the AP inbox July 1 at 07:42 PT. It reports $46,800 gross, "
                "$13,003.31 of prior billings, and $33,796.69 currently due. The "
                "accounting system identifies that prior draw as vendor invoice 88102 / "
                "AP-0004287 in June job cost, "
                "AP, and the journal, while current invoice 88417 is absent. Its tickets and "
                "matching vendor release assign $28,800 and $11,400 to the June 22-28 "
                "field window and $6,600 to the July 2 night shift. No "
                "matching June job-cost entry, vendor bill, or posted project journal exists "
                "for the current amount due. The AP intake row and invoice match PO-2409-117, "
                "vendor, project, package, and amount to a PM Commitment Detail row showing "
                "$33,796.69 open but $46,800 forecast. CCS-CM-0629-17 is dated June 30 and "
                "conditionally applies $27,000 against previously billed work. The agreement "
                "requires its delivery to ARM Commercial Records, page 1 records receipt on "
                "July 2, and no matching June ledger record exists."
            ),
            "required_treatment": (
                "Record or approve a separate gross cutoff entry for the $27,196.69 unbilled "
                "June invoice portion, include that amount in the WIP close cost basis, and "
                "do not record the conditional $27,000 credit in June. Remove both the $13,003.31 prior-billing "
                "duplicate and $27,196.69 current June transfer from remaining ETC, "
                "retain the $6,600 July 2 line in ETC, and do not reduce current cost or ETC "
                "for either conditional credit component at June 30. "
                "Retie before WIP release; the invoice transfer has no EAC effect."
            ),
            "grading_boundary": (
                "Judge the accounting logic semantically. Accept normal business wording "
                "and a conservative posting hold; do not require the words cutoff, transfer, "
                "or double count if the same treatment is clear."
            ),
        },
        category="auditability",
        weight=10,
        evidence_scope=(
            "cutoff population row, cutoff issue row, adjustment table, accounting "
            "conclusion, and invoice/credit source records"
        ),
    )

    ap_issue_row = row_for(
        "judgment",
        "issue",
        ("AP-0004290", "Vendor payment hold", "AP payment hold"),
    )
    add_semantic(
        "population__ap_hold_treatment",
        "The memo correctly distinguishes AP-0004290's payment hold from recognition of its posted June cost",
        "\n\n".join(
            value for value in (
                anchored_row_evidence("AP-0004290"),
                population_table_evidence,
                action_like_evidence,
                narrative_evidence,
                compact_evidence,
            ) if value
        ).strip(),
        {
            "expected_treatment": (
                "AP-0004290 is posted subcontract cost. Its vendor-payment hold affects "
                "settlement timing, not whether the cost remains in June job cost."
            ),
        },
        category="auditability",
        weight=5,
        evidence_scope="AP issue row and accounting-completeness conclusion",
    )
    billing_issue_row = row_for(
        "judgment",
        "issue",
        ("June invoice", "Billing collection", "Customer collection", "June billing"),
    )
    add_semantic(
        "population__billing_collection_treatment",
        "The memo correctly separates the invoice's partial collection from recognition of the full non-voided June billing",
        "\n\n".join(
            value for value in (
                anchored_row_evidence("PAYAPP-ARM-2409-202606"),
                numeric_candidate_evidence((79_513.55, 1, .02)),
                numeric_candidate_evidence((9_064.54, 1, .02)),
                numeric_candidate_evidence((70_449.01, 1, .02)),
                population_table_evidence,
                action_like_evidence,
                narrative_evidence,
                compact_evidence,
            ) if value
        ).strip(),
        {
            "expected_treatment": (
                "The $79,513.55 invoice is non-voided and fully included in cumulative "
                "billings. $9,064.54 was collected and $70,449.01 remained open; that "
                "cash status does not reduce billing for WIP."
            ),
            "source_distinction": (
                "Customer collection is different from the AP vendor-payment hold."
            ),
        },
        category="auditability",
        weight=5,
        evidence_scope="billing/collection issue row and accounting-completeness conclusion",
    )

    # Judgment is entirely semantic; regex must not decide authority or policy meaning.
    pco_row = row_for(
        "judgment",
        "issue",
        ("PCO-011", "Concrete trade recovery", "Backcharge recovery", "Commercial recovery"),
    )
    pco_evidence = "\n\n".join(
        value for value in (
            anchored_row_evidence("PCO-011"),
            numeric_candidate_evidence((forecast_recovery, 1, .02)),
            numeric_candidate_evidence((executed_credit, 1, .02)),
            numeric_candidate_evidence((current_cost_credit, 1, .02)),
            numeric_candidate_evidence((-current_cost_credit, 1, .02)),
            numeric_candidate_evidence((remaining_cost_credit, 1, .02)),
            numeric_candidate_evidence((commercial_correction, 1, .02)),
            numeric_candidate_evidence((project["current_contract"], 1, .02)),
            narrative_evidence,
            adjustment_evidence,
            compact_evidence,
        ) if value
    ).strip()
    add_semantic(
        "commercial__authority_at_cutoff",
        "The memo accurately assesses PCO-011's authority and evidence at June 30",
        pco_evidence,
        {
            "source_facts": (
                "The commercial master and draft packet retain the full $185,000 working "
                "position. Separate June 29 correspondence is signed by Cascade on June "
                "29 and ARM on June 30. A separate June 30 routing transmittal records ARM's "
                "June 29 package release and return of the executed counterpart to both parties "
                "on June 30. The agreement also requires Cascade to deliver CCS-CM-0629-17 "
                "to ARM Commercial Records before the reduction becomes effective. Page 1's "
                "ordinary receipt field shows July 2 at 07:41 PT and states that the memo "
                "has no independent effect before the agreement conditions are satisfied. "
                "The $72,000 adjustment therefore is not effective at June 30. It conditionally "
                "allocates $45,000 to the remaining unbilled balance and $27,000 against "
                "previously billed work."
            ),
            "grading_boundary": (
                "Judge whether the authority assessment follows from the raw evidence. "
                "Both signatures alone are insufficient because the agreement also "
                "requires package release, return of the executed counterpart, and delivery "
                "of the referenced credit memo to ARM Commercial Records. Fail when the "
                "submission recognizes the $72,000 in the June close despite the July 2 "
                "delivery evidence, or leaves any operative condition unevidenced. "
                "Do not require the words draft, disputed, unaccepted, or unsupported "
                "when equivalent business wording is used."
            ),
        },
        evidence_scope="PCO issue row, adjustment schedule, and judgment conclusion",
    )
    add_semantic(
        "commercial__record_conflict",
        "The memo resolves the earlier working records and later signed correspondence by source authority and cutoff rather than by status wording alone",
        pco_evidence,
        {
            "expected_resolution": (
                "The June 30 master export and draft packet remain evidence of the gross "
                "$185,000 working position. The separate correspondence records Cascade's "
                "June 29 and ARM's June 30 signatures. A contemporaneous routing transmittal "
                "records ARM's June 29 package release and return of the executed counterpart "
                "to both parties June 30. The agreement's remaining operative condition is "
                "delivery of CCS-CM-0629-17 to ARM Commercial Records. The memo was received "
                "July 2 and has no independent effect before the agreement conditions are "
                "satisfied. Therefore the later records do not support a June credit, though "
                "they conditionally allocate $27,000 to previously billed work and $45,000 "
                "to the unbilled balance for subsequent treatment. The full $185,000 forecast "
                "recovery remains unsupported at June 30."
            ),
            "grading_boundary": (
                "Accept any accurate professional explanation of the dates, signature and "
                "delivery conditions, scope, and source hierarchy. Do not require the words stale, "
                "superseded, controlling, executed, or conflict. Do not award credit "
                "merely because the later signed document is called controlling: the "
                "submission must establish the package release, executed return, and credit-"
                "memo delivery chronology. Fail when it recognizes the $72,000 in June "
                "despite the July 2 delivery, because that does not resolve the earlier "
                "working record against the later evidence."
            ),
        },
        evidence_scope="PCO issue row, adjustment schedule, and judgment conclusion",
    )
    add_semantic(
        "commercial__contract_treatment",
        "The memo keeps PCO-011 out of June current contract value unless supported authorization exists",
        "\n\n".join(
            value for value in (
                numeric_candidate_evidence((project["current_contract"], 1, .02)),
                pco_evidence,
            ) if value
        ),
        {
            "expected_policy_application": (
                "Current contract value remains $5,526,000. The conditional $72,000 "
                "subcontract adjustment and its billed-work credit memorandum do not amend the "
                "Northline customer contract or authorize customer billing."
            ),
        },
        evidence_scope="PCO issue row, adjustment schedule, and judgment conclusion",
    )
    add_semantic(
        "commercial__etc_treatment",
        "The memo excludes the conditional $72,000 subcontract adjustment at June 30, resolves the unsupported forecast credit without duplicating PO-2409-117 or its prior billing, and makes the separate cutoff transfer",
        pco_evidence,
        {
            "expected_policy_application": (
                "The PM detail shows a $185,000 project-specific exposure and an equal "
                "forecast credit. PM Commitment Detail shows PO-2409-117 at $33,796.69 "
                "open but still forecasts $46,800 in base subcontract ETC. Invoice 88417 "
                "deducts $13,003.31 of prior billings; accounting identifies that draw as "
                "vendor invoice 88102 / AP-0004287. "
                "In the July 2 contract-file batch, the "
                "PCO-011 agreement and the separate Rogue Valley release identify the same "
                "field-report package; the release carries that PO and the invoice line "
                "items. The signed agreement conditionally allocates $27,000 of the "
                "adjustment to previously billed work through CCS-CM-0629-17 and $45,000 "
                "to the remaining unbilled balance, but delivery of that memo to ARM "
                "Commercial Records did not occur until July 2. Finance therefore adds "
                "only the $138,200 incremental PCO cost, applies neither conditional "
                "credit component at June 30, removes the stale "
                "$13,003.31 prior-billing duplicate and transfers the $27,196.69 unbilled "
                "June portion to incurred cost while retaining "
                "the $6,600 July 2 line in ETC, leaving $682,000 final close ETC."
            ),
            "grading_boundary": (
                "No subjective allocation of the remaining $138,200 among labor, material, "
                "or markup is required. The shared PO, conditional $27,000/$45,000 split, "
                "and July 2 delivery chronology are fixed by cross-referenced source records. "
                "Accept any normal business language that reaches the same cutoff treatment."
            ),
        },
        evidence_scope="PCO issue row, adjustment schedule, and judgment conclusion",
    )
    review_threshold_row = row_for(
        "judgment",
        "issue",
        (
            "Review threshold", "Controller review", "Margin review",
            "Review / posting boundary", "Review and posting boundary",
        ),
    )
    add_semantic(
        "commercial__controller_review",
        "The memo identifies the applicable Controller-review trigger and keeps disposition pending",
        "\n\n".join(
            value for value in (
                recommendation,
                judgment_conclusion,
                metadata_evidence,
                action_like_evidence,
                authored_row_text(review_threshold_row, "issue"),
                pco_evidence,
            ) if value
        ).strip(),
        {
            "expected_trigger": (
                f"The {abs(margin_change_bps):,.2f}-basis-point margin-rate decline "
                "exceeds the policy's 100-basis-point quantitative threshold. The low-"
                "confidence estimate and unresolved recovery independently require "
                "Controller review as well."
            ),
            "grading_boundary": (
                "Any one supported independent trigger is sufficient when the memo "
                "keeps Controller disposition pending. A low-confidence estimate or "
                "the remaining unresolved recovery may therefore satisfy this criterion "
                "even when the memo calculates the margin movement incorrectly. The exact "
                "margin rate and basis-point movement are scored by separate criteria and "
                "must not be required again here. Read the complete scoped evidence: "
                "phrases such as submitted for Controller review, Controller review "
                "required or mandatory, Controller disposition blank/open/pending, and "
                "no journal posted establish the pending boundary even when they appear "
                "in different memo sections or an assigned Controller action row. Do not "
                "claim that a visible trigger or pending disposition is absent."
            ),
        },
        evidence_scope=(
            "review-threshold row, judgment conclusion, recommendation, review-status "
            "metadata, posting conclusion, and Controller action evidence"
        ),
    )

    # Proposed entry: May already reversed, so the new entry is the gross June balance.
    posting_rows = canonical_tables["posting"]

    def posting_row(
        entry_aliases: tuple[str, ...],
        account_tokens: tuple[str, ...],
        *,
        debit: float | None = None,
        credit: float | None = None,
    ) -> dict[str, str]:
        entry_wanted = _normalize(entry_aliases[0])
        account_wanted = _normalize(account_tokens[0])
        entry_matches = []
        for row in posting_rows:
            entry = _normalize(row.get("entry"))
            account = _normalize(row.get("account"))
            entry_met = entry == entry_wanted
            account_met = bool(account_wanted) and account_wanted in account
            if entry_met and account_met:
                entry_matches.append(row)
        if len(entry_matches) == 1:
            return entry_matches[0]
        return {}

    proposed_aliases = (
        "Proposed June 30 contribution", "Proposed June contribution",
        "June 30 WIP", "June WIP contribution", "Establish June close",
    )
    proposed_1200 = posting_row(
        proposed_aliases,
        ("1200", "contract asset", "costs and earnings in excess"),
        debit=contract_asset,
    )
    proposed_4300 = posting_row(
        proposed_aliases,
        ("4300", "wip revenue adjustment", "revenue adjustment"),
        credit=contract_asset,
    )
    proposed_1200_met = (
        bool(proposed_1200)
        and _text_contains_number(proposed_1200.get("debit"), contract_asset)
        and _task_001_v24_zero(proposed_1200.get("credit"), allow_blank=True)
    )
    proposed_1200_evidence = anchored_numeric_candidate_evidence(
        ("1200",),
        (contract_asset, 1, .02),
    )
    add_semantic(
        "journal__proposed_1200",
        f"The proposed June project contribution debits account 1200 for ${contract_asset:,.2f}",
        proposed_1200_evidence,
        {
            "entry_status": "proposed June project WIP contribution",
            "account": "1200 Costs and Estimated Earnings in Excess of Billings",
            "debit": contract_asset,
            "credit": 0.0,
            "grading_boundary": (
                "The code gate checks only visible account/value candidates. Decide "
                "whether the amount is a debit on the proposed June WIP contribution, "
                "not another journal or a narrative sensitivity."
            ),
        },
        hard_gate=bool(proposed_1200_evidence),
        evidence_scope="submitted rows containing account 1200 or the exact proposed amount",
        legacy_met=proposed_1200_met,
        always_judge=False,
    )
    proposed_4300_met = (
        bool(proposed_4300)
        and _task_001_v24_zero(proposed_4300.get("debit"), allow_blank=True)
        and _text_contains_number(proposed_4300.get("credit"), contract_asset)
    )
    proposed_4300_evidence = anchored_numeric_candidate_evidence(
        ("4300",),
        (contract_asset, 1, .02),
    )
    add_semantic(
        "journal__proposed_4300",
        f"The proposed June project contribution credits account 4300 for ${contract_asset:,.2f}",
        proposed_4300_evidence,
        {
            "entry_status": "proposed June project WIP contribution",
            "account": "4300 WIP Revenue Adjustment",
            "debit": 0.0,
            "credit": contract_asset,
            "grading_boundary": (
                "The code gate checks only visible account/value candidates. Decide "
                "whether the amount is a credit on the proposed June WIP contribution, "
                "not another journal or a narrative sensitivity."
            ),
        },
        hard_gate=bool(proposed_4300_evidence),
        evidence_scope="submitted rows containing account 4300 or the exact proposed amount",
        legacy_met=proposed_4300_met,
        always_judge=False,
    )
    control_row = row_for(
        "posting",
        "entry",
        (
            "Proposed-entry control", "Entry control", "Control",
            "Debits equal credits", "Project contribution subtotal",
            "Contribution subtotal", "Tie-out control",
        ),
    )
    explicit_control_met = bool(control_row) and all((
        _text_contains_number(control_row.get("debit"), proposed_debits),
        _text_contains_number(control_row.get("credit"), proposed_credits),
    ))
    derived_control_met = bool(proposed_1200 and proposed_4300) and all((
        _text_contains_number(proposed_1200.get("debit"), contract_asset),
        _task_001_v24_zero(proposed_1200.get("credit"), allow_blank=True),
        _task_001_v24_zero(proposed_4300.get("debit"), allow_blank=True),
        _text_contains_number(proposed_4300.get("credit"), contract_asset),
    ))
    proposed_control_evidence = "\n\n".join(
        value for value in (
            proposed_1200_evidence,
            proposed_4300_evidence,
            anchored_row_evidence("1200", "4300"),
        ) if value
    )
    add_semantic(
        "journal__proposed_control",
        f"The proposed June project contribution balances at ${contract_asset:,.2f}",
        proposed_control_evidence,
        {
            "expected_total_debits": proposed_debits,
            "expected_total_credits": proposed_credits,
            "expected_lines": (
                "Dr 1200 and Cr 4300 for the same gross June contract-asset amount"
            ),
            "grading_boundary": (
                "Accept either an explicit subtotal/control row or the two visible "
                "balanced account lines. The code gate does not infer which candidate "
                "rows form the proposed entry; decide that association semantically."
            ),
        },
        category="auditability",
        weight=5,
        hard_gate=bool(proposed_1200_evidence and proposed_4300_evidence),
        evidence_scope="candidate proposed-entry account and control rows",
        legacy_met=explicit_control_met or derived_control_met,
        always_judge=False,
    )
    posting_conclusion = _task_001_v24_box(
        document,
        "Posting conclusion and period-effect bridge",
        "Posting conclusion",
        "Posting status and approval boundary",
        "Posting status, approval boundary, and tie to the May-to-June change",
    )
    posting_evidence = "\n\n".join(
        value for value in (
            proposed_control_evidence,
            anchored_row_evidence("1200", "4300", "REV-WIP-2026-05", "88417"),
            numeric_candidate_evidence((contract_asset, 1, .02)),
            numeric_candidate_evidence((abs(period_effect), 1, .02)),
            narrative_evidence,
            action_like_evidence,
            compact_evidence,
        ) if value
    )
    period_effect_met = (
        _text_contains_number(posting_conclusion, abs(period_effect))
        or _text_contains_number(posting_conclusion, -abs(period_effect))
        or _text_contains_number(row_text(bridge_rows["contract_asset"]), abs(period_effect))
        or _text_contains_number(row_text(bridge_rows["contract_asset"]), -abs(period_effect))
    )
    period_effect_evidence = numeric_candidate_evidence(
        (abs(period_effect), 1, .02)
    )
    add_semantic(
        "journal__period_effect_amount",
        f"The posting support reports the ${abs(period_effect):,.2f} net June WIP revenue effect or equivalent contract-asset movement",
        period_effect_evidence,
        {
            "expected_net_june_effect": period_effect,
            "acceptable_association": (
                "net June WIP revenue effect or the equivalent May-to-June contract-asset movement"
            ),
            "grading_boundary": (
                "The code gate verifies only that the exact amount is visible. Decide "
                "whether it is the period effect rather than the gross June journal or "
                "another unrelated balance."
            ),
        },
        hard_gate=bool(period_effect_evidence),
        evidence_scope="all rows containing the exact net-period-effect amount",
        legacy_met=period_effect_met,
        always_judge=False,
    )
    add_semantic(
        "journal__auto_reversal_and_gross_basis",
        "The memo recognizes that May WIP already auto-reversed and treats the June WIP journal as a cumulative gross re-establishment rather than the month-over-month change",
        posting_evidence,
        {
            "existing_ledger_fact": (
                "REV-WIP-2026-05 is posted on June 1. At the project-support level it "
                "reverses the $274,805.16 May contract asset."
            ),
            "required_posting_logic": (
                "After the posted May reversal, the June 30 entry re-establishes the "
                "memo's cumulative gross project contract-asset balance through Dr 1200 / "
                "Cr 4300. The May-to-June movement is the net June income-statement effect, "
                "not a replacement net-balance journal."
            ),
            "grading_boundary": (
                "This criterion tests reversal awareness and gross-versus-net posting "
                "logic. The exact June gross amount and exact period effect are scored "
                "separately by the proposed-account, control, amount, and bridge criteria; "
                "a memo may satisfy this convention criterion even when those amounts are wrong."
            ),
        },
        evidence_scope="posting table and posting conclusion",
    )
    add_semantic(
        "journal__period_effect_bridge",
        f"The memo correctly connects the posted May reversal, proposed gross June contribution, and ${abs(period_effect):,.2f} May-to-June contract-asset {period_direction}",
        posting_evidence,
        {
            "expected_relationship": (
                f"${contract_asset:,.2f} June contract asset less "
                f"${prior_contract_asset:,.2f} May contract asset equals a "
                f"${abs(period_effect):,.2f} {period_direction}; the same amount is the "
                f"net June {revenue_effect_direction} to WIP revenue adjustment after "
                "the May reversal and June re-establishment."
            ),
        },
        evidence_scope="posting conclusion and contract-asset bridge row",
    )
    add_semantic(
        "journal__posting_boundary",
        "The proposed entry remains unposted, subject to the consolidated WIP schedule tie and documented Controller approval",
        "\n\n".join(
            value for value in (posting_evidence, metadata_evidence) if value
        ),
        {
            "expected_status": journal_gold["status"],
            "required_boundary": (
                "This memo supports a project contribution only. Finance must tie it "
                "into the consolidated June WIP entry, keep the separate gross invoice "
                "cutoff out of the WIP journal, exclude the post-cutoff conditional credit "
                "from June, and obtain "
                "Controller approval before any posting is released."
            ),
            "grading_boundary": (
                "Accept distributed, professionally equivalent release-control wording. "
                "A memo satisfies this boundary when it states that it supports only the "
                "project contribution, that the proposed lines are draft and unposted, "
                "that they tie to the supporting close schedule and post only inside the "
                "consolidated June WIP journal, and that documented Controller disposition "
                "is required before posting. Do not require the literal sentence 'Finance "
                "must first tie it.' Correctness of the proposed amount, period-effect "
                "bridge, and separate cutoff accounting is scored by the neighboring "
                "journal criteria and must not erase an otherwise explicit posting boundary."
            ),
            "distributed_equivalence_examples": (
                "Taken together, wording that this is a single-project contribution only, that the "
                "consolidated entry remains a separate Controller action, that no journal has been "
                "created or posted and nothing authorizes posting, and an owned completion step whose "
                "evidence is an updated WIP schedule tie-out plus Controller review satisfies this "
                "criterion. Do not require one sentence to repeat the sequence, the word documented, "
                "or the literal phrase before release when those controls are unambiguously distributed "
                "across the scoped recommendation, posting conclusion, status, and action evidence."
            ),
        },
        evidence_scope="posting conclusion, recommendation, and review status",
    )
    add_semantic(
        "journal__separate_ap_accrual_boundary",
        "The posting support keeps the $27,196.69 current-invoice cutoff separate from the gross WIP journal, excludes the conditional $27,000 credit from June, and requires a retie before release",
        "\n\n".join(
            value for value in (posting_evidence, cutoff_evidence) if value
        ),
        {
            "required_boundary": (
                "The proposed WIP entry contains only Dr 1200 / Cr 4300 for the "
                f"${contract_asset:,.2f} gross contract asset. Invoice 88417 requires "
                "a separate $27,196.69 cutoff for unbilled June work; its $13,003.31 prior "
                "draw is already posted, while its $6,600 July 2 line remains in ETC. "
                "CCS-CM-0629-17 was delivered after the June 30 effectiveness cutoff, so "
                "its conditional $27,000 billed-work component and $45,000 remaining-balance "
                "component have no June entry. The invoice cutoff belongs in the close cost "
                "basis, its matching commitment leaves ETC, and it does not belong inside "
                "the WIP journal. Project Accounting "
                "must retie before release."
            ),
            "grading_boundary": (
                "Accept equivalent journal-boundary and release-control wording. No "
                "particular AP expense or liability account number is required. If the "
                "submission visibly provides separate invoice and conditional-credit rows, "
                "acknowledge that work and identify the actual unmet timing, amount, ETC, WIP, or retie conjunct "
                "rather than claiming a present row is absent. Wording that the entry must "
                "be recomputed or tied to the final consolidated schedule before posting "
                "satisfies the retie conjunct; do not require the literal word 'retie'. "
                "If the submission visibly identifies CCS-CM-0629-17 and its $27,000 "
                "amount, acknowledge that identification. Treating it as a future-period "
                "item because of the July 2 delivery is correct; describe a June current-cost "
                "or ETC reduction as premature rather than saying the credit is unidentified."
            ),
        },
        category="auditability",
        weight=10,
        evidence_scope=(
            "posting table, posting conclusion, cutoff evidence, and recommendation"
        ),
    )

    # Source authority. Names, versions, and uses are semantic so equivalent
    # professional labels remain acceptable.
    source_specs = (
        (
            "accounting",
            (
                "Company accounting records", "Current accounting records",
                "Accounting system", "Vista ERP", "Posted records",
            ),
            ("accounting system", "vista erp", "posted records"),
            "Posted system records as of June 30 for contract, cost, billings, and transaction status.",
        ),
        (
            "may_close",
            (
                "WIP 5.31.26_FINAL_v7_revised NB.xlsx", "Final May WIP",
                "May final WIP schedule", "May WIP schedule", "May close",
            ),
            ("5.31.26", "May final WIP", "May close"),
            "The final May v7 workpaper supplies the prior-close WIP position; its dump tabs are not June actuals.",
        ),
        (
            "pm_forecast",
            (
                "PM ETC updates_6.29 530pm_COMBINED_v3.xlsx", "PM forecast",
                "June PM submission", "PM June ETC submission",
                "June ETC submission",
            ),
            ("PM ETC", "PM forecast", "June ETC submission"),
            (
                "The June 29 5:30 PM file supplies the $584,000 forecast and its "
                "$185,000 exposure/credit build. Its Commitment Detail sheet shows "
                "PO-2409-117 / Rogue Valley / freezer-slab rework at $33,796.69 open "
                "but $46,800 forecast, leaving the already-posted prior draw duplicated "
                "inside the base subcontract forecast."
            ),
        ),
        (
            "commercial_working",
            (
                "PCO-011 commercial records", "CO master and PCO-011 draft packet", "Commercial Log",
                "Commercial change-order log", "PCO-011 commercial backup",
                "CO log MASTER", "PCO-011", "CO master",
            ),
            (
                "PCO-011", "CO log", "Commercial", "draft packet",
                "backcharge backup",
            ),
            (
                "The master and draft packet establish the full working position and "
                "$185,000 chronology, but do not themselves authorize an accounting "
                "recovery or customer-contract increase."
            ),
        ),
        (
            "executed_credit",
            (
                "Contract-records scan and correspondence archive",
                "scan batch 20260702-01.pdf",
                "sent 1706.eml",
                "Signed cost participation", "Concrete correspondence",
                "Executed subcontract credit",
            ),
            (
                "contract records", "correspondence archive", "signed letter",
                "cost participation", "executed credit",
            ),
            (
                "The June 29 correspondence contains both parties' signatures and the "
                "separate routing record establishes release and executed return by June "
                "30. The agreement also conditions effectiveness on delivery of "
                "CCS-CM-0629-17 to ARM Commercial Records. The memo's ordinary receipt "
                "field records July 2 at 07:41 PT and says it has no independent effect "
                "before the agreement conditions are satisfied. The signed records "
                "conditionally allocate $45,000 to the remaining unbilled balance and "
                "$27,000 against previously billed work, but support no June credit and "
                "no customer-contract amendment."
            ),
        ),
        (
            "policy",
            (
                "FIN-REV-04 signed WIP policy", "FIN-REV-04", "Signed WIP policy",
                "Signed accounting / WIP policy", "Signed accounting policy",
                "WIP policy", "Policy / status crosswalk",
            ),
            ("FIN-REV-04", "WIP policy", "status crosswalk"),
            "Signed FIN-REV-04 controls contract value, ETC, gross presentation, auto reversal, and review thresholds.",
        ),
        (
            "late_invoice",
            (
                "AP intake queue and invoice scan batch p. 4", "Invoice 88417",
                "batch 20260701-02 page 4", "Rogue Valley invoice", "Late vendor invoice", "Cutoff invoice",
            ),
            ("88417", "Rogue Valley", "late vendor invoice", "cutoff invoice"),
            (
                "Invoice 88417 and its AP intake row supply $46,800 gross billings, "
                "$13,003.31 of prior billings, the $33,796.69 current amount due, June 29 "
                "invoice date, July 1 receipt, PO-2409-117, vendor, and project. Accounting "
                "records identify that prior draw as vendor invoice 88102 / AP-0004287 and prove the current invoice "
                "is unposted. Its field-ticket lines and the matching vendor release allocate "
                "$40,200 gross work to June and $6,600 to July 2. The PM Commitment Detail "
                "row shows the open amount but still forecasts the full package."
            ),
        ),
    )
    # Source authority and use are semantic questions. Route the complete,
    # compact source record plus the relevant workpaper section to the judge;
    # row-label aliases may help isolate a row, but they must never decide
    # whether an otherwise authored answer reaches semantic review.
    source_table_evidence = "\n\n".join(
        value for value in (four_column_evidence, compact_evidence) if value
    )
    source_use_evidence = {
        "accounting": "\n".join([
            five_column_evidence,
            population_table_evidence,
        ]).strip(),
        "may_close": movement_evidence,
        "pm_forecast": "\n".join([
            central_evidence,
            adjustment_evidence,
            cutoff_evidence,
        ]).strip(),
        "commercial_working": pco_evidence,
        "executed_credit": pco_evidence,
        "policy": "\n".join([
            pco_evidence,
            cutoff_evidence,
            recommendation,
            posting_conclusion,
        ]).strip(),
        "late_invoice": cutoff_evidence,
    }
    source_rows: dict[str, dict[str, str]] = {}
    for key, aliases, matching_tokens, expected_use in source_specs:
        row = row_for("source", "source", aliases)
        source_rows[key] = row
        submitted_source_evidence = "\n".join([
            anchored_row_evidence(*matching_tokens),
            source_table_evidence,
            source_use_evidence[key],
        ]).strip()
        source_context: dict[str, Any] = {
            "expected_source_use": expected_use,
            "grading_boundary": (
                "Accept normal abbreviations, file-name shortening, and equivalent "
                "professional source labels. Judge the source's authority and use. "
                "Do not award credit from a correct number or conclusion alone: "
                "the authored evidence must identify the relevant source and connect "
                "it to the workpaper's use."
            ),
        }
        source_scope = (
            "complete source record plus the criterion-relevant workpaper section"
        )
        if key == "late_invoice":
            source_context["distributed_evidence_boundary"] = (
                "The association may be established across ordinary source, population, "
                "forecast, and workpaper sections. Identifying the PM forecast file and "
                "its open PO-2409-117 commitment together with the invoice/AP evidence, "
                "vendor, project, cumulative gross/prior/current amounts, the prior accounting "
                "match, line-level June/July timing, matching release, and current-invoice "
                "ledger absence is sufficient; do not require "
                "the literal sheet label 'Commitment Detail' or all facts in one row. "
                "Source traceability is separate from whether the memo ultimately chooses "
                "the correct cutoff accounting treatment, which is scored elsewhere."
            )
            source_scope = (
                "complete source record plus the late-invoice, open-commitment, "
                "population, and cutoff workpaper evidence"
            )
        if key == "pm_forecast":
            source_context["criterion_isolation_boundary"] = (
                "A MET verdict requires the authored memo to connect the PM source's "
                "$33,796.69 open commitment with its $46,800 forecast amount and "
                "identify the already-posted $13,003.31 prior draw as remaining inside "
                "that forecast. Merely naming the PM file, citing the $584,000 headline "
                "ETC, or describing the separate $185,000 exposure/recovery build is "
                "insufficient. Accept distributed professional wording and do not "
                "require the words duplicate, gross, Commitment Detail, or prior draw. "
                "The downstream correction amount and final recommendation are scored "
                "separately and must not decide this source-use criterion. For a MET "
                "verdict, the explanation must identify the visible open-versus-forecast "
                "amount association rather than citing only the headline PM baseline."
            )
        if key == "executed_credit":
            source_context["criterion_isolation_boundary"] = (
                "This criterion is only source identification, authority, and the documented "
                "conditional $27,000 billed-work / $45,000 remaining-balance split. Award it when the "
                "authored source record and workpaper connect the signed correspondence, "
                "the separate package-release/return routing evidence, the credit-memo "
                "delivery record, and those split facts to a $72,000 subcontract-only "
                "adjustment that was not effective at June 30. Both signatures or the "
                "scan batch alone are insufficient under the document's operative conditions. "
                "Fail when the submission leaves package release, executed return, or credit-"
                "memo delivery unevidenced, or uses the adjustment in June despite the July 2 receipt. "
                "Downstream classification and arithmetic are scored separately, so an ETC "
                "calculation may still misapply one component without changing this source-use "
                "decision only when that downstream error does not contradict the source's "
                "June 30 effectiveness conclusion. For a MET verdict, explicitly assess the "
                "submission's stated Finance treatment; identifying the July 2 condition while "
                "calling it administrative and recognizing either credit component in June is "
                "UNMET. Do not claim correspondence, routing evidence, "
                "or the split is absent when it is visibly present in the supplied evidence."
            )
        add_semantic(
            f"source__{key}",
            f"The source record identifies and correctly uses the controlling {key.replace('_', ' ')} evidence",
            submitted_source_evidence,
            source_context,
            category="provenance",
            weight=3,
            hard_gate=bool(submitted_source_evidence),
            evidence_scope=source_scope,
        )

    action_specs = (
        (
            "ap_cutoff",
            (
                "AP / Project Accounting", "AP and Project Accounting",
                "Accounts Payable / Project Accounting", "Accounts Payable",
            ),
            (
                "Match prior invoice 88102 to AP-0004287, approve or post the separate $27,196.69 "
                "current-invoice cutoff, exclude the conditional $27,000 credit from June, remove the stale prior "
                "billing and current June transfer from PM subcontract ETC, retain the $6,600 July line, "
                "and retie the close cost basis before WIP release."
            ),
        ),
        (
            "finance",
            ("Finance / Project Accounting", "Finance", "Project Accounting"),
            (
                "Use the net $98,000 ETC correction, carry the $301,726.52 gross project "
                "contribution into the consolidated June WIP entry, keep the AP cutoff "
                "outside that journal, complete the schedule-to-entry tie, and "
                "retain the memo."
            ),
        ),
        (
            "commercial",
            ("Commercial / Project Team", "Commercial", "Project Team"),
            "Document the July 2 effectiveness, apply the $27,000 billed-work credit and $45,000 remaining-balance reduction subsequently, then pursue support for the remaining $113,000 before any additional recovery treatment without changing customer contract value.",
        ),
        (
            "controller",
            ("Controller", "Corporate Controller"),
            (
                "Review the invoice cutoff, conditional-credit timing, ETC transfer, "
                "commercial correction, margin movement, and consolidated "
                "WIP entry; document approve or return disposition before any posting."
            ),
        ),
    )
    action_table_evidence = "\n\n".join(
        value for value in (action_like_evidence, compact_evidence) if value
    )
    reviewer_disposition_evidence = raw_table_scope(columns={4}, max_rows=3)
    for key, aliases, expected_action in action_specs:
        row = row_for("action", "owner", aliases)
        action_context: dict[str, Any] = {
            "expected_action": expected_action,
            "grading_boundary": (
                "Accept a professionally equivalent action, owner label, timing, and "
                "completion-evidence description; do not require reference wording. "
                "A generic action or another workstream's row is not sufficient."
            ),
        }
        action_evidence = action_table_evidence
        action_scope = (
            "complete action table; judge the professionally equivalent "
            "owner/action row for this criterion"
        )
        if key == "controller":
            action_context["preparer_boundary"] = (
                "This is a future reviewer action. Open, pending, or not-started is the "
                "correct status in a preparer memorandum; do not require completed "
                "Controller sign-off. A completed or simulated approval would violate "
                "the task boundary."
            )
            action_context["acceptable_equivalent"] = (
                "A direction to review this project support and issue a documented "
                "Controller disposition before the consolidated WIP entry is posted is "
                "sufficient. Evidence may be distributed across owner/clearance rows, "
                "the posting boundary, and a blank Controller disposition block. A "
                "dedicated Controller row in the action table is not required when those "
                "sections together assign future Controller review, its timing or release "
                "condition, and a documented approve/return disposition; no analysis "
                "section must be repeated by name."
            )
            action_evidence = "\n\n".join([
                reviewer_disposition_evidence,
                anchored_row_evidence("Controller"),
                posting_conclusion,
                metadata_evidence,
                action_like_evidence,
            ]).strip()
            action_scope = (
                "complete action, owner/clearance, posting-boundary, and "
                "Controller-disposition evidence"
            )
        add_semantic(
            f"action__{key}",
            f"The {aliases[0]} action is specific, assigned, and consistent with the preparer boundary",
            "\n".join([
                authored_row_text(row, "owner"),
                action_evidence,
            ]).strip(),
            action_context,
            category="auditability",
            weight=5,
            hard_gate=bool(action_table_evidence),
            evidence_scope=action_scope,
        )

    consistency_sections = (
        ("review status", metadata_evidence, 700),
        ("recommendation", central_evidence, 1_800),
        ("June close", five_column_evidence, 2_800),
        ("May-to-June bridge", movement_evidence, 1_400),
        ("population controls", population_evidence, 1_600),
        ("cutoff document treatment", cutoff_evidence, 1_800),
        ("commercial judgment", pco_evidence, 1_600),
        (
            "proposed posting",
            posting_evidence,
            2_000,
        ),
    )
    consistency_evidence = "\n\n".join(
        f"{label}:\n{value[:limit]}"
        for label, value, limit in consistency_sections
        if str(value or "").strip()
    )
    add_semantic(
        "artifact__internal_consistency",
        "The completed memo is review-ready and internally consistent across its recommendation, calculations, population controls, judgment, sources, and proposed entry",
        consistency_evidence,
        {
            "required_quality": (
                f"No material contradiction among the ${close_etc:,.2f} ETC, "
                f"${float(project['estimated_cost_at_completion']):,.2f} EAC, "
                f"${contract_asset:,.2f} contract asset, "
                f"${estimated_margin:,.2f} estimated margin, unposted "
                "status, gross June posting logic, the separate $27,196.69 current-invoice "
                "cutoff, no June recognition of the conditional credit, the $13,003.31 prior draw already posted, the "
                "$6,600 post-cutoff line retained in ETC, and the invoice's cumulative-billing "
                "reconciliation. Neither the $27,000 billed-work component nor the $45,000 "
                "unbilled-balance component may reduce June cost because the delivery "
                "condition was satisfied July 2. The PO cost must not be omitted, retained in ETC after accrual, "
                "duplicated through the PCO overlay, or inserted into the WIP journal. "
                "Unsupported certainty or an invented approval is "
                "not review-ready."
            ),
            "grading_boundary": (
                "This is a holistic consistency check, not a formatting-style test. "
                "Harmless layout, label, ordering, unit, and sign-presentation changes "
                "must not reduce credit. The explanation must name an actual material "
                "contradiction visible in the submitted evidence. If the conditional credit "
                "split and timing are correct, acknowledge them and identify a different "
                "inconsistency rather than claiming they are missing. Distinguish the $46,800 gross "
                "billings or PM forecast, $13,003.31 posted prior draw, $33,796.69 current "
                "amount due, $40,200 gross June work, and $27,196.69 current June accrual. "
                "Do not claim the full invoice or amount due was accrued merely because one "
                "of those values is cited. If the primary close omits the current accrual or "
                "presents it only as a sensitivity, identify that actual contradiction instead."
            ),
        },
        category="auditability",
        weight=5,
        hard_gate=bool(consistency_evidence.strip()),
        evidence_scope=(
            "compact structured record of the recommendation, calculation, bridge, "
            "population, judgment, and posting sections"
        ),
    )

    result = _result(criteria)
    result["artifact_review"] = _task_001_v24_artifact_review(
        document, path, parse_error
    )
    result = _attach_semantic_review(
        result,
        task_id="task_001",
        evidence=_legacy_artifact_evidence(path) if document is not None else "",
        artifact_type="controller sign-off memorandum",
        specs=semantic_specs,
        decision_failure_cap=None,
        always_judge=True,
        execution_mode="scoped_per_criterion",
    )
    result["task_grading_revision"] = dict(TASK_GRADING_REVISIONS["task_001"])
    return result


FILE_GRADERS: dict[str, Callable[[Path], dict[str, Any]]] = {
    "task_004": _grade_task_004,
    "task_015": _grade_task_015,
}


def grade_apex_task(task_id: str, answer: Any, workspace_root: str | Path) -> dict[str, Any]:
    if task_id not in SAMPLE_TASK_IDS:
        raise KeyError(f"Task is not part of this sample: {task_id}")
    if task_id in {"task_035", "task_068"}:
        from runtime.grading.corporate_finance import grade_corporate_finance_task
        result = grade_corporate_finance_task(task_id, answer, workspace_root)
    elif task_id == "task_001":
        result = _grade_task_001_v27(Path(workspace_root), answer)
    elif task_id in FILE_GRADERS:
        result = FILE_GRADERS[task_id](Path(workspace_root))
        result = _canonicalize_legacy_file_policy(task_id, result)
    else:
        raise AssertionError(f"No sample grader registered for {task_id}")
    if task_id in TASK_GRADING_REVISIONS:
        result.setdefault(
            "task_grading_revision", dict(TASK_GRADING_REVISIONS[task_id])
        )
    return apply_reward_policy(task_id, attach_default_policy(result))
