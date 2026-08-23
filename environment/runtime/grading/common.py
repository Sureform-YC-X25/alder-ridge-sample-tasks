from __future__ import annotations

import json
import math
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


GOLD = Path(__file__).resolve().parent / "gold"


def load_gold(name: str) -> dict[str, Any]:
    return json.loads((GOLD / name).read_text(encoding="utf-8"))


def normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def score_close(actual: float | None, expected: float, *, abs_tol: float, rel_tol: float) -> float:
    if actual is None or not math.isfinite(actual):
        return 0.0
    tol = max(abs_tol, abs(expected) * rel_tol)
    error = abs(actual - expected)
    if error <= tol:
        return 1.0
    if error >= tol * 6:
        return 0.0
    return max(0.0, 1.0 - (error - tol) / (tol * 5))


def number(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str) and not value.startswith("="):
        cleaned = value.strip().replace("$", "").replace(",", "")
        negative = cleaned.startswith("(") and cleaned.endswith(")")
        cleaned = cleaned.strip("()")
        try:
            result = float(cleaned)
            return -result if negative else result
        except ValueError:
            return None
    return None


def find_header_row(ws, required_groups: list[set[str]], max_rows: int = 30) -> tuple[int, dict[str, int]] | None:
    for row in range(1, min(ws.max_row, max_rows) + 1):
        mapping = {normalize(ws.cell(row, col).value): col for col in range(1, ws.max_column + 1) if ws.cell(row, col).value is not None}
        if all(any(candidate in mapping for candidate in group) for group in required_groups):
            return row, mapping
    return None


def find_column(mapping: dict[str, int], candidates: set[str]) -> int | None:
    for candidate in candidates:
        if candidate in mapping:
            return mapping[candidate]
    return None


def docx_text(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def pptx_text(path: Path) -> tuple[str, int]:
    inspection = inspect_pptx(path)
    return "\n".join(inspection["slide_texts"]), inspection["slide_count"]


def inspect_pptx(path: Path) -> dict[str, Any]:
    """Open a real OOXML presentation and collect per-slide structure.

    Parsing slide XML from any ZIP is insufficient: a few XML fragments can
    impersonate a deck without being openable by PowerPoint or LibreOffice.
    ``python-pptx`` validates the package relationships and presentation part.
    """

    presentation = Presentation(str(path))
    slide_texts: list[str] = []
    shape_counts: list[int] = []
    visual_counts: list[int] = []
    for slide in presentation.slides:
        chunks: list[str] = []
        visuals = 0
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                chunks.append(text)
            if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
                visuals += 1
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                visuals += 1
        slide_texts.append("\n".join(chunks))
        shape_counts.append(len(slide.shapes))
        visual_counts.append(visuals)
    return {
        "slide_count": len(presentation.slides),
        "slide_texts": slide_texts,
        "shape_counts": shape_counts,
        "visual_counts": visual_counts,
        "slide_width": int(presentation.slide_width),
        "slide_height": int(presentation.slide_height),
    }


def inspect_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    return {
        "page_count": len(reader.pages),
        "page_texts": page_texts,
        "nonblank_ratio": (
            sum(len(normalize(text).split()) >= 12 for text in page_texts) / len(page_texts)
            if page_texts else 0.0
        ),
    }


def token_overlap(left: str, right: str) -> float:
    left_tokens = {token for token in normalize(left).split() if len(token) >= 3}
    right_tokens = {token for token in normalize(right).split() if len(token) >= 3}
    if not left_tokens or not right_tokens:
        return 0.0
    return len(left_tokens & right_tokens) / min(len(left_tokens), len(right_tokens))


def formula_ratio(ws, rows: list[int], columns: list[int]) -> float:
    cells = [ws.cell(r, c).value for r in rows for c in columns]
    if not cells:
        return 0.0
    return sum(isinstance(v, str) and v.startswith("=") for v in cells) / len(cells)


def find_sheet(wb, candidates: set[str]):
    normalized = {normalize(name): name for name in wb.sheetnames}
    for candidate in candidates:
        if candidate in normalized:
            return wb[normalized[candidate]]
    for norm, name in normalized.items():
        if any(candidate in norm for candidate in candidates):
            return wb[name]
    return None


def weighted(parts: dict[str, tuple[float, float]]) -> tuple[float, dict[str, float]]:
    detail = {name: max(0.0, min(1.0, score)) for name, (score, weight) in parts.items()}
    total_weight = sum(weight for score, weight in parts.values())
    reward = sum(detail[name] * parts[name][1] for name in parts) / total_weight if total_weight else 0.0
    return round(reward, 6), detail
