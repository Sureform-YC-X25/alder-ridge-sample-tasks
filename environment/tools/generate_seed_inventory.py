from __future__ import annotations

import csv
import email
import json
import math
import re
import sqlite3
import statistics
import subprocess
import tempfile
from collections import Counter, defaultdict
from email import policy
from pathlib import Path
from urllib.parse import quote

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
SEED_ROOT = ROOT / "environment" / "seed"
SOURCE_ROOT = SEED_ROOT / "sources"
OUTPUT = ROOT / "SEED_DATA_INVENTORY.md"
REGISTRY = SEED_ROOT / "controls" / "file_registry.json"


def human_size(byte_count: int) -> str:
    value = float(byte_count)
    for unit in ("B", "KiB", "MiB", "GiB"):
        if value < 1024 or unit == "GiB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.2f} {unit}"
        value /= 1024
    raise AssertionError("unreachable")


def relative_link(path: Path) -> str:
    relative = path.relative_to(ROOT).as_posix()
    return quote(relative, safe="/")


def stat_line(values: list[int]) -> str:
    return (
        f"total {sum(values):,} / average {statistics.mean(values):,.2f} / "
        f"median {statistics.median(values):,} / range {min(values):,}-{max(values):,}"
    )


def detected_list_records(worksheet) -> int:
    """Count the strongest list-like range on one worksheet.

    A list-like range has a primarily textual header spanning at least three
    populated columns and at least eight contiguous populated rows below it.
    This keeps titles and summary blocks out while providing a reproducible
    approximation for schedule/list volume in heterogeneous finance workbooks.
    """

    row_values: dict[int, list[object]] = {}
    for row in worksheet.iter_rows():
        values = [cell.value for cell in row]
        if any(value is not None for value in values):
            row_values[row[0].row] = values

    best = 0
    for row_number, values in row_values.items():
        populated = [(index, value) for index, value in enumerate(values) if value is not None]
        if len(populated) < 3:
            continue
        text_count = sum(
            isinstance(value, str) and not value.startswith("=")
            for _, value in populated
        )
        if text_count < 2 or text_count / len(populated) < 0.5:
            continue

        first_column = min(index for index, _ in populated)
        last_column = max(index for index, _ in populated)
        required_values = max(2, math.ceil(len(populated) * 0.4))
        records = 0
        consecutive_gaps = 0
        for data_row in range(row_number + 1, worksheet.max_row + 1):
            candidate = row_values.get(data_row, [])
            populated_count = sum(
                1
                for index in range(first_column, min(last_column + 1, len(candidate)))
                if candidate[index] is not None
            )
            if populated_count >= required_values:
                records += 1
                consecutive_gaps = 0
            else:
                consecutive_gaps += 1
                if consecutive_gaps >= 2:
                    break
        if records >= 8:
            best = max(best, records)
    return best


def workbook_metrics(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    populated_rows = populated_cells = formula_cells = list_records = 0
    for worksheet in workbook.worksheets:
        sheet_rows = 0
        for row in worksheet.iter_rows():
            row_populated = False
            for cell in row:
                if cell.value is None:
                    continue
                row_populated = True
                populated_cells += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells += 1
            sheet_rows += int(row_populated)
        populated_rows += sheet_rows
        list_records += detected_list_records(worksheet)
    return {
        "sheets": len(workbook.worksheets),
        "populated_rows": populated_rows,
        "populated_cells": populated_cells,
        "formula_cells": formula_cells,
        "list_records": list_records,
    }


def email_metrics(path: Path) -> dict[str, int]:
    with path.open("rb") as stream:
        message = email.message_from_binary_file(stream, policy=policy.default)
    body_parts: list[str] = []
    attachments = 0
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_disposition() == "attachment":
                attachments += 1
            elif part.get_content_type() == "text/plain":
                try:
                    body_parts.append(str(part.get_content()))
                except (LookupError, UnicodeError):
                    continue
    else:
        try:
            body_parts.append(str(message.get_content()))
        except (LookupError, UnicodeError):
            pass
    words = len(re.findall(r"\b[\w'-]+\b", "\n".join(body_parts)))
    return {"words": words, "attachments": attachments}


def render_docx_page_counts(paths: list[Path]) -> dict[Path, int | None]:
    counts: dict[Path, int | None] = {}
    with tempfile.TemporaryDirectory(prefix="alder-seed-docx-pages-") as temporary:
        temp_root = Path(temporary)
        for index, path in enumerate(paths):
            output = temp_root / f"doc-{index:02d}"
            output.mkdir()
            completed = subprocess.run(
                [
                    "soffice",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output),
                    str(path),
                ],
                capture_output=True,
                check=False,
                timeout=60,
            )
            rendered = list(output.glob("*.pdf"))
            counts[path] = (
                len(PdfReader(str(rendered[0])).pages)
                if completed.returncode == 0 and rendered
                else None
            )
    return counts


def source_metrics() -> tuple[list[dict[str, object]], dict[str, list[int]]]:
    registry = json.loads(REGISTRY.read_text(encoding="utf-8"))
    status_by_path = {
        row["path"]: row["version_status"]
        for row in registry["files"]
    }
    docx_paths = sorted(SOURCE_ROOT.rglob("*.docx"))
    docx_pages = render_docx_page_counts(docx_paths)
    rows: list[dict[str, object]] = []
    aggregates: dict[str, list[int]] = defaultdict(list)

    for path in sorted(candidate for candidate in SOURCE_ROOT.rglob("*") if candidate.is_file()):
        relative = path.relative_to(SOURCE_ROOT).as_posix()
        extension = path.suffix.lower().lstrip(".")
        metrics: dict[str, object] = {}
        if extension == "pdf":
            metrics["pages"] = len(PdfReader(str(path)).pages)
            aggregates["pdf_pages"].append(int(metrics["pages"]))
        elif extension == "docx":
            document = Document(str(path))
            metrics.update(
                pages=docx_pages[path],
                tables=len(document.tables),
                table_rows=sum(len(table.rows) for table in document.tables),
            )
            if metrics["pages"] is not None:
                aggregates["docx_pages"].append(int(metrics["pages"]))
        elif extension == "pptx":
            presentation = Presentation(str(path))
            metrics["slides"] = len(presentation.slides)
            aggregates["pptx_slides"].append(int(metrics["slides"]))
        elif extension == "xlsx":
            metrics.update(workbook_metrics(path))
            aggregates["xlsx_sheets"].append(int(metrics["sheets"]))
            aggregates["xlsx_rows"].append(int(metrics["populated_rows"]))
            if int(metrics["list_records"]) > 0:
                aggregates["xlsx_list_records"].append(int(metrics["list_records"]))
        elif extension == "csv":
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
                content = list(csv.reader(stream))
            metrics["records"] = sum(
                1 for row in content[1:] if any(str(value).strip() for value in row)
            )
            metrics["columns"] = len(content[0]) if content else 0
            aggregates["csv_records"].append(int(metrics["records"]))
        elif extension == "eml":
            metrics.update(email_metrics(path))
            aggregates["email_words"].append(int(metrics["words"]))
        elif extension == "txt":
            text = path.read_text(encoding="utf-8", errors="replace")
            metrics["lines"] = len(text.splitlines())

        rows.append(
            {
                "path": path,
                "relative": relative,
                "extension": extension,
                "status": status_by_path.get(relative, "unclassified"),
                "bytes": path.stat().st_size,
                "metrics": metrics,
            }
        )
    return rows, aggregates


def metric_description(row: dict[str, object]) -> str:
    extension = str(row["extension"])
    metrics = dict(row["metrics"])
    parts = [extension.upper(), human_size(int(row["bytes"])), str(row["status"])]
    if extension == "pdf":
        parts.append(f"{metrics['pages']:,} pages")
    elif extension == "docx":
        pages = metrics.get("pages")
        parts.append(f"{pages:,} rendered pages" if pages is not None else "page count unavailable")
        parts.append(f"{metrics['tables']:,} tables")
        parts.append(f"{metrics['table_rows']:,} table rows")
    elif extension == "pptx":
        parts.append(f"{metrics['slides']:,} slides")
    elif extension == "xlsx":
        parts.extend(
            [
                f"{metrics['sheets']:,} sheets",
                f"{metrics['populated_rows']:,} populated rows",
                f"{metrics['formula_cells']:,} formulas",
            ]
        )
        if int(metrics["list_records"]) > 0:
            parts.append(f"{metrics['list_records']:,} detected list records")
    elif extension == "csv":
        parts.append(f"{metrics['records']:,} records")
        parts.append(f"{metrics['columns']:,} columns")
    elif extension == "eml":
        parts.append(f"{metrics['words']:,} body words")
    elif extension == "txt":
        parts.append(f"{metrics['lines']:,} lines")
    return " · ".join(parts)


def database_rows() -> list[tuple[str, int]]:
    connection = sqlite3.connect(SEED_ROOT / "accounting.db")
    try:
        tables = [
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master "
                "WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            )
        ]
        return [
            (table, connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0])
            for table in tables
        ]
    finally:
        connection.close()


def build_markdown() -> str:
    sources, aggregates = source_metrics()
    source_counts = Counter(str(row["extension"]) for row in sources)
    status_counts = Counter(str(row["status"]) for row in sources)
    database = database_rows()
    seed_files = [path for path in SEED_ROOT.rglob("*") if path.is_file()]
    seed_bytes = sum(path.stat().st_size for path in seed_files)
    source_bytes = sum(int(row["bytes"]) for row in sources)
    db_bytes = (SEED_ROOT / "accounting.db").stat().st_size

    lines = [
        "# Alder Ridge Seed Data Inventory",
        "",
        "This catalog links directly to every file included in the complete Alder Ridge sample seed. "
        "It covers the full 132-file shared company world, the accounting database, and the seed control/support files.",
        "",
        "## Summary",
        "",
        "| Measure | Volume |",
        "|---|---:|",
        f"| Company-world source files | {len(sources):,} |",
        f"| Complete seed files, including controls/database | {len(seed_files):,} |",
        f"| Complete seed size | {human_size(seed_bytes)} |",
        f"| Source-artifact size | {human_size(source_bytes)} |",
        f"| Accounting database size | {human_size(db_bytes)} |",
        f"| Accounting database rows | {sum(count for _, count in database):,} |",
        f"| Excel workbooks / worksheets | {source_counts['xlsx']:,} / {sum(aggregates['xlsx_sheets']):,} |",
        f"| Detected Excel list records | {sum(aggregates['xlsx_list_records']):,} |",
        f"| PDF files / pages | {source_counts['pdf']:,} / {sum(aggregates['pdf_pages']):,} |",
        f"| Word files / rendered pages | {source_counts['docx']:,} / {sum(aggregates['docx_pages']):,} |",
        f"| PowerPoint files / slides | {source_counts['pptx']:,} / {sum(aggregates['pptx_slides']):,} |",
        f"| CSV files / records | {source_counts['csv']:,} / {sum(aggregates['csv_records']):,} |",
        "",
        "### File types",
        "",
        "| Type | Files |",
        "|---|---:|",
    ]
    for extension, count in sorted(source_counts.items(), key=lambda item: (-item[1], item[0])):
        lines.append(f"| `{extension}` | {count:,} |")

    lines.extend(["", "### Source authority/version status", "", "| Status | Files |", "|---|---:|"])
    for status, count in status_counts.most_common():
        lines.append(f"| {status} | {count:,} |")

    lines.extend(
        [
            "",
            "### Native-format volume",
            "",
            f"- PDF pages: {stat_line(aggregates['pdf_pages'])}",
            f"- Word rendered pages: {stat_line(aggregates['docx_pages'])}",
            f"- PowerPoint slides: {stat_line(aggregates['pptx_slides'])}",
            f"- Excel worksheets: {stat_line(aggregates['xlsx_sheets'])}",
            f"- Excel populated rows: {stat_line(aggregates['xlsx_rows'])}",
            f"- Excel detected list records per list-containing workbook: {stat_line(aggregates['xlsx_list_records'])}",
            f"- CSV records: {stat_line(aggregates['csv_records'])}",
            f"- Email body words: {stat_line(aggregates['email_words'])}",
            "",
            "Excel list records use a structural estimate: a mainly textual header spanning at least three columns followed by at least eight contiguous populated data rows. Model output rows, repeated versions, and reconciled copies can represent the same underlying business fact, so these are workload-volume measures rather than unique-event counts.",
            "",
            "## Accounting database",
            "",
            f"[`accounting.db`]({relative_link(SEED_ROOT / 'accounting.db')}) contains **{sum(count for _, count in database):,} rows across {len(database)} business tables**.",
            "",
            "| Table | Rows |",
            "|---|---:|",
        ]
    )
    for table, count in database:
        lines.append(f"| `{table}` | {count:,} |")
    lines.append(f"| **Total** | **{sum(count for _, count in database):,}** |")

    lines.extend(
        [
            "",
            "## Seed documentation and controls",
            "",
            "These files document or verify the seed package. Control files are not mounted into the agent-visible workspace.",
            "",
        ]
    )
    support_paths = sorted(
        path for path in SEED_ROOT.iterdir() if path.is_file()
    ) + sorted((SEED_ROOT / "controls").glob("*.json"))
    for path in support_paths:
        lines.append(
            f"- [`{path.relative_to(SEED_ROOT).as_posix()}`]({relative_link(path)}) — {human_size(path.stat().st_size)}"
        )

    lines.extend(["", "## Company-world source files", ""])
    by_parent: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in sources:
        by_parent[str(Path(str(row["relative"])).parent)] .append(row)
    for parent in sorted(by_parent):
        lines.extend([f"### `{parent}`", ""])
        for row in by_parent[parent]:
            path = Path(row["path"])
            lines.append(
                f"- [`{path.name}`]({relative_link(path)}) — {metric_description(row)}"
            )
        lines.append("")

    lines.extend(
        [
            "## Measurement notes",
            "",
            "- PDF pages are read from the PDF page tree.",
            "- Word pages are measured after headless LibreOffice rendering and may differ slightly from Microsoft Word pagination.",
            "- Excel populated rows include model rows, summaries, assumptions, checks, and source data—not only independent records.",
            "- Accounting counts are physical table rows. A single economic event can appear in a header table, one or more detail tables, and the general ledger.",
            "- File sizes reflect compressed source files and do not measure analytical difficulty.",
            "",
            "Generated by `environment/tools/generate_seed_inventory.py`.",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    OUTPUT.write_text(build_markdown(), encoding="utf-8")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
